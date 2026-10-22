"""Build the small, generic form template."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = Path(__file__).resolve().parent
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def add_checkbox(paragraph, label, identifier):
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    control.append(properties)
    for tag, value in (("w:id", str(identifier)), ("w:alias", label)):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        properties.append(element)
    checkbox = OxmlElement("w14:checkbox")
    checked = OxmlElement("w14:checked")
    checked.set(f"{{{W14}}}val", "0")
    checkbox.append(checked)
    for kind, glyph in (("checkedState", "2612"), ("uncheckedState", "2610")):
        state = OxmlElement(f"w14:{kind}")
        state.set(f"{{{W14}}}val", glyph)
        state.set(f"{{{W14}}}font", "MS Gothic")
        checkbox.append(state)
    properties.append(checkbox)
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "☐"
    run.append(text)
    content.append(run)
    control.append(content)
    paragraph._p.append(control)
    paragraph.add_run(" " + label + "   ")


def normalize_docx(path):
    """Remove ZIP timestamps so the same document produces identical bytes."""
    buffer = BytesIO()
    with ZipFile(path) as source, ZipFile(buffer, "w") as output:
        for item in source.infolist():
            normalized = ZipInfo(item.filename, (1980, 1, 1, 0, 0, 0))
            normalized.compress_type = item.compress_type
            normalized.external_attr = item.external_attr
            output.writestr(normalized, source.read(item.filename))
    path.write_bytes(buffer.getvalue())


def main():
    document = Document()
    document.add_heading("Form", 0)
    document.add_paragraph(
        "Guidance: Enter a name and summary, then add a goal and metric for each item. Choose one service option."
    )

    fields = document.add_table(rows=2, cols=2)
    fields.style = "Table Grid"
    for row, label in enumerate(("Name", "Summary")):
        fields.cell(row, 0).text = label
        fields.cell(row, 1).text = "[...]"

    document.add_paragraph("Goals")
    goals = document.add_table(rows=3, cols=3)
    goals.style = "Table Grid"
    for col, title in enumerate(("Item", "Goal", "Metric")):
        goals.cell(0, col).text = title
    for row, label in enumerate(("Item A", "Item B"), 1):
        goals.cell(row, 0).text = label
        goals.cell(row, 1).text = ""
        goals.cell(row, 2).text = ""

    document.add_paragraph("Service option")
    options = document.add_paragraph()
    for identifier, label in ((1001, "Standard"), (1002, "Expedited")):
        add_checkbox(options, label, identifier)

    document.sections[0].top_margin = Inches(0.7)
    document.sections[0].bottom_margin = Inches(0.7)
    document.core_properties.title = "Form template"
    document.core_properties.subject = "Generic form example"
    document.core_properties.author = "DocXUnravel"
    output = HERE / "form-template.docx"
    document.save(output)
    normalize_docx(output)
    print(output)


if __name__ == "__main__":
    main()
