"""Build and run the fictional DocXUnravel Azure migration HLD demonstration."""

from __future__ import annotations

import copy
import hashlib
import json
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from docxunravel import create_fidelity_report, export_bundle, validate_docx_package
from docxunravel.template_contract import (
    populate_migration_template,
    template_contract_binding,
)
from docxunravel.visual_regression import compare_pdfs, verify_visuals

FICTIONAL_NOTICE = "FICTIONAL DEMONSTRATION — NOT A CUSTOMER DOCUMENT"


def _set_cell_shading(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _set_cell_text(cell: Any, text: str, *, bold: bool = False, size: int = 8) -> None:
    cell.text = text
    run = cell.paragraphs[0].runs[0]
    run.bold = bold
    run.font.size = Pt(size)


def _add_table(document: Any, headers: list[str], template_values: list[str]) -> Any:
    table = document.add_table(rows=2, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        _set_cell_text(table.cell(0, index), header, bold=True, size=8)
        _set_cell_shading(table.cell(0, index), "1F4E78")
        table.cell(0, index).paragraphs[0].runs[0].font.color.rgb = RGBColor(
            255, 255, 255
        )
    for index, value in enumerate(template_values):
        _set_cell_text(table.cell(1, index), value, size=8)
        _set_cell_shading(table.cell(1, index), "D9EAF7")
    return table


def _add_page_number(paragraph: Any) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def create_architecture_illustration(path: str | Path) -> Path:
    """Create a fictional pre-existing architecture image for the template."""

    path = Path(path)
    image = Image.new("RGB", (1600, 700), "white")
    drawing = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=26)
    title_font = ImageFont.load_default(size=34)
    drawing.text(
        (55, 35),
        "FICTIONAL TARGET ARCHITECTURE ILLUSTRATION",
        fill=(31, 78, 120),
        font=title_font,
    )
    boxes = [
        (70, 220, 360, 430, "Fictional users\nand laboratories"),
        (
            470,
            150,
            790,
            500,
            "Azure landing zone\n\nApp Service\nAKS\nPrivate endpoints",
        ),
        (900, 220, 1200, 430, "Managed data\n\nSQL Managed Instance\nAzure Files"),
        (1300, 220, 1530, 430, "Monitoring\nand security"),
    ]
    fills = [(224, 235, 247), (213, 232, 212), (255, 242, 204), (234, 209, 220)]
    for box, fill in zip(boxes, fills):
        left, top, right, bottom, label = box
        drawing.rounded_rectangle(
            (left, top, right, bottom),
            radius=24,
            fill=fill,
            outline=(31, 78, 120),
            width=5,
        )
        lines = label.split("\n")
        for line_index, line in enumerate(lines):
            drawing.text(
                (left + 28, top + 35 + line_index * 42),
                line,
                fill=(25, 45, 65),
                font=font,
            )
    for start, end in [
        ((360, 325), (470, 325)),
        ((790, 325), (900, 325)),
        ((1200, 325), (1300, 325)),
    ]:
        drawing.line((start, end), fill=(31, 78, 120), width=8)
        drawing.polygon(
            [(end[0], end[1]), (end[0] - 22, end[1] - 14), (end[0] - 22, end[1] + 14)],
            fill=(31, 78, 120),
        )
    drawing.text(
        (55, 630),
        "All names and topology shown are fictional demonstration content.",
        fill=(140, 40, 40),
        font=font,
    )
    image.save(path)
    return path


def create_hld_template(path: str | Path, illustration: str | Path) -> Path:
    path = Path(path)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header.paragraphs[
        0
    ].text = "DocXUnravel v0.4 | Fictional Azure Migration High-Level Design"
    section.header.paragraphs[0].runs[0].font.size = Pt(8)
    section.footer.paragraphs[0].add_run(FICTIONAL_NOTICE + " | ")
    _add_page_number(section.footer.paragraphs[0])

    styles = document.styles
    if "Demo Caption" not in styles:
        caption: Any = styles.add_style("Demo Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.size = Pt(8)
        caption.font.italic = True
        caption.font.color.rgb = RGBColor(89, 89, 89)

    title = document.add_paragraph(
        "Azure Migration HLD — {{ CUSTOMER }}", style="Title"
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Project Aurora | High-Level Design")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice = document.add_paragraph(FICTIONAL_NOTICE)
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.runs[0].bold = True
    notice.runs[0].font.color.rgb = RGBColor(192, 0, 0)
    document.add_paragraph(
        "This source DOCX is the authoritative template. DocXUnravel populates only contract-bound elements."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("Document Control", level=1)
    control = _add_table(
        document, ["Field", "Value"], ["Customer", "Template customer"]
    )
    for label, value in [
        ("Project", "Template project"),
        ("Sponsor", "Template sponsor"),
        ("Document owner", "Template owner"),
        ("Target start", "YYYY-MM-DD"),
        ("Target completion", "YYYY-MM-DD"),
    ]:
        cells = control.add_row().cells
        _set_cell_text(cells[0], label, bold=True)
        _set_cell_text(cells[1], value)

    document.add_heading("Executive Summary", level=1)
    document.add_paragraph(
        "Template executive summary — replace with approved migration scope."
    )
    document.add_page_break()

    document.add_heading("Current-State Assessment", level=1)
    document.add_paragraph(
        "The observations below are populated from validated migration-discovery records."
    )
    _add_table(
        document,
        ["Category", "Observation", "Impact"],
        ["Template category", "Template observation", "Template impact"],
    )
    document.add_heading("Application Inventory", level=1)
    _add_table(
        document,
        [
            "ID",
            "Application",
            "Owner",
            "Criticality",
            "Current hosting",
            "Data class",
            "Disposition",
        ],
        [
            "APP-ID",
            "Template application",
            "Template owner",
            "medium",
            "Template hosting",
            "internal",
            "rehost",
        ],
    )
    document.add_page_break()

    document.add_heading("Target Architecture", level=1)
    document.add_paragraph(
        "The existing illustration is a read-only package asset retained during population."
    )
    picture = document.add_picture(str(illustration), width=Inches(6.9))
    picture._inline.docPr.set(
        "descr", "Fictional Azure target architecture illustration"
    )
    document.add_paragraph(
        "Figure 1 — Fictional Azure target architecture (pre-existing template asset)",
        style="Demo Caption",
    )
    _add_table(
        document,
        ["Workload", "Azure service", "Region", "Rationale"],
        [
            "Template workload",
            "Template service",
            "Template region",
            "Template rationale",
        ],
    )
    document.add_page_break()

    document.add_heading("Migration Strategy", level=1)
    document.add_paragraph(
        "Migration waves are sequenced by dependency, risk and operational readiness."
    )
    _add_table(
        document,
        ["Wave", "Name", "Applications", "Start", "End", "Exit criteria"],
        [
            "1",
            "Template wave",
            "APP-ID",
            "YYYY-MM-DD",
            "YYYY-MM-DD",
            "Template criteria",
        ],
    )
    document.add_page_break()

    document.add_heading("Security Considerations", level=1)
    document.add_heading("Security requirements", level=2)
    _add_table(
        document,
        ["ID", "Requirement", "Priority"],
        ["SEC-ID", "Template security requirement", "must"],
    )
    document.add_heading("Networking requirements", level=2)
    _add_table(
        document,
        ["ID", "Requirement", "Priority"],
        ["NET-ID", "Template networking requirement", "must"],
    )
    document.add_page_break()

    document.add_heading("Risks and Assumptions", level=1)
    document.add_heading("Risks", level=2)
    _add_table(
        document,
        ["ID", "Description", "Likelihood", "Impact", "Mitigation"],
        ["RISK-ID", "Template risk", "medium", "medium", "Template mitigation"],
    )
    document.add_heading("Assumptions", level=2)
    _add_table(
        document,
        ["ID", "Description", "Validation owner"],
        ["ASM-ID", "Template assumption", "Template owner"],
    )
    document.add_page_break()

    document.add_heading("Appendix", level=1)
    document.add_heading("Outstanding design decisions", level=2)
    _add_table(
        document,
        ["ID", "Decision", "Owner", "Due date", "Status"],
        ["DD-ID", "Template decision", "Template owner", "YYYY-MM-DD", "open"],
    )
    document.add_heading("Source infrastructure inventory", level=2)
    _add_table(
        document,
        ["ID", "Name", "Category", "Location", "Quantity", "OS", "Specification"],
        [
            "INF-ID",
            "Template asset",
            "compute",
            "Template location",
            "1",
            "Template OS",
            "Template specification",
        ],
    )
    document.add_heading("Application dependencies", level=2)
    _add_table(
        document,
        ["Source", "Target", "Type", "Notes"],
        ["APP-A", "APP-B", "synchronous", "Template dependency"],
    )

    properties = document.core_properties
    properties.title = "Fictional Azure Migration High-Level Design"
    properties.subject = "DocXUnravel v0.4 semantic template demonstration"
    properties.author = "DocXUnravel demonstration generator"
    properties.comments = FICTIONAL_NOTICE
    document.save(str(path))
    return path


def create_hld_contract(profile: dict[str, Any]) -> dict[str, Any]:
    paragraph_ids = {
        element.get("text"): element["id"]
        for element in profile["body"]
        if element.get("type") == "paragraph"
    }
    tables = [element for element in profile["body"] if element.get("type") == "table"]

    def cell(table_index: int, row: int, column: int) -> str:
        return f"{tables[table_index]['id']}.r{row}.c{column}"

    fields: dict[str, Any] = {
        "customer_brand": {
            "type": "placeholder",
            "target": paragraph_ids["Azure Migration HLD — {{ CUSTOMER }}"],
            "token": "CUSTOMER",
            "required": True,
            "data_path": "/customer/name",
        },
        "control_customer": {
            "type": "table_cell",
            "target": cell(0, 1, 1),
            "required": True,
            "data_path": "/customer/name",
        },
        "control_project": {
            "type": "table_cell",
            "target": cell(0, 2, 1),
            "required": True,
            "data_path": "/project/name",
        },
        "control_sponsor": {
            "type": "table_cell",
            "target": cell(0, 3, 1),
            "required": True,
            "data_path": "/project/sponsor",
        },
        "control_owner": {
            "type": "table_cell",
            "target": cell(0, 4, 1),
            "required": True,
            "data_path": "/project/owner",
        },
        "control_start": {
            "type": "table_cell",
            "target": cell(0, 5, 1),
            "required": True,
            "data_path": "/project/target_start",
        },
        "control_completion": {
            "type": "table_cell",
            "target": cell(0, 6, 1),
            "required": True,
            "data_path": "/project/target_completion",
        },
        "executive_summary": {
            "type": "text",
            "target": paragraph_ids[
                "Template executive summary — replace with approved migration scope."
            ],
            "required": True,
            "data_path": "/project/scope",
        },
    }
    row_fields = [
        (
            "current_state_observations",
            1,
            "/current_state_observations",
            ["category", "observation", "impact"],
        ),
        (
            "application_inventory",
            2,
            "/application_inventory",
            [
                "id",
                "name",
                "owner",
                "business_criticality",
                "current_hosting",
                "data_classification",
                "migration_disposition",
            ],
        ),
        (
            "target_services",
            3,
            "/proposed_azure_target_services",
            ["workload", "azure_service", "region", "rationale"],
        ),
        (
            "migration_waves",
            4,
            "/migration_waves",
            [
                "wave",
                "name",
                "application_ids",
                "target_start",
                "target_end",
                "exit_criteria",
            ],
        ),
        (
            "security_requirements",
            5,
            "/security_and_networking_requirements/security",
            ["id", "requirement", "priority"],
        ),
        (
            "networking_requirements",
            6,
            "/security_and_networking_requirements/networking",
            ["id", "requirement", "priority"],
        ),
        (
            "risks",
            7,
            "/risks_and_assumptions/risks",
            ["id", "description", "likelihood", "impact", "mitigation"],
        ),
        (
            "assumptions",
            8,
            "/risks_and_assumptions/assumptions",
            ["id", "description", "validation_owner"],
        ),
        (
            "design_decisions",
            9,
            "/outstanding_design_decisions",
            ["id", "decision", "owner", "due_date", "status"],
        ),
        (
            "source_infrastructure",
            10,
            "/source_infrastructure_inventory",
            [
                "id",
                "name",
                "category",
                "location",
                "quantity",
                "operating_system",
                "specification",
            ],
        ),
        (
            "application_dependencies",
            11,
            "/application_dependencies",
            [
                "source_application_id",
                "target_application_id",
                "dependency_type",
                "notes",
            ],
        ),
    ]
    for name, table_index, data_path, columns in row_fields:
        fields[name] = {
            "type": "table_rows",
            "target": tables[table_index]["id"],
            "required": True,
            "template_row": 1,
            "columns": {column: index for index, column in enumerate(columns)},
            "data_path": data_path,
        }
    return {
        "contract_version": "0.1",
        "document_type": "azure_migration_hld",
        "binding": template_contract_binding(profile),
        "fields": fields,
    }


def generate_demo(output_directory: str | Path) -> dict[str, Any]:
    output_directory = Path(output_directory)
    output_directory.mkdir(parents=True, exist_ok=True)
    template = output_directory / "fictional-azure-migration-hld-template.docx"
    with tempfile.TemporaryDirectory(
        prefix="docxunravel-hld-illustration-"
    ) as illustration_directory:
        illustration = create_architecture_illustration(
            Path(illustration_directory) / "fictional-architecture-illustration.png"
        )
        create_hld_template(template, illustration)
    discovery_path = output_directory / "fictional-migration-discovery-0.1.json"
    discovery = json.loads(discovery_path.read_text(encoding="utf-8"))
    generated = output_directory / "fictional-azure-migration-hld-output.docx"
    contract_path = output_directory / "fictional-hld-template-contract-0.1.json"
    preview_path = output_directory / "fictional-hld-dry-run-preview.json"
    patch_path = output_directory / "fictional-hld-compiled-patch.json"
    fidelity_path = output_directory / "fictional-hld-fidelity-report.json"
    integrity_path = output_directory / "fictional-hld-package-integrity-report.json"
    visual_report_path = output_directory / "fictional-hld-visual-report.json"

    with tempfile.TemporaryDirectory(prefix="docxunravel-hld-demo-") as temporary:
        bundle = Path(temporary) / "bundle"
        profile = export_bundle(template, bundle)
        contract = create_hld_contract(profile)
        contract_path.write_text(
            json.dumps(contract, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        result = populate_migration_template(
            bundle,
            generated,
            contract,
            discovery,
            patch_output=patch_path,
            fidelity_output=fidelity_path,
        )
        preview_path.write_text(
            json.dumps(result["preview"], indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        visual_directory = Path(temporary) / "visual"
        initial_visual = verify_visuals(
            template,
            generated,
            visual_directory,
            expected_changed_pages=set(range(1, 50)),
            expected_difference_threshold=0.65,
        )
        changed_pages = {
            page["page"]
            for page in initial_visual["pages"]
            if page.get("difference_ratio", 0) > 0.001
            or not page.get("dimension_match", True)
        }
        visual = compare_pdfs(
            visual_directory / "original.pdf",
            visual_directory / "generated.pdf",
            visual_directory,
            expected_changed_pages=changed_pages,
            expected_difference_threshold=0.65,
        )
        visual["renderer"] = initial_visual["renderer"]
        visual["original_docx"] = str(template)
        visual["generated_docx"] = str(generated)
        compact_visual = copy.deepcopy(visual)
        compact_visual.pop("original_pdf", None)
        compact_visual.pop("generated_pdf", None)
        compact_visual["renderer"] = Path(compact_visual["renderer"]).name
        compact_visual["original_docx"] = template.name
        compact_visual["generated_docx"] = generated.name
        compact_visual["artifacts_retained"] = False
        for page in compact_visual["pages"]:
            page.pop("images", None)
        visual_report_path.write_text(
            json.dumps(compact_visual, indent=2), encoding="utf-8"
        )
        fidelity = create_fidelity_report(
            template,
            generated,
            bundle=bundle,
            patch=result["preview"]["patch"],
            visual_report=compact_visual,
        )
        fidelity_path.write_text(
            json.dumps(fidelity, indent=2, ensure_ascii=False), encoding="utf-8"
        )

        with (
            zipfile.ZipFile(template) as original,
            zipfile.ZipFile(generated) as updated,
        ):
            asset_checks = []
            for asset in profile["assets"]["items"]:
                location = asset.get("package_location")
                if not location:
                    continue
                asset_checks.append(
                    {
                        "asset_id": asset["id"],
                        "kind": asset["kind"],
                        "package_location": location,
                        "source_sha256": hashlib.sha256(
                            original.read(location)
                        ).hexdigest(),
                        "generated_sha256": hashlib.sha256(
                            updated.read(location)
                        ).hexdigest(),
                        "preserved": original.read(location) == updated.read(location),
                    }
                )
        integrity = {
            "schema_version": "1.0",
            "passed": fidelity["package_integrity"]["passed"]
            and fidelity["untouched_part_hash_preservation"]["passed"]
            and all(item["preserved"] for item in asset_checks),
            "source_sha256": hashlib.sha256(template.read_bytes()).hexdigest(),
            "generated_sha256": hashlib.sha256(generated.read_bytes()).hexdigest(),
            "generated_package": validate_docx_package(generated),
            "part_hash_preservation": fidelity["untouched_part_hash_preservation"],
            "read_only_assets": asset_checks,
        }
        integrity_path.write_text(
            json.dumps(integrity, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    return {
        "template": str(template),
        "contract": str(contract_path),
        "discovery": str(discovery_path),
        "generated": str(generated),
        "preview": str(preview_path),
        "patch": str(patch_path),
        "fidelity": str(fidelity_path),
        "integrity": str(integrity_path),
        "visual": str(visual_report_path),
        "passed": integrity["passed"]
        and fidelity["passed"]
        and compact_visual["passed"],
    }


if __name__ == "__main__":
    print(json.dumps(generate_demo(Path(__file__).parent), indent=2))
