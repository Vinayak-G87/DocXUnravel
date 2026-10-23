import contextlib
import hashlib
import io
import json
import tempfile
import tomllib
import unittest
import zipfile
from importlib.metadata import version as distribution_version
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

import docxunravel
from docxunravel import (
    create_fidelity_report,
    export_bundle,
    preview_bundle,
    profile_document,
    render_bundle,
    render_document,
    validate_docx_package,
)
from docxunravel.cli import main
from docxunravel.visual_regression import (
    compare_pdfs,
    verify_visuals,
    visual_dependencies,
)


class DocXUnravelTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "template.docx"

        document = Document()
        title = document.add_paragraph(style="Title")
        title.add_run("Architecture for ").bold = True
        title.add_run("{{ CUST")
        title.add_run("OMER }}").italic = True

        summary = document.add_paragraph("Original summary")
        summary.runs[0].font.size = Pt(11)

        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Component"
        table.cell(0, 1).text = "Technology"
        table.cell(1, 0).text = "Compute"
        table.cell(1, 1).text = "{{TECH}}"

        header = document.sections[0].header
        header.paragraphs[0].text = "Confidential - {{CUSTOMER}}"
        document.save(self.source)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_release_version_matches_package_metadata(self):
        with (Path(__file__).resolve().parents[1] / "pyproject.toml").open(
            "rb"
        ) as file:
            version = tomllib.load(file)["project"]["version"]
        self.assertEqual(docxunravel.__version__, version)
        self.assertEqual(distribution_version("docxunravel"), version)

    def test_profile_and_render(self):
        profile = profile_document(self.source)
        self.assertEqual(profile["schema_version"], "1.1")
        self.assertIn("CUSTOMER", profile["summary"]["tokens"])
        self.assertIn("TECH", profile["summary"]["tokens"])
        self.assertEqual(profile["body"][0]["style"], "Title")
        self.assertEqual(profile["body"][2]["type"], "table")
        self.assertIn("body.p1", profile["replacement_contract"]["elements"])

        output = self.root / "generated.docx"
        result = render_document(
            self.source,
            output,
            {
                "tokens": {"CUSTOMER": "Acme", "TECH": "Kubernetes"},
                "elements": {"body.p1": "Generated system summary"},
            },
            verify_sha256=profile["source"]["sha256"],
        )

        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[0].text, "Architecture for Acme")
        self.assertEqual(rendered.paragraphs[1].text, "Generated system summary")
        self.assertEqual(rendered.tables[0].cell(1, 1).text, "Kubernetes")
        self.assertEqual(
            rendered.sections[0].header.paragraphs[0].text, "Confidential - Acme"
        )
        self.assertEqual(result["changed_elements"], 1)
        self.assertEqual(result["token_replacements"], 3)

    def test_explicit_literal_placeholders_in_cells_and_split_runs(self):
        document = Document(self.source)
        paragraph = document.add_paragraph("Owner: ")
        paragraph.add_run("[Enter ")
        paragraph.add_run("owner]")
        document.tables[0].cell(1, 0).text = "<Lead>"
        document.add_paragraph("[Keep this] is not a placeholder")
        document.save(self.source)

        bundle = self.root / "literal-bundle"
        profile = export_bundle(
            self.source,
            bundle,
            placeholders={"OWNER": "[Enter owner]", "LEAD": "<Lead>"},
        )
        self.assertEqual(profile["placeholder_locations"]["OWNER"], ["body.p2"])
        self.assertEqual(profile["placeholder_locations"]["LEAD"], ["body.t0.r1.c0.p0"])
        self.assertIn("OWNER", profile["summary"]["tokens"])
        self.assertIn("LEAD", profile["replacement_contract"]["tokens"])
        self.assertEqual(profile["custom_placeholders"]["LEAD"], "<Lead>")
        patch = {
            "contract": profile["patch_contract"],
            "tokens": {"OWNER": "Avery Patel", "LEAD": "Migration lead"},
        }
        self.assertTrue(preview_bundle(bundle, patch)["valid"])
        output = self.root / "literal-output.docx"
        result = render_bundle(bundle, output, patch)
        self.assertEqual(result["token_replacements"], 2)
        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[2].text, "Owner: Avery Patel")
        self.assertEqual(rendered.tables[0].cell(1, 0).text, "Migration lead")
        self.assertEqual(
            rendered.paragraphs[3].text, "[Keep this] is not a placeholder"
        )
        self.assertTrue(
            create_fidelity_report(self.source, output, bundle=bundle, patch=patch)[
                "passed"
            ]
        )

        with self.assertRaisesRegex(ValueError, "not present"):
            preview_bundle(
                bundle,
                {"contract": profile["patch_contract"], "tokens": {"UNKNOWN": "x"}},
            )
        with self.assertRaisesRegex(ValueError, "not found"):
            profile_document(self.source, placeholders={"NONE": "[Missing]"})
        with self.assertRaisesRegex(ValueError, "overlap"):
            profile_document(self.source, placeholders={"A": "<Lead>", "B": "Lead"})

    def test_replacement_containing_its_own_placeholder_is_not_rescanned(self):
        document = Document(self.source)
        document.add_paragraph("[X] and [X]")
        document.save(self.source)
        bundle = self.root / "self-reference-bundle"
        profile = export_bundle(self.source, bundle, placeholders={"X": "[X]"})
        patch = {
            "contract": profile["patch_contract"],
            "tokens": {"X": "keep [X]", "TECH": "still {{TECH}}"},
        }
        output = self.root / "self-reference.docx"
        result = render_bundle(bundle, output, patch)
        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[2].text, "keep [X] and keep [X]")
        self.assertEqual(rendered.tables[0].cell(1, 1).text, "still {{TECH}}")
        self.assertEqual(result["token_replacements"], 3)
        self.assertTrue(
            create_fidelity_report(self.source, output, bundle=bundle, patch=patch)[
                "passed"
            ]
        )

    def test_token_replacement_preserves_spaces_across_runs(self):
        document = Document(self.source)
        paragraph = document.add_paragraph()
        paragraph.add_run("[Enter ")
        paragraph.add_run("owner]")
        document.tables[0].cell(1, 0).text = "<Lead>"
        document.save(self.source)
        bundle = self.root / "spaces-bundle"
        profile = export_bundle(
            self.source,
            bundle,
            placeholders={"OWNER": "[Enter owner]", "LEAD": "<Lead>"},
        )
        output = self.root / "spaces.docx"
        render_bundle(
            bundle,
            output,
            {
                "contract": profile["patch_contract"],
                "tokens": {"OWNER": " leading and trailing ", "LEAD": " padded "},
            },
        )
        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[2].text, " leading and trailing ")
        self.assertEqual(rendered.tables[0].cell(1, 0).text, " padded ")
        with zipfile.ZipFile(output) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        xml_space = "{http://www.w3.org/XML/1998/namespace}space"
        text_nodes = root.iter(qn("w:t"))
        preserved = [
            node.text for node in text_nodes if node.get(xml_space) == "preserve"
        ]
        self.assertIn(" leading and trailing ", preserved)
        self.assertIn(" padded ", preserved)

    def test_fidelity_ignores_new_literal_outside_original_placeholder_locations(self):
        document = Document(self.source)
        document.add_paragraph("[X]")
        document.save(self.source)
        bundle = self.root / "literal-location-bundle"
        profile = export_bundle(self.source, bundle, placeholders={"X": "[X]"})
        patch = {
            "contract": profile["patch_contract"],
            "elements": {
                "body.p1": {"expected": "Original summary", "value": "Keep [X] here"}
            },
            "tokens": {"X": "Ada"},
        }
        output = self.root / "literal-location.docx"
        render_bundle(bundle, output, patch)
        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[1].text, "Keep [X] here")
        self.assertEqual(rendered.paragraphs[2].text, "Ada")
        self.assertTrue(
            create_fidelity_report(self.source, output, bundle=bundle, patch=patch)[
                "passed"
            ]
        )

    def test_literal_backslashes_in_token_values_are_not_regex_expressions(self):
        bundle = self.root / "backslash-bundle"
        profile = export_bundle(self.source, bundle)
        patch = {
            "contract": profile["patch_contract"],
            "tokens": {"CUSTOMER": r"C:\Users\Demo"},
        }
        output = self.root / "backslashes.docx"
        render_bundle(bundle, output, patch)
        self.assertIn(r"C:\Users\Demo", Document(output).paragraphs[0].text)
        self.assertTrue(
            create_fidelity_report(self.source, output, bundle=bundle, patch=patch)[
                "passed"
            ]
        )

    def test_explicit_only_ignores_unapproved_brace_tokens(self):
        bundle = self.root / "explicit-bundle"
        profile = export_bundle(
            self.source,
            bundle,
            auto_placeholders=False,
            placeholders={"TITLE_CUSTOMER": "{{ CUSTOMER }}"},
        )
        self.assertFalse(profile["auto_placeholders"])
        self.assertEqual(profile["summary"]["tokens"], ["TITLE_CUSTOMER"])
        self.assertEqual(
            profile["placeholder_locations"]["TITLE_CUSTOMER"], ["body.p0"]
        )
        with self.assertRaisesRegex(ValueError, "not present"):
            preview_bundle(
                bundle,
                {
                    "contract": profile["patch_contract"],
                    "tokens": {"TECH": "Unapproved"},
                },
            )
        patch = {
            "contract": profile["patch_contract"],
            "tokens": {"TITLE_CUSTOMER": "Fictional Customer"},
        }
        output = self.root / "explicit.docx"
        render_bundle(bundle, output, patch)
        document = Document(output)
        self.assertEqual(
            document.paragraphs[0].text, "Architecture for Fictional Customer"
        )
        self.assertEqual(
            document.sections[0].header.paragraphs[0].text,
            "Confidential - {{CUSTOMER}}",
        )
        self.assertEqual(document.tables[0].cell(1, 1).text, "{{TECH}}")

    def test_cli_placeholder_flags_extract_and_profile(self):
        document = Document(self.source)
        document.add_paragraph("[Replace this with owner]")
        document.save(self.source)
        bundle = self.root / "cli-bundle"
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(
                main(
                    [
                        "extract",
                        str(self.source),
                        str(bundle),
                        "--placeholder",
                        "OWNER=[Replace this with owner]",
                    ]
                ),
                0,
            )
        profile = json.loads((bundle / "profile.json").read_text(encoding="utf-8"))
        self.assertEqual(profile["placeholder_locations"]["OWNER"], ["body.p2"])
        output = self.root / "cli-profile.json"
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(
                main(
                    [
                        "profile",
                        str(self.source),
                        str(output),
                        "--placeholder",
                        "OWNER=[Replace this with owner]",
                    ]
                ),
                0,
            )
        self.assertEqual(
            json.loads(output.read_text())["placeholder_locations"]["OWNER"],
            ["body.p2"],
        )
        explicit = self.root / "explicit-profile.json"
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(
                main(
                    [
                        "profile",
                        str(self.source),
                        str(explicit),
                        "--explicit-placeholders-only",
                    ]
                ),
                0,
            )
        self.assertEqual(json.loads(explicit.read_text())["summary"]["tokens"], [])
        explicit_bundle = self.root / "explicit-cli-bundle"
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(
                main(
                    [
                        "extract",
                        str(self.source),
                        str(explicit_bundle),
                        "--explicit-placeholders-only",
                        "--placeholder",
                        "OWNER=[Replace this with owner]",
                    ]
                ),
                0,
            )
        explicit_profile = json.loads((explicit_bundle / "profile.json").read_text())
        self.assertEqual(explicit_profile["summary"]["tokens"], ["OWNER"])
        with self.assertRaisesRegex(ValueError, "Duplicate --placeholder"):
            with contextlib.redirect_stdout(io.StringIO()):
                main(
                    [
                        "profile",
                        str(self.source),
                        str(output),
                        "--placeholder",
                        "OWNER=first",
                        "--placeholder",
                        "OWNER=second",
                    ]
                )

    def test_lossless_bundle_and_structural_rendering(self):
        bundle = self.root / "bundle"
        profile = export_bundle(self.source, bundle)
        self.assertTrue((bundle / "package" / "word" / "document.xml").exists())
        self.assertGreater(profile["package"]["entry_count"], 5)
        self.assertIn("locator", profile["body"][0])

        output = self.root / "bundle-generated.docx"
        result = render_bundle(
            bundle,
            output,
            {
                "contract": profile["patch_contract"],
                "tokens": {"CUSTOMER": "Acme", "TECH": "EKS"},
                "elements": {"body.p1": "Bundle-generated summary"},
                "operations": [
                    {
                        "op": "insert_after",
                        "target": "body.p1",
                        "element": {
                            "type": "paragraph",
                            "style": "Normal",
                            "text": "Inserted architecture detail",
                        },
                    },
                    {
                        "op": "insert_after",
                        "target": "body.p0",
                        "element": {
                            "type": "table",
                            "style": "Table Grid",
                            "rows": [
                                ["Layer", "Details"],
                                [
                                    "Compute",
                                    {
                                        "elements": [
                                            {"type": "paragraph", "text": "Cluster"},
                                            {
                                                "type": "table",
                                                "style": "Table Grid",
                                                "rows": [["Nodes", "3"]],
                                            },
                                        ]
                                    },
                                ],
                            ],
                        },
                    },
                ],
            },
        )

        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[0].text, "Architecture for Acme")
        self.assertEqual(rendered.paragraphs[1].text, "Bundle-generated summary")
        self.assertEqual(rendered.paragraphs[2].text, "Inserted architecture detail")
        self.assertEqual(rendered.tables[0].cell(1, 1).tables[0].cell(0, 1).text, "3")
        self.assertEqual(rendered.tables[1].cell(1, 1).text, "EKS")
        self.assertEqual(result["structural_changes"], 2)
        self.assertIn("word/document.xml", result["modified_parts"])
        self.assertTrue(result["integrity"]["valid"])
        self.assertTrue(validate_docx_package(output)["valid"])

    def test_patch_preview_conflicts_and_atomic_output(self):
        bundle = self.root / "preview-bundle"
        profile = export_bundle(self.source, bundle)
        patch = {
            "contract": profile["patch_contract"],
            "elements": {
                "body.p1": {
                    "expected": "Original summary",
                    "value": "New summary",
                    "format_policy": "preserve_first_run",
                }
            },
        }
        preview = preview_bundle(bundle, patch)
        self.assertTrue(preview["valid"])
        self.assertEqual(preview["changes"][0]["before"], "Original summary")

        with self.assertRaises(ValueError):
            render_bundle(
                bundle,
                self.root / "conflict.docx",
                {
                    **patch,
                    "contract": {**profile["patch_contract"], "source_sha256": "bad"},
                },
            )
        self.assertFalse((self.root / "conflict.docx").exists())

        with self.assertRaises(ValueError):
            preview_bundle(
                bundle,
                {
                    "contract": profile["patch_contract"],
                    "elements": {
                        "body.p1": {"expected": "stale", "value": "New summary"}
                    },
                },
            )

    def test_no_op_patch_preserves_every_package_part(self):
        bundle = self.root / "noop-bundle"
        profile = export_bundle(self.source, bundle)
        output = self.root / "noop.docx"
        result = render_bundle(
            bundle,
            output,
            {"contract": profile["patch_contract"]},
        )
        self.assertEqual(result["modified_parts"], [])
        with (
            zipfile.ZipFile(self.source) as original,
            zipfile.ZipFile(output) as generated,
        ):
            self.assertEqual(set(original.namelist()), set(generated.namelist()))
            for name in original.namelist():
                self.assertEqual(original.read(name), generated.read(name), name)

    def test_mixed_run_formatting_policies(self):
        bundle = self.root / "format-bundle"
        profile = export_bundle(self.source, bundle)
        contract = profile["patch_contract"]

        first_run_output = self.root / "first-run.docx"
        result = render_bundle(
            bundle,
            first_run_output,
            {
                "contract": contract,
                "elements": {
                    "body.p0": {
                        "value": "Replacement title",
                        "format_policy": "preserve_first_run",
                    }
                },
            },
        )
        first_run = Document(first_run_output).paragraphs[0]
        self.assertEqual(first_run.text, "Replacement title")
        self.assertTrue(first_run.runs[0].bold)
        self.assertTrue(result["warnings"])

        common_output = self.root / "common.docx"
        render_bundle(
            bundle,
            common_output,
            {
                "contract": contract,
                "elements": {
                    "body.p0": {
                        "value": "Common-format title",
                        "format_policy": "preserve_common_formatting",
                    }
                },
            },
        )
        common = Document(common_output).paragraphs[0]
        self.assertEqual(common.text, "Common-format title")
        self.assertIsNone(common.runs[0].bold)

        text_only_output = self.root / "text-only.docx"
        render_bundle(
            bundle,
            text_only_output,
            {
                "contract": contract,
                "elements": {
                    "body.p0": {
                        "value": "Text distributed over existing formatted runs",
                        "format_policy": "replace_text_only",
                    }
                },
            },
        )
        text_only = Document(text_only_output).paragraphs[0]
        self.assertEqual(
            text_only.text, "Text distributed over existing formatted runs"
        )
        self.assertTrue(text_only.runs[0].bold)
        self.assertTrue(text_only.runs[2].italic)

        with self.assertRaises(ValueError):
            render_bundle(
                bundle,
                self.root / "rejected.docx",
                {
                    "contract": contract,
                    "elements": {
                        "body.p0": {
                            "value": "Unsafe replacement",
                            "format_policy": "reject_mixed_runs",
                        }
                    },
                },
            )

    def test_run_level_replacement_preserves_other_runs(self):
        bundle = self.root / "run-bundle"
        profile = export_bundle(self.source, bundle)
        output = self.root / "run-output.docx"
        result = render_bundle(
            bundle,
            output,
            {
                "contract": profile["patch_contract"],
                "runs": {
                    "body.p0": [
                        {
                            "index": 0,
                            "expected": "Architecture for ",
                            "value": "Platform for ",
                        }
                    ]
                },
            },
        )
        paragraph = Document(output).paragraphs[0]
        self.assertEqual(paragraph.text, "Platform for {{ CUSTOMER }}")
        self.assertTrue(paragraph.runs[0].bold)
        self.assertEqual(paragraph.runs[1].text, "{{ CUST")
        self.assertTrue(paragraph.runs[2].italic)
        self.assertEqual(result["changed_runs"], 1)

    def test_stale_profile_and_structural_conflicts_are_rejected(self):
        source_bundle = self.root / "source-conflict-bundle"
        source_profile = export_bundle(self.source, source_bundle)
        document_xml = source_bundle / "package" / "word" / "document.xml"
        document_xml.write_bytes(document_xml.read_bytes() + b" ")
        with self.assertRaisesRegex(ValueError, "package part has changed"):
            preview_bundle(
                source_bundle,
                {"contract": source_profile["patch_contract"]},
            )
        source_profile_path = source_bundle / "profile.json"
        modified_profile = json.loads(source_profile_path.read_text(encoding="utf-8"))
        changed_bytes = document_xml.read_bytes()
        document_entry = next(
            entry
            for entry in modified_profile["package"]["entries"]
            if entry["path"] == "word/document.xml"
        )
        document_entry["sha256"] = hashlib.sha256(changed_bytes).hexdigest()
        document_entry["size"] = len(changed_bytes)
        source_profile_path.write_text(json.dumps(modified_profile), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Semantic profile is stale"):
            preview_bundle(
                source_bundle,
                {"contract": source_profile["patch_contract"]},
            )

        conflict_bundle = self.root / "conflict-bundle"
        profile = export_bundle(self.source, conflict_bundle)
        contract = profile["patch_contract"]
        with self.assertRaises(ValueError):
            preview_bundle(
                conflict_bundle,
                {
                    "contract": contract,
                    "operations": [
                        {
                            "op": "insert_before",
                            "target": "body.p1",
                            "element": {"type": "paragraph", "text": "Before"},
                        },
                        {
                            "op": "delete",
                            "target": "body.p1",
                        },
                    ],
                },
            )
        with self.assertRaisesRegex(ValueError, "inside deleted element"):
            preview_bundle(
                conflict_bundle,
                {
                    "contract": contract,
                    "elements": {"body.t0.r1.c1.p0": "Conflicting cell edit"},
                    "operations": [{"op": "delete", "target": "body.t0"}],
                },
            )

        profile_path = conflict_bundle / "profile.json"
        stale_profile = json.loads(profile_path.read_text(encoding="utf-8"))
        stale_profile["body"][1]["text"] = "Tampered semantic profile"
        profile_path.write_text(json.dumps(stale_profile), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Semantic profile is stale"):
            preview_bundle(conflict_bundle, {"contract": contract})

    def test_control_nodes_are_rejected_for_full_and_run_replacement(self):
        source = self.root / "breaks.docx"
        document = Document()
        run = document.add_paragraph().add_run("before")
        run.add_break()
        run.add_text("after")
        document.save(source)
        bundle = self.root / "break-bundle"
        profile = export_bundle(source, bundle)
        self.assertIn(
            "line_or_page_break", profile["body"][0]["unsupported_edit_structures"]
        )
        with self.assertRaisesRegex(ValueError, "Unsupported OOXML structures"):
            preview_bundle(
                bundle,
                {
                    "contract": profile["patch_contract"],
                    "elements": {"body.p0": "replacement"},
                },
            )
        with self.assertRaisesRegex(ValueError, "cannot be safely edited at run level"):
            preview_bundle(
                bundle,
                {
                    "contract": profile["patch_contract"],
                    "runs": {"body.p0": [{"index": 0, "value": "replacement"}]},
                },
            )

    @unittest.skipUnless(
        visual_dependencies()["imaging"],
        "PyMuPDF and Pillow are optional visual-test dependencies",
    )
    def test_pdf_visual_comparison_without_libreoffice(self):
        import pymupdf

        original = self.root / "original.pdf"
        pdf = pymupdf.open()
        page = pdf.new_page(width=200, height=200)
        page.insert_text((24, 50), "Sample text")
        pdf.save(str(original))
        pdf.close()

        report = compare_pdfs(
            original, original, self.root / "pdf-comparison", no_op=True
        )
        self.assertTrue(report["passed"])
        self.assertEqual(report["pages"][0]["status"], "no_change")
        self.assertTrue((self.root / "pdf-comparison" / "page-001-diff.png").is_file())

    @unittest.skipUnless(
        visual_dependencies()["available"],
        "LibreOffice, PyMuPDF, and Pillow are optional visual-test dependencies",
    )
    def test_optional_visual_regression_for_no_op_output(self):
        def verify_or_skip(*args, **kwargs):
            try:
                return verify_visuals(*args, **kwargs)
            except RuntimeError as error:
                if str(error).startswith(
                    (
                        "DOCX-to-PDF rendering failed for ",
                        "DOCX-to-PDF rendering timed out after ",
                    )
                ) and f" for {self.source}" in str(error):
                    self.skipTest(
                        f"LibreOffice cannot render the source fixture: {error}"
                    )
                raise
            except OSError as error:
                if error.filename == visual_dependencies()["renderer"]:
                    self.skipTest(f"LibreOffice could not be started: {error}")
                raise

        bundle = self.root / "visual-bundle"
        profile = export_bundle(self.source, bundle)
        output = self.root / "visual-noop.docx"
        render_bundle(bundle, output, {"contract": profile["patch_contract"]})
        report = verify_or_skip(
            self.source,
            output,
            self.root / "visual-report",
            difference_threshold=0.0,
            expected_difference_threshold=0.0,
            no_op=True,
        )
        self.assertTrue(report["passed"])
        self.assertEqual(report["mode"], "no-op")
        self.assertTrue((self.root / "visual-report" / "visual-report.json").exists())
        fidelity = create_fidelity_report(
            self.source,
            output,
            bundle=bundle,
            patch={"contract": profile["patch_contract"]},
            visual_report=report,
        )
        self.assertTrue(fidelity["passed"])
        self.assertTrue(fidelity["visual_differences"]["performed"])

        changed_output = self.root / "visual-changed.docx"
        render_bundle(
            bundle,
            changed_output,
            {
                "contract": profile["patch_contract"],
                "elements": {"body.p1": "A substantially different generated summary."},
            },
        )
        bounded = verify_or_skip(
            self.source,
            changed_output,
            self.root / "visual-bounded-report",
            difference_threshold=0.0,
            expected_difference_threshold=0.0,
            expected_changed_pages={1},
        )
        self.assertFalse(bounded["passed"])
        self.assertIn(1, bounded["unexpected_changed_pages"])

    def test_repository_regression_fixture_with_patch(self):
        project_root = Path(__file__).resolve().parents[1]
        fixture = (
            project_root
            / "tests"
            / "fixtures"
            / "DocXUnravel_Regression_Test_Fixture.docx"
        )
        patch_path = project_root / "tests" / "fixtures" / "regression-patch.json"
        bundle = self.root / "regression-bundle"
        output = self.root / "regression-output.docx"

        profile = export_bundle(fixture, bundle)
        with patch_path.open(encoding="utf-8") as handle:
            patch = json.load(handle)
        self.assertEqual(patch["contract"], profile["patch_contract"])
        self.assertTrue(preview_bundle(bundle, patch)["valid"])
        with self.assertRaisesRegex(ValueError, "Unsupported OOXML structures"):
            preview_bundle(
                bundle,
                {
                    "contract": profile["patch_contract"],
                    "elements": {"body.p12": "Unsafe hyperlink replacement"},
                },
            )

        result = render_bundle(bundle, output, patch)
        rendered = Document(output)
        body_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
        self.assertIn(
            "Customer: Acme Corporation  |  Platform: Order Platform", body_text
        )
        self.assertIn("Inserted regression test paragraph", body_text)
        self.assertIn("event-driven architecture", body_text)
        self.assertNotIn("TARGET_DELETE:", body_text)
        self.assertEqual(len(rendered.tables), 2)
        self.assertEqual(len(rendered.inline_shapes), 1)
        self.assertTrue(result["integrity"]["valid"])
        self.assertEqual(result["token_replacements"], 4)
        self.assertEqual(result["structural_changes"], 3)
        modified = set(result["modified_parts"])
        with zipfile.ZipFile(output) as generated:
            for entry in profile["package"]["entries"]:
                if entry["path"] not in modified:
                    self.assertEqual(
                        hashlib.sha256(generated.read(entry["path"])).hexdigest(),
                        entry["sha256"],
                        entry["path"],
                    )
            self.assertIn("word/media/image1.png", generated.namelist())
            relationships = generated.read("word/_rels/document.xml.rels")
            self.assertIn(b"relationships/hyperlink", relationships)
            self.assertIn(b"relationships/image", relationships)

    def test_hash_and_unknown_element_are_rejected(self):
        with self.assertRaises(ValueError):
            render_document(
                self.source, self.root / "bad.docx", {}, verify_sha256="bad"
            )

        with self.assertRaises(KeyError):
            render_document(
                self.source,
                self.root / "bad.docx",
                {"elements": {"body.p999": "no"}},
            )


if __name__ == "__main__":
    unittest.main()
