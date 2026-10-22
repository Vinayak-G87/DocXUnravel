"""Deterministic regression-fixture builders for DocXUnravel tests.

This module is test support. It deliberately creates OOXML constructs that the
production patch API must detect and reject rather than edit unsafely.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


def _add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def create_complex_fixture(path: str | Path, logo: str | Path | None = None) -> Path:
    path = Path(path)
    document = Document()
    document.settings.odd_and_even_pages_header_footer = True
    first = document.sections[0]
    first.different_first_page_header_footer = True
    first.header.paragraphs[0].text = "Section One Header"
    first.first_page_header.paragraphs[0].text = "Section One First Page"
    first.even_page_header.paragraphs[0].text = "Section One Even Header"
    first.footer.paragraphs[0].text = "Section One Footer"

    document.add_heading("Complex Architecture", level=0)
    bookmarked = document.add_paragraph()
    bookmarked.add_run("Bookmarked architecture overview")
    _add_bookmark(bookmarked, "ArchitectureOverview", 42)
    comment_run = document.add_paragraph().add_run(
        "This statement has a review comment."
    )
    document.add_comment(
        comment_run, "Verify this architecture decision.", "Reviewer", "RV"
    )

    for index in range(4):
        document.add_paragraph(f"Numbered requirement {index + 1}", style="List Number")
    for index in range(3):
        document.add_paragraph(f"Design consideration {index + 1}", style="List Bullet")

    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(0, 2).text = "Status"
    table.cell(1, 0).text = "Compute"
    nested_cell = table.cell(1, 1)
    nested_cell.text = ""
    nested = nested_cell.add_table(rows=2, cols=2)
    nested.style = "Table Grid"
    nested.cell(0, 0).text = "Runtime"
    nested.cell(0, 1).text = "Python"
    nested.cell(1, 0).text = "Replicas"
    nested.cell(1, 1).text = "3"
    table.cell(1, 2).text = "Approved"
    table.cell(2, 0).text = "Data"
    table.cell(2, 1).text = "PostgreSQL"
    table.cell(2, 2).text = "Approved"

    if logo is not None:
        document.add_picture(str(logo), width=Inches(0.6))

    second = document.add_section(WD_SECTION.NEW_PAGE)
    second.header.is_linked_to_previous = False
    second.footer.is_linked_to_previous = False
    second.header.paragraphs[0].text = "Section Two Header"
    second.footer.paragraphs[0].text = "Section Two Footer"
    second.orientation = WD_ORIENT.LANDSCAPE
    second.page_width, second.page_height = second.page_height, second.page_width
    second.left_margin = Inches(0.7)
    second.right_margin = Inches(0.7)
    document.add_heading("Landscape Deployment Matrix", level=1)
    matrix = document.add_table(rows=8, cols=5)
    matrix.style = "Table Grid"
    for row_index, row in enumerate(matrix.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.text = f"R{row_index + 1}C{cell_index + 1}"

    third = document.add_section(WD_SECTION.NEW_PAGE)
    third.header.is_linked_to_previous = False
    third.footer.is_linked_to_previous = False
    third.header.paragraphs[0].text = "Section Three Header"
    third.footer.paragraphs[0].text = "Section Three Footer"
    document.add_heading("Revision Test", level=1)
    document.add_paragraph("TRACKED_INSERT_MARKER")
    document.add_paragraph("TRACKED_DELETE_MARKER")
    document.save(str(path))
    _inject_tracked_revisions(path)
    return path


def _inject_tracked_revisions(path: Path) -> None:
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info)) for info in source.infolist()]
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    updated = []
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for info, data in entries:
        if info.filename != "word/document.xml":
            updated.append((info, data))
            continue
        root = etree.fromstring(data, parser)
        for marker, wrapper_name in (
            ("TRACKED_INSERT_MARKER", "ins"),
            ("TRACKED_DELETE_MARKER", "del"),
        ):
            for paragraph in root.xpath(".//w:p", namespaces=namespace):
                text = "".join(paragraph.xpath(".//w:t/text()", namespaces=namespace))
                if text != marker:
                    continue
                run = paragraph.find(qn("w:r"))
                if run is None:
                    continue
                wrapper = etree.Element(qn(f"w:{wrapper_name}"))
                wrapper.set(qn("w:id"), "100" if wrapper_name == "ins" else "101")
                wrapper.set(qn("w:author"), "Fixture Builder")
                wrapper.set(qn("w:date"), "2026-01-01T00:00:00Z")
                run.addprevious(wrapper)
                paragraph.remove(run)
                if wrapper_name == "del":
                    for text_node in run.iter(qn("w:t")):
                        text_node.tag = qn("w:delText")
                wrapper.append(run)
                break
        updated.append(
            (
                info,
                etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                ),
            )
        )
    with zipfile.ZipFile(path, "w") as destination:
        for info, data in updated:
            destination.writestr(info, data)


def create_performance_fixture(
    path: str | Path,
    *,
    paragraph_count: int,
    table_count: int,
    rows_per_table: int = 5,
) -> Path:
    path = Path(path)
    document = Document()
    for index in range(paragraph_count):
        paragraph = document.add_paragraph()
        normal = paragraph.add_run(f"Paragraph {index}: architecture component ")
        normal.font.size = Pt(10)
        emphasized = paragraph.add_run(f"service-{index % 25}")
        emphasized.bold = True
        paragraph.add_run(" uses managed infrastructure and audited interfaces.")
    for table_index in range(table_count):
        table = document.add_table(rows=rows_per_table, cols=4)
        table.style = "Table Grid"
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.text = f"T{table_index} R{row_index} C{column_index}"
    document.save(str(path))
    return path
