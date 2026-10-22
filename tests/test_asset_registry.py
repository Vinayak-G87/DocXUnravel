import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches

from docxunravel import export_bundle, render_bundle
from docxunravel.schemas import schema_path


class AssetRegistryTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "asset-template.docx"
        document = Document()
        document.add_paragraph("Editable summary")
        document.add_picture(
            str(Path(__file__).resolve().parents[1] / "assets" / "docxunravel-logo.png"),
            width=Inches(1.25),
        )
        document.save(str(self.source))
        self.bundle = self.root / "bundle"
        self.profile = export_bundle(self.source, self.bundle)

    def tearDown(self):
        self.temporary.cleanup()

    def test_profile_has_hash_backed_read_only_image_metadata(self):
        registry = self.profile["assets"]
        self.assertEqual(registry["schema_version"], "0.1")
        self.assertTrue(registry["read_only"])
        self.assertFalse(registry["binary_content_included"])
        images = [item for item in registry["items"] if item["kind"] == "image"]
        self.assertEqual(len(images), 1)
        image = images[0]
        self.assertEqual(image["media_type"], "image/png")
        self.assertTrue(image["package_location"].startswith("word/media/"))
        self.assertEqual(len(image["sha256"]), 64)
        self.assertEqual(image["owning_part"], "word/document.xml")
        self.assertTrue(image["relationship_id"].startswith("rId"))
        self.assertEqual(len(image["placements"]), 1)
        dimensions = image["placements"][0]["dimensions"]
        self.assertGreater(dimensions["width_emu"], 0)
        self.assertGreater(dimensions["height_emu"], 0)
        serialized = json.dumps(registry)
        self.assertNotIn("base64", serialized.lower())
        self.assertNotIn("binary_data", serialized.lower())

    def test_asset_part_is_preserved_by_text_population(self):
        image = next(
            item for item in self.profile["assets"]["items"] if item["kind"] == "image"
        )
        output = self.root / "generated.docx"
        patch = {
            "contract": self.profile["patch_contract"],
            "elements": {
                "body.p0": {
                    "expected": "Editable summary",
                    "value": "Updated summary",
                    "format_policy": "preserve_first_run",
                }
            },
        }
        render_bundle(self.bundle, output, patch)
        with (
            zipfile.ZipFile(self.source) as original,
            zipfile.ZipFile(output) as generated,
        ):
            self.assertEqual(
                original.read(image["package_location"]),
                generated.read(image["package_location"]),
            )

    def test_asset_registry_schema_is_versioned(self):
        schema = json.loads(
            schema_path("asset-registry-0.1.schema.json").read_text(encoding="utf-8")
        )
        self.assertEqual(schema["properties"]["schema_version"]["const"], "0.1")
        self.assertEqual(schema["properties"]["read_only"]["const"], True)
        self.assertEqual(
            schema["properties"]["binary_content_included"]["const"], False
        )


if __name__ == "__main__":
    unittest.main()
