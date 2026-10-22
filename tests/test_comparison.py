import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from docxunravel import export_bundle, profile_document
from docxunravel.comparison import compare_documents
from docxunravel.core import main


class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.original = self.root / "original.docx"
        document = Document()
        document.add_paragraph("Architecture overview")
        document.add_paragraph("Original summary")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Table content"
        document.add_paragraph("")
        document.sections[0].header.paragraphs[0].text = "Confidential"
        document.save(str(self.original))

    def tearDown(self):
        self.temporary.cleanup()

    def test_bundle_writes_compact_paragraph_content(self):
        bundle = self.root / "bundle"
        profile = export_bundle(self.original, bundle)
        content_path = bundle / "content.json"
        content = json.loads(content_path.read_text(encoding="utf-8"))

        self.assertEqual(content["artifact"], "docxunravel.paragraph-content")
        self.assertEqual(content["patch_contract"], profile["patch_contract"])
        self.assertEqual(content["summary"]["omitted_empty_paragraph_count"], 1)
        self.assertEqual(
            [paragraph["text"] for paragraph in content["paragraphs"]],
            [
                "Architecture overview",
                "Original summary",
                "Table content",
                "Confidential",
            ],
        )
        self.assertEqual(
            set(content),
            {
                "schema_version",
                "artifact",
                "source",
                "summary",
                "paragraphs",
                "patch_contract",
            },
        )
        self.assertLess(
            content_path.stat().st_size, (bundle / "profile.json").stat().st_size
        )

    def test_compare_accepts_docx_profile_content_and_bundle(self):
        revised = self.root / "revised.docx"
        document = Document()
        document.add_paragraph("Architecture overview")
        document.add_paragraph("Revised summary")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Table content"
        document.add_paragraph("New appendix")
        document.sections[0].header.paragraphs[0].text = "Confidential"
        document.save(str(revised))

        report = compare_documents(self.original, revised)
        self.assertFalse(report["equal"])
        self.assertEqual(report["summary"]["modified"], 1)
        self.assertEqual(report["summary"]["added"], 1)
        self.assertEqual(report["summary"]["removed"], 0)
        self.assertEqual(
            [change["kind"] for change in report["changes"]],
            ["modified", "added"],
        )

        original_profile_path = self.root / "original-profile.json"
        original_profile_path.write_text(
            json.dumps(profile_document(self.original)), encoding="utf-8"
        )
        revised_bundle = self.root / "revised-bundle"
        export_bundle(revised, revised_bundle)
        mixed_report = compare_documents(original_profile_path, revised_bundle)
        self.assertEqual(mixed_report["summary"], report["summary"])

        content_report = compare_documents(
            revised_bundle / "content.json", revised_bundle / "content.json"
        )
        self.assertTrue(content_report["equal"])
        self.assertEqual(content_report["changes"], [])

    def test_compare_cli_writes_report_and_can_fail_on_difference(self):
        revised = self.root / "revised.docx"
        document = Document(str(self.original))
        document.paragraphs[1].text = "Changed summary"
        document.save(str(revised))
        output = self.root / "comparison.json"

        stdout = io.StringIO()
        with contextlib.redirect_stdout(stdout):
            status = main(
                [
                    "compare",
                    str(self.original),
                    str(revised),
                    "--output",
                    str(output),
                    "--fail-on-difference",
                ]
            )

        self.assertEqual(status, 2)
        self.assertEqual(json.loads(stdout.getvalue()), json.loads(output.read_text()))
        self.assertEqual(json.loads(output.read_text())["summary"]["modified"], 1)


if __name__ == "__main__":
    unittest.main()
