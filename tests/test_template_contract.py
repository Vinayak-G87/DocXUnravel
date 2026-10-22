import copy
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from docxunravel import export_bundle
from docxunravel.schemas import schema_path
from docxunravel.template_contract import (
    template_contract_binding,
    validate_template_contract,
)


class TemplateContractTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "template.docx"
        document = Document()
        document.add_paragraph("Prepared for {{ CUSTOMER }}", style="Title")
        document.add_paragraph("Original executive summary")
        table = document.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = "Application"
        table.cell(0, 1).text = "Owner"
        table.cell(1, 0).text = "Template application"
        table.cell(1, 1).text = "Template owner"
        document.save(str(self.source))
        self.bundle = self.root / "bundle"
        self.profile = export_bundle(self.source, self.bundle)

    def tearDown(self):
        self.temporary.cleanup()

    def contract(self, fields):
        return {
            "contract_version": "0.1",
            "document_type": "azure_migration_hld",
            "binding": template_contract_binding(self.profile),
            "fields": fields,
        }

    def test_valid_contract_canonicalizes_targets_and_supports_optional_fields(self):
        stable_summary = next(
            element["stable_id"]
            for element in self.profile["body"]
            if element.get("id") == "body.p1"
        )
        contract = self.contract(
            {
                "customer_name": {
                    "type": "placeholder",
                    "target": "body.p0",
                    "token": "CUSTOMER",
                    "required": True,
                },
                "executive_summary": {
                    "type": "text",
                    "target": stable_summary,
                    "required": True,
                    "format_policy": "preserve_first_run",
                },
                "application_owner": {
                    "type": "table_cell",
                    "target": "body.t0.r1.c1",
                    "required": False,
                },
            }
        )
        normalized = validate_template_contract(self.profile, contract)
        self.assertEqual(normalized["fields"]["executive_summary"]["target"], "body.p1")
        self.assertFalse(normalized["fields"]["application_owner"]["required"])
        self.assertEqual(
            contract["fields"]["executive_summary"]["target"], stable_summary
        )

    def test_table_rows_contract_validates_template_row_and_columns(self):
        contract = self.contract(
            {
                "application_inventory": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {"name": 0, "owner": 1},
                    "data_path": "/applications",
                }
            }
        )
        normalized = validate_template_contract(self.profile, contract)
        self.assertEqual(
            normalized["fields"]["application_inventory"]["columns"],
            {"name": 0, "owner": 1},
        )

    def test_incorrect_or_stale_binding_is_rejected(self):
        contract = self.contract(
            {
                "summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                }
            }
        )
        contract["binding"]["source_sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "does not match"):
            validate_template_contract(self.profile, contract)

        contract = self.contract(
            {
                "summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                }
            }
        )
        contract["binding"]["profile_sha256"] = "f" * 64
        with self.assertRaisesRegex(ValueError, "does not match"):
            validate_template_contract(self.profile, contract)

    def test_modified_profile_is_rejected_as_stale(self):
        contract = self.contract(
            {
                "summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                }
            }
        )
        stale_profile = copy.deepcopy(self.profile)
        stale_profile["body"][1]["text"] = "Tampered profile text"
        with self.assertRaisesRegex(ValueError, "stale"):
            validate_template_contract(stale_profile, contract)

    def test_unknown_duplicate_and_conflicting_targets_are_rejected(self):
        unknown = self.contract(
            {
                "summary": {
                    "type": "text",
                    "target": "body.p999",
                    "required": True,
                }
            }
        )
        with self.assertRaisesRegex(KeyError, "unknown target"):
            validate_template_contract(self.profile, unknown)

        duplicate = self.contract(
            {
                "summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                },
                "other_summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": False,
                },
            }
        )
        with self.assertRaisesRegex(ValueError, "same target"):
            validate_template_contract(self.profile, duplicate)

        conflict = self.contract(
            {
                "applications": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {"name": 0, "owner": 1},
                },
                "owner": {
                    "type": "table_cell",
                    "target": "body.t0.r1.c1",
                    "required": False,
                },
            }
        )
        with self.assertRaisesRegex(ValueError, "conflict"):
            validate_template_contract(self.profile, conflict)

    def test_invalid_field_mappings_are_rejected(self):
        wrong_type = self.contract(
            {
                "summary": {
                    "type": "text",
                    "target": "body.t0",
                    "required": True,
                }
            }
        )
        with self.assertRaisesRegex(ValueError, "paragraph target"):
            validate_template_contract(self.profile, wrong_type)

        duplicate_column = self.contract(
            {
                "applications": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {"name": 0, "owner": 0},
                }
            }
        )
        with self.assertRaisesRegex(ValueError, "multiple values"):
            validate_template_contract(self.profile, duplicate_column)

        absent_token = self.contract(
            {
                "brand": {
                    "type": "placeholder",
                    "target": "body.p1",
                    "token": "CUSTOMER",
                    "required": True,
                }
            }
        )
        with self.assertRaisesRegex(ValueError, "not present"):
            validate_template_contract(self.profile, absent_token)

    def test_versioned_schema_is_valid_json_and_separate_from_patch_contract(self):
        schema = json.loads(
            schema_path("template-contract-0.1.schema.json").read_text(encoding="utf-8")
        )
        self.assertEqual(schema["properties"]["contract_version"]["const"], "0.1")
        self.assertNotIn("operations", schema["properties"])


if __name__ == "__main__":
    unittest.main()
