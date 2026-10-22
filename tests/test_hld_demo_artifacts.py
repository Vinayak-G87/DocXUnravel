import json
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from docxunravel import export_bundle, profile_document, validate_docx_package
from docxunravel.template_contract import (
    preview_migration_template_population,
    validate_template_contract,
)


class HldDemoArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.examples = Path(__file__).resolve().parents[1] / "examples" / "hld"
        cls.template = cls.examples / "fictional-azure-migration-hld-template.docx"
        cls.generated = cls.examples / "fictional-azure-migration-hld-output.docx"
        cls.contract = json.loads(
            (cls.examples / "fictional-hld-template-contract-0.1.json").read_text(
                encoding="utf-8"
            )
        )
        cls.discovery = json.loads(
            (cls.examples / "fictional-migration-discovery-0.1.json").read_text(
                encoding="utf-8"
            )
        )

    def test_demo_packages_and_expected_document_features(self):
        self.assertTrue(validate_docx_package(self.template)["valid"])
        self.assertTrue(validate_docx_package(self.generated)["valid"])
        template_profile = profile_document(self.template)
        self.assertTrue(
            any(item["kind"] == "image" for item in template_profile["assets"]["items"])
        )
        generated = Document(str(self.generated))
        text = "\n".join(paragraph.text for paragraph in generated.paragraphs)
        for heading in (
            "Executive Summary",
            "Current-State Assessment",
            "Application Inventory",
            "Target Architecture",
            "Migration Strategy",
            "Security Considerations",
            "Risks and Assumptions",
            "Appendix",
        ):
            self.assertIn(heading, text)
        self.assertIn("Fictional Northstar Research Group", text)
        self.assertGreaterEqual(len(generated.tables), 12)
        with (
            zipfile.ZipFile(self.template) as original,
            zipfile.ZipFile(self.generated) as populated,
        ):
            self.assertIn(b"PAGE", original.read("word/footer1.xml"))
            self.assertIn(b"DocXUnravel v0.4", original.read("word/header1.xml"))
            self.assertEqual(
                original.read("word/media/image1.png"),
                populated.read("word/media/image1.png"),
            )

    def test_contract_rebinds_and_preview_is_reproducible(self):
        with tempfile.TemporaryDirectory() as temporary:
            bundle = Path(temporary) / "bundle"
            profile = export_bundle(self.template, bundle)
            normalized = validate_template_contract(profile, self.contract)
            self.assertEqual(normalized["document_type"], "azure_migration_hld")
            preview = preview_migration_template_population(
                bundle, self.contract, self.discovery
            )
            self.assertTrue(preview["valid"])
            self.assertEqual(preview["unsupported_operations"], [])
            persisted = json.loads(
                (self.examples / "fictional-hld-dry-run-preview.json").read_text(
                    encoding="utf-8"
                )
            )
            self.assertEqual(preview["patch"], persisted["patch"])

    def test_persisted_integrity_fidelity_and_visual_reports_pass(self):
        integrity = json.loads(
            (self.examples / "fictional-hld-package-integrity-report.json").read_text(
                encoding="utf-8"
            )
        )
        fidelity = json.loads(
            (self.examples / "fictional-hld-fidelity-report.json").read_text(
                encoding="utf-8"
            )
        )
        visual = json.loads(
            (self.examples / "fictional-hld-visual-report.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertTrue(integrity["passed"])
        self.assertEqual(
            integrity["part_hash_preservation"]["changed_parts"],
            ["word/document.xml"],
        )
        self.assertTrue(
            all(asset["preserved"] for asset in integrity["read_only_assets"])
        )
        self.assertTrue(fidelity["passed"])
        self.assertTrue(
            fidelity["semantic_text_preservation"]["matches_expected_patch"]
        )
        self.assertTrue(visual["passed"])
        self.assertTrue(visual["page_count_match"])
        self.assertEqual(visual["unexpected_changed_pages"], [])


if __name__ == "__main__":
    unittest.main()
