import asyncio
import base64
import json
import os
import tempfile
import time
import unittest
from pathlib import Path

import jwt
from cryptography.hazmat.primitives.asymmetric import rsa
from docx import Document
from mcp.server.auth.middleware.auth_context import auth_context_var
from mcp.server.auth.middleware.bearer_auth import AuthenticatedUser
from mcp.server.auth.provider import AccessToken
from starlette.testclient import TestClient

from docxunravel.mcp_auth import OIDCJWTVerifier
from docxunravel.mcp_server import (
    PREVIEW_SCOPE,
    READ_SCOPE,
    WRITE_SCOPE,
    DocumentWorkspace,
    HTTPAuthConfig,
    _require_scopes,
    create_mcp_server,
    main,
)
from docxunravel.template_contract import template_contract_binding


def _base64url_integer(value: int) -> str:
    size = (value.bit_length() + 7) // 8
    return base64.urlsafe_b64encode(value.to_bytes(size, "big")).rstrip(b"=").decode()


class JWTVerifierTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.private_key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
        numbers = cls.private_key.public_key().public_numbers()
        cls.jwks = {
            "keys": [
                {
                    "kty": "RSA",
                    "kid": "test-key",
                    "use": "sig",
                    "alg": "RS256",
                    "n": _base64url_integer(numbers.n),
                    "e": _base64url_integer(numbers.e),
                }
            ]
        }
        cls.issuer = "https://issuer.example.test"
        cls.audience = "api://docxunravel"

    def token(self, **overrides):
        now = int(time.time())
        claims = {
            "iss": self.issuer,
            "aud": self.audience,
            "sub": "user-123",
            "azp": "mcp-client",
            "iat": now,
            "nbf": now - 1,
            "exp": now + 300,
            "scope": f"{READ_SCOPE} {PREVIEW_SCOPE}",
        }
        claims.update(overrides)
        return jwt.encode(
            claims,
            self.private_key,
            algorithm="RS256",
            headers={"kid": "test-key", "typ": "at+jwt"},
        )

    def verifier(self):
        return OIDCJWTVerifier(
            issuer=self.issuer,
            audience=self.audience,
            algorithms=("RS256",),
            allowed_client_ids=("mcp-client",),
            jwks=self.jwks,
        )

    def verify(self, token):
        return asyncio.run(self.verifier().verify_token(token))

    def test_valid_jwt_returns_minimal_access_metadata(self):
        access = self.verify(self.token())
        self.assertIsNotNone(access)
        assert access is not None
        self.assertEqual(access.client_id, "mcp-client")
        self.assertEqual(access.subject, "user-123")
        self.assertEqual(set(access.scopes), {READ_SCOPE, PREVIEW_SCOPE})
        self.assertIsNotNone(access.claims)
        assert access.claims is not None
        self.assertEqual(set(access.claims), {"iss", "aud", "azp"})

    def test_expired_wrong_issuer_and_wrong_audience_are_rejected(self):
        now = int(time.time())
        self.assertIsNone(self.verify(self.token(exp=now - 120)))
        self.assertIsNone(self.verify(self.token(nbf=now + 120)))
        self.assertIsNone(self.verify(self.token(iss="https://other.example.test")))
        self.assertIsNone(self.verify(self.token(aud="api://other")))

    def test_unapproved_client_and_algorithm_are_rejected(self):
        self.assertIsNone(self.verify(self.token(azp="unapproved-client")))
        unsigned = jwt.encode(
            {
                "iss": self.issuer,
                "aud": self.audience,
                "azp": "mcp-client",
                "exp": int(time.time()) + 300,
            },
            key="",
            algorithm="none",
            headers={"kid": "test-key"},
        )
        self.assertIsNone(self.verify(unsigned))


class StaticVerifier:
    async def verify_token(self, token: str):
        if token != "header-token":
            return None
        return AccessToken(
            token=token,
            client_id="test-client",
            scopes=[READ_SCOPE],
            resource="api://docxunravel",
        )


class MCPHTTPTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.workspace = DocumentWorkspace(self.temporary.name)
        self.auth = HTTPAuthConfig(
            issuer="https://issuer.example.test",
            audience="api://docxunravel",
            resource_server_url="https://documents.example.test/mcp",
        )
        self.server = create_mcp_server(
            self.workspace, auth=self.auth, token_verifier=StaticVerifier()
        )

    def tearDown(self):
        self.temporary.cleanup()

    def test_rfc9728_metadata_and_header_only_bearer_auth(self):
        app = self.server.streamable_http_app(stateless_http=True, json_response=True)
        with TestClient(app, base_url="http://127.0.0.1:8000") as client:
            metadata = client.get("/.well-known/oauth-protected-resource/mcp")
            self.assertEqual(metadata.status_code, 200)
            value = metadata.json()
            self.assertEqual(value["resource"], "https://documents.example.test/mcp")
            self.assertEqual(
                value["authorization_servers"], ["https://issuer.example.test/"]
            )
            self.assertEqual(value["bearer_methods_supported"], ["header"])

            query_only = client.get("/mcp?access_token=header-token")
            self.assertEqual(query_only.status_code, 401)
            self.assertIn("resource_metadata", query_only.headers["www-authenticate"])

            header = client.post(
                "/mcp",
                headers={"Authorization": "Bearer header-token"},
                json={},
            )
            self.assertNotEqual(header.status_code, 401)

    def test_operation_scopes_are_cumulative(self):
        access = AccessToken(
            token="not-logged",
            client_id="client",
            scopes=[READ_SCOPE, PREVIEW_SCOPE],
        )
        context_token = auth_context_var.set(AuthenticatedUser(access))
        try:
            _require_scopes({READ_SCOPE}, True)
            _require_scopes({READ_SCOPE, PREVIEW_SCOPE}, True)
            with self.assertRaisesRegex(PermissionError, WRITE_SCOPE):
                _require_scopes({READ_SCOPE, PREVIEW_SCOPE, WRITE_SCOPE}, True)
        finally:
            auth_context_var.reset(context_token)

    def test_stdio_configuration_needs_no_oauth_but_http_refuses_it(self):
        server = create_mcp_server(self.workspace)
        tools = asyncio.run(server.list_tools())
        self.assertEqual(
            {tool.name for tool in tools},
            {
                "list_documents",
                "profile_document",
                "preview_template_population",
                "populate_template",
                "validate_output",
            },
        )
        with self.assertRaisesRegex(SystemExit, "unauthenticated HTTP is disabled"):
            main(
                [
                    "--transport",
                    "streamable-http",
                    "--workspace",
                    self.temporary.name,
                ]
            )


class MCPWorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.workspace = DocumentWorkspace(self.root / "workspace")
        self.source = self.workspace.documents / "template_01.docx"
        document = Document()
        document.add_paragraph("Project for {{ CUSTOMER }}", style="Title")
        document.add_paragraph("Original scope")
        table = document.add_table(rows=2, cols=4)
        table.style = "Table Grid"
        for index, value in enumerate(("ID", "Application", "Owner", "Criticality")):
            table.cell(0, index).text = value
        for index, value in enumerate(("APP-ID", "Template", "Owner", "medium")):
            table.cell(1, index).text = value
        document.save(str(self.source))

        self.data = {
            "customer": {"name": "Northstar Research Group"},
            "project": {"scope": "A project summary."},
            "application_inventory": [
                {
                    "id": "APP-1",
                    "name": "Portal",
                    "owner": "Team A",
                    "business_criticality": "high",
                }
            ],
        }
        self.profile = self.workspace.profile("template_01")
        self.contract = {
            "contract_version": "0.1",
            "document_type": "project_form",
            "binding": template_contract_binding(self.profile),
            "fields": {
                "customer_brand": {
                    "type": "placeholder",
                    "target": "body.p0",
                    "token": "CUSTOMER",
                    "required": True,
                    "data_path": "/customer/name",
                },
                "executive_summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                    "data_path": "/project/scope",
                },
                "application_inventory": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {
                        "id": 0,
                        "name": 1,
                        "owner": 2,
                        "business_criticality": 3,
                    },
                    "data_path": "/application_inventory",
                },
            },
        }

    def tearDown(self):
        self.temporary.cleanup()

    def test_workspace_rejects_traversal_and_symlink_documents(self):
        for invalid in ("../outside", "a/b", ".", "", "x.docx"):
            with self.assertRaises(ValueError):
                self.workspace.profile(invalid)
        outside = self.root / "outside.docx"
        outside.write_bytes(self.source.read_bytes())
        link = self.workspace.documents / "linked.docx"
        os.symlink(outside, link)
        with self.assertRaisesRegex(ValueError, "Symbolic-link"):
            self.workspace.profile("linked")

    def test_profile_is_semantic_and_contains_no_binary_base64(self):
        serialized = json.dumps(self.profile)
        self.assertIn("assets", self.profile)
        self.assertNotIn("base64", serialized.lower())
        self.assertNotIn("data:application/", serialized.lower())
        listed = self.workspace.list_documents()
        self.assertEqual(listed["count"], 1)
        self.assertEqual(listed["documents"][0]["document_id"], "template_01")

    def test_preview_is_read_only_and_population_is_atomic_and_validated(self):
        before = list(self.workspace.outputs.iterdir())
        preview = self.workspace.preview("template_01", self.contract, self.data)
        self.assertTrue(preview["valid"])
        self.assertEqual(list(self.workspace.outputs.iterdir()), before)

        result = self.workspace.populate("template_01", self.contract, self.data)
        self.assertRegex(result["output_id"], r"^[0-9a-f]{32}$")
        self.assertTrue(result["fidelity_passed"])
        self.assertTrue(result["render"]["integrity"]["valid"])
        output = self.workspace.outputs / f"{result['output_id']}.docx"
        report = self.workspace.reports / f"{result['output_id']}.json"
        self.assertTrue(output.is_file())
        self.assertTrue(report.is_file())
        validation = self.workspace.validate(result["output_id"])
        self.assertTrue(validation["package"]["valid"])
        self.assertTrue(validation["fidelity"]["passed"])

        count = len(list(self.workspace.outputs.glob("*.docx")))
        stale_contract = json.loads(json.dumps(self.contract))
        stale_contract["binding"]["source_sha256"] = "0" * 64
        with self.assertRaises(ValueError):
            self.workspace.populate("template_01", stale_contract, self.data)
        self.assertEqual(len(list(self.workspace.outputs.glob("*.docx"))), count)
        self.assertFalse(list(self.workspace.outputs.glob("*.tmp")))
        self.assertFalse(list(self.workspace.outputs.glob("render-*")))


if __name__ == "__main__":
    unittest.main()
