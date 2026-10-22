import json
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from docxunravel import (
    PackageLimits,
    create_fidelity_report,
    export_bundle,
    inspect_docx_archive,
    preview_bundle,
    profile_document,
    render_bundle,
    validate_docx_package,
)
from tests.support.fixture_factory import create_complex_fixture


class PackageSecurityTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "source.docx"
        document = Document()
        document.add_paragraph("Safe package")
        document.save(self.source)

    def tearDown(self):
        self.temporary.cleanup()

    def _rewrite_part(self, destination: Path, part: str, content: bytes) -> None:
        with zipfile.ZipFile(self.source) as archive:
            entries = [(info, archive.read(info)) for info in archive.infolist()]
        with zipfile.ZipFile(destination, "w") as archive:
            for info, data in entries:
                archive.writestr(info, content if info.filename == part else data)

    def test_package_limits_are_enforced(self):
        with self.assertRaisesRegex(ValueError, "entries"):
            inspect_docx_archive(self.source, PackageLimits(max_entries=1))
        with self.assertRaisesRegex(ValueError, "entry exceeds size limit"):
            inspect_docx_archive(
                self.source,
                PackageLimits(
                    max_entry_size=128, max_total_uncompressed_size=10_000_000
                ),
            )
        aggregate = self.root / "aggregate.docx"
        with zipfile.ZipFile(aggregate, "w", compression=zipfile.ZIP_STORED) as archive:
            archive.writestr("one.bin", b"a" * 400)
            archive.writestr("two.bin", b"b" * 400)
            archive.writestr("three.bin", b"c" * 400)
        with self.assertRaisesRegex(ValueError, "total uncompressed size"):
            inspect_docx_archive(
                aggregate,
                PackageLimits(max_entry_size=500, max_total_uncompressed_size=1_000),
                validate_xml=False,
            )
        with self.assertRaisesRegex(ValueError, "compression ratio"):
            inspect_docx_archive(
                self.source,
                PackageLimits(max_compression_ratio=1.0),
            )

    def test_traversal_duplicate_and_invalid_zip_are_rejected(self):
        traversal = self.root / "traversal.docx"
        with zipfile.ZipFile(traversal, "w") as archive:
            archive.writestr("../escape.xml", "<x/>")
        with self.assertRaisesRegex(ValueError, "Unsafe DOCX package path"):
            inspect_docx_archive(traversal)

        duplicate = self.root / "duplicate.docx"
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(duplicate, "w") as archive:
                archive.writestr("word/document.xml", "<x/>")
                archive.writestr("word/document.xml", "<x/>")
        with self.assertRaisesRegex(ValueError, "duplicate"):
            inspect_docx_archive(duplicate)

        invalid = self.root / "invalid.docx"
        invalid.write_bytes(b"not a zip file")
        with self.assertRaisesRegex(ValueError, "Invalid DOCX ZIP"):
            inspect_docx_archive(invalid)

    def test_malformed_xml_and_entities_are_rejected(self):
        malformed = self.root / "malformed.docx"
        self._rewrite_part(malformed, "word/document.xml", b"<w:document")
        with self.assertRaisesRegex(ValueError, "Invalid XML part"):
            inspect_docx_archive(malformed)

        entity = self.root / "entity.docx"
        payload = (
            b'<?xml version="1.0"?>'
            b'<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
            b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            b"<w:body><w:p><w:r><w:t>&xxe;</w:t></w:r></w:p></w:body></w:document>"
        )
        self._rewrite_part(entity, "word/document.xml", payload)
        with self.assertRaisesRegex(ValueError, "DTD or entity declarations"):
            inspect_docx_archive(entity)
        with self.assertRaisesRegex(ValueError, "DTD or entity declarations"):
            profile_document(entity)

    def test_missing_parts_and_broken_relationships_are_rejected(self):
        missing = self.root / "missing.docx"
        with zipfile.ZipFile(missing, "w") as archive:
            archive.writestr("data.xml", "<data/>")
        with self.assertRaisesRegex(ValueError, "missing required parts"):
            validate_docx_package(missing)

        broken = self.root / "broken-relationship.docx"
        with zipfile.ZipFile(self.source) as archive:
            relationships = etree.fromstring(
                archive.read("word/_rels/document.xml.rels")
            )
        relationship = relationships[0]
        relationship.set("Target", "missing-part.xml")
        self._rewrite_part(
            broken,
            "word/_rels/document.xml.rels",
            etree.tostring(relationships, xml_declaration=True, encoding="UTF-8"),
        )
        with self.assertRaisesRegex(ValueError, "Broken relationship"):
            validate_docx_package(broken)


class FidelityFixtureTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)

    def tearDown(self):
        self.temporary.cleanup()

    def test_complex_fixture_contains_and_preserves_advanced_constructs(self):
        fixture = create_complex_fixture(self.root / "complex.docx")
        inspection = inspect_docx_archive(fixture)
        self.assertTrue(inspection["valid"])
        with zipfile.ZipFile(fixture) as archive:
            names = set(archive.namelist())
            document_xml = archive.read("word/document.xml")
            self.assertIn("word/comments.xml", names)
            self.assertIn(b"bookmarkStart", document_xml)
            self.assertIn(b"<w:ins", document_xml)
            self.assertIn(b"<w:del", document_xml)
        document = Document(fixture)
        self.assertEqual(len(document.sections), 3)
        self.assertGreaterEqual(len(document.tables), 2)
        self.assertTrue(document.comments)

        profile = profile_document(fixture)
        revision_paragraphs = [
            paragraph
            for paragraph in profile["body"]
            if paragraph.get("type") == "paragraph" and paragraph.get("has_revisions")
        ]
        self.assertEqual(len(revision_paragraphs), 2)
        bundle = self.root / "bundle"
        bundle_profile = export_bundle(fixture, bundle)
        with self.assertRaisesRegex(ValueError, "Tracked revisions"):
            preview_bundle(
                bundle,
                {
                    "contract": bundle_profile["patch_contract"],
                    "elements": {revision_paragraphs[0]["id"]: "unsafe"},
                },
            )

        output = self.root / "noop.docx"
        render_bundle(bundle, output, {"contract": bundle_profile["patch_contract"]})
        with zipfile.ZipFile(fixture) as original, zipfile.ZipFile(output) as generated:
            for name in original.namelist():
                self.assertEqual(original.read(name), generated.read(name), name)

    def test_machine_readable_noop_and_patch_fidelity_reports(self):
        source = self.root / "report.docx"
        document = Document()
        document.add_paragraph("Title", style="Title")
        document.add_paragraph("Original summary")
        document.add_paragraph("Numbered item", style="List Number")
        document.save(source)
        bundle = self.root / "report-bundle"
        profile = export_bundle(source, bundle)

        noop = self.root / "noop.docx"
        noop_patch = {"contract": profile["patch_contract"]}
        render_bundle(bundle, noop, noop_patch)
        noop_report = create_fidelity_report(
            source, noop, bundle=bundle, patch=noop_patch
        )
        self.assertTrue(noop_report["passed"])
        self.assertTrue(noop_report["semantic_text_preservation"]["exact_match"])
        self.assertTrue(noop_report["style_and_numbering_preservation"]["passed"])
        self.assertFalse(noop_report["visual_differences"]["performed"])

        patch = {
            "contract": profile["patch_contract"],
            "elements": {
                "body.p1": {
                    "expected": "Original summary",
                    "value": "Updated summary",
                    "format_policy": "preserve_first_run",
                }
            },
        }
        generated = self.root / "patched.docx"
        render_bundle(bundle, generated, patch)
        report = create_fidelity_report(source, generated, bundle=bundle, patch=patch)
        self.assertTrue(report["passed"])
        self.assertFalse(report["semantic_text_preservation"]["exact_match"])
        self.assertTrue(report["semantic_text_preservation"]["matches_expected_patch"])
        self.assertEqual(
            report["untouched_part_hash_preservation"]["unexpected_changed_parts"], []
        )
        self.assertIn(
            "word/document.xml",
            report["untouched_part_hash_preservation"]["expected_modified_parts"],
        )


if __name__ == "__main__":
    unittest.main()
