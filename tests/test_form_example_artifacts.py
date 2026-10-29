import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

from docxunravel import (
    create_fidelity_report,
    export_bundle,
    inspect_controls,
    validate_docx_package,
)
from scripts.package_examples import EXAMPLE_FILES, build_example_archive


class FormExampleArtifactTests(unittest.TestCase):
    def test_release_archive_runs_without_repository_examples(self):
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            first = directory / "examples.zip"
            second = directory / "again.zip"
            build_example_archive(first)
            build_example_archive(second)
            self.assertEqual(first.read_bytes(), second.read_bytes())
            with zipfile.ZipFile(first) as archive:
                self.assertEqual(archive.namelist(), list(EXAMPLE_FILES))
                self.assertIsNone(archive.testzip())
                archive.extractall(directory / "downloaded")

            example = directory / "downloaded" / "examples" / "form"
            patch = json.loads((example / "patch.json").read_text(encoding="utf-8"))
            downloaded = directory / "downloaded"
            commands = (
                (
                    "extract",
                    "examples/form/form-template.docx",
                    "form-bundle",
                    "--explicit-placeholders-only",
                ),
                ("preview", "form-bundle", "examples/form/patch.json"),
                (
                    "render-bundle",
                    "form-bundle",
                    "examples/form/patch.json",
                    "filled.docx",
                ),
            )
            for command in commands:
                subprocess.run(
                    [sys.executable, "-m", "docxunravel", *command],
                    cwd=downloaded,
                    capture_output=True,
                    text=True,
                    check=True,
                )
            profile = json.loads(
                (downloaded / "form-bundle" / "profile.json").read_text(
                    encoding="utf-8"
                )
            )
            self.assertEqual(patch["contract"], profile["patch_contract"])
            rendered = downloaded / "filled.docx"
            self.assertTrue(validate_docx_package(rendered)["valid"])
            source = Document(str(example / "form-template.docx"))
            result = Document(str(rendered))
            self.assertEqual(len(source.sections), len(result.sections))
            for before, after in zip(source.sections, result.sections):
                for property_name in (
                    "orientation",
                    "page_width",
                    "page_height",
                    "top_margin",
                    "bottom_margin",
                    "left_margin",
                    "right_margin",
                    "header_distance",
                    "footer_distance",
                    "gutter",
                ):
                    with self.subTest(property=property_name):
                        self.assertEqual(
                            getattr(before, property_name),
                            getattr(after, property_name),
                        )

    def test_form_artifacts(self):
        example = Path(__file__).resolve().parents[1] / "examples" / "form"
        template = example / "form-template.docx"
        filled = example / "form-filled.docx"
        field_map = json.loads((example / "field-map.json").read_text(encoding="utf-8"))
        patch = json.loads((example / "patch.json").read_text(encoding="utf-8"))

        for path in (template, filled):
            self.assertTrue(validate_docx_package(path)["valid"])
        source = Document(str(template))
        result = Document(str(filled))
        self.assertEqual(
            [source.tables[0].cell(row, 1).text for row in (0, 1)], ["[...]"] * 2
        )
        self.assertNotEqual(field_map["fields"]["name"], field_map["fields"]["summary"])
        for row, key in enumerate(("name", "summary")):
            target = field_map["fields"][key]
            self.assertEqual(patch["elements"][target]["expected"], "[...]")
            self.assertEqual(
                result.tables[0].cell(row, 1).text, patch["elements"][target]["value"]
            )
        for row, label in enumerate(("Item A", "Item B"), 1):
            self.assertEqual(result.tables[1].cell(row, 0).text, label)
            for column, key in ((1, "goal"), (2, "metric")):
                target = field_map["goals"][label][key]
                self.assertEqual(
                    result.tables[1].cell(row, column).text,
                    patch["elements"][target]["value"],
                )

        before = {item["id"]: item for item in inspect_controls(template)["checkboxes"]}
        after = {item["id"]: item for item in inspect_controls(filled)["checkboxes"]}
        self.assertEqual(set(before), set(field_map["checkboxes"].values()))
        for label, target in field_map["checkboxes"].items():
            self.assertFalse(before[target]["checked"])
            self.assertEqual(after[target]["checked"], label == "Standard")
        self.assertEqual(
            patch["operations"][0]["target"], field_map["checkboxes"]["Standard"]
        )

        with tempfile.TemporaryDirectory() as temporary:
            bundle = Path(temporary) / "bundle"
            profile = export_bundle(template, bundle, auto_placeholders=False)
            self.assertEqual(field_map["source_sha256"], profile["source"]["sha256"])
            self.assertEqual(patch["contract"], profile["patch_contract"])
            report = create_fidelity_report(
                template, filled, bundle=bundle, patch=patch
            )
            self.assertTrue(report["passed"], report)


if __name__ == "__main__":
    unittest.main()
