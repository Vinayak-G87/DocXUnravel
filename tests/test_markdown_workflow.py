import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from docxunravel import export_bundle
from docxunravel.markdown_ingestion import (
    compare_markdown,
    parse_markdown,
    render_markdown_template,
)
from docxunravel.markdown_workflow import (
    compare_markdown_to_docx,
    populate_markdown_docx,
    preview_markdown_population,
    render_markdown_docx,
)


class MarkdownWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "template.docx"

        document = Document()
        document.add_paragraph("Original summary")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "ID"
        table.cell(0, 1).text = "Application"
        table.cell(1, 0).text = "template-id"
        table.cell(1, 1).text = "Keep this cell prefix"
        table.cell(1, 1).add_paragraph("template-name")
        document.add_paragraph("☐ Azure")
        document.add_paragraph("☐ AWS")
        document.save(str(self.source))

        self.bundle = self.root / "bundle"
        profile = export_bundle(self.source, self.bundle)
        checkbox_ids = [item["id"] for item in profile["checkboxes"]["items"]]
        self.contract = {
            "contract_version": "0.1",
            "document_type": "markdown_test",
            "binding": profile["patch_contract"],
            "fields": {
                "summary": {
                    "type": "text",
                    "target": "body.p0",
                    "required": True,
                },
                "applications": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "empty_policy": "remove_template_row",
                    "row_targets": {
                        "id": {"target": "c0.p0", "data_path": "/id"},
                        "name": {"target": "c1.p1", "data_path": "/name"},
                    },
                },
                "hosting": {
                    "type": "checkbox_group",
                    "required": True,
                    "checkboxes": {
                        "azure": checkbox_ids[0],
                        "aws": checkbox_ids[1],
                    },
                },
            },
        }
        self.contract_path = self.root / "contract.json"
        self.contract_path.write_text(json.dumps(self.contract), encoding="utf-8")
        self.markdown = self.root / "generated.md"
        self.markdown.write_text(
            """<!-- docxunravel:field summary -->
Cloud migration summary.

<!-- docxunravel:table applications -->
| id | name |
|---|---|
| APP-001 | Portal |
| APP-002 | Billing |

<!-- docxunravel:field azure -->
- [x] Azure

<!-- docxunravel:field aws -->
- [ ] AWS
""",
            encoding="utf-8",
        )

    def tearDown(self):
        self.temporary.cleanup()

    def test_markdown_populates_text_rows_and_checkboxes(self):
        preview = preview_markdown_population(self.bundle, self.contract, self.markdown)
        self.assertTrue(preview["valid"])
        self.assertEqual(preview["markdown"]["unmapped"], {})

        output = self.root / "output.docx"
        patch = self.root / "patch.json"
        result = populate_markdown_docx(
            self.bundle,
            output,
            self.contract,
            self.markdown,
            patch_output=patch,
        )
        self.assertEqual(result["render"]["structural_changes"], 1)

        rendered = Document(str(output))
        self.assertEqual(rendered.paragraphs[0].text, "Cloud migration summary.")
        self.assertEqual(len(rendered.tables[0].rows), 3)
        self.assertEqual(rendered.tables[0].cell(1, 0).text, "APP-001")
        self.assertEqual(
            rendered.tables[0].cell(1, 1).paragraphs[0].text, "Keep this cell prefix"
        )
        self.assertEqual(rendered.tables[0].cell(1, 1).paragraphs[1].text, "Portal")
        self.assertEqual(rendered.tables[0].cell(2, 1).paragraphs[1].text, "Billing")
        self.assertEqual(rendered.paragraphs[1].text, "☑ Azure")
        self.assertEqual(rendered.paragraphs[2].text, "☐ AWS")

        comparison = compare_markdown_to_docx(
            self.markdown, output, self.bundle, self.contract
        )
        self.assertTrue(comparison["equal"])
        self.assertEqual(comparison["summary"]["different"], 0)

    def test_markdown_updates_predefined_keyed_rows(self):
        source = self.root / "keyed-template.docx"
        document = Document()
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Type"
        table.cell(0, 1).text = "Output Goal"
        table.cell(1, 0).text = "Agility"
        table.cell(2, 0).text = "Security"
        document.save(str(source))
        bundle = self.root / "keyed-bundle"
        profile = export_bundle(source, bundle)
        contract = {
            "contract_version": "0.1",
            "document_type": "keyed_rows",
            "binding": profile["patch_contract"],
            "fields": {
                "output_goals": {
                    "type": "table_keyed_rows",
                    "target": "body.t0",
                    "required": True,
                    "key": {
                        "column": 0,
                        "data_path": "/type",
                        "match": "case_insensitive",
                    },
                    "row_targets": {
                        "output": {"target": "c1.p0", "data_path": "/output"}
                    },
                }
            },
        }
        markdown = self.root / "keyed.md"
        markdown.write_text(
            """<!-- docxunravel:table output_goals -->
| type | output |
|---|---|
| agility | Scale automatically |
| security | Encrypt customer data |
""",
            encoding="utf-8",
        )
        output = self.root / "keyed-output.docx"
        populate_markdown_docx(bundle, output, contract, markdown)
        rendered = Document(str(output))
        self.assertEqual(len(rendered.tables[0].rows), 3)
        self.assertEqual(rendered.tables[0].cell(1, 1).text, "Scale automatically")
        self.assertEqual(rendered.tables[0].cell(2, 1).text, "Encrypt customer data")

    def test_markdown_template_repeats_rows_and_renders_standalone_docx(self):
        template = self.root / "template.md"
        template.write_text(
            """# {{project.name}}

| ID | Name |
|---|---|
| {{applications.id}} | {{applications.name}} |
""",
            encoding="utf-8",
        )
        data = {
            "project": {"name": "Migration Plan"},
            "applications": [
                {"id": "A-1", "name": "Portal"},
                {"id": "A-2", "name": "Billing"},
            ],
        }
        rendered_markdown = render_markdown_template(template, data)
        self.assertIn("| A-1 | Portal |", rendered_markdown)
        self.assertIn("| A-2 | Billing |", rendered_markdown)

        output = self.root / "standalone.docx"
        result = render_markdown_docx(template, data, output)
        self.assertEqual(result["output"], str(output))
        document = Document(str(output))
        self.assertEqual(document.paragraphs[0].text, "Migration Plan")
        self.assertEqual(len(document.tables[0].rows), 3)

    def test_markdown_parser_and_comparison_are_structural(self):
        model = parse_markdown(self.markdown)
        self.assertEqual(model["tables"]["applications"]["columns"], ["id", "name"])
        revised = self.root / "revised.md"
        revised.write_text(
            self.markdown.read_text(encoding="utf-8").replace("Billing", "Payments"),
            encoding="utf-8",
        )
        report = compare_markdown(self.markdown, revised)
        self.assertFalse(report["equal"])
        self.assertEqual(report["summary"]["modified"], 1)


if __name__ == "__main__":
    unittest.main()
