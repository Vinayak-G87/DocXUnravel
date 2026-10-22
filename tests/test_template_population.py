import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from docxunravel import export_bundle, render_bundle
from docxunravel.template_contract import (
    compile_template_patch,
    preview_template_population,
    template_contract_binding,
)


def _xml(node):
    return node.xml if node is not None else None


class TemplatePopulationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.source = self.root / "hld-template.docx"
        document = Document()
        document.add_paragraph("HLD for {{ CUSTOMER }}", style="Title")
        document.add_paragraph("Original executive summary")
        document.add_paragraph("Optional current-state narrative")

        control = document.add_table(rows=2, cols=2)
        control.style = "Table Grid"
        control.cell(0, 0).text = "Field"
        control.cell(0, 1).text = "Value"
        control.cell(1, 0).text = "Owner"
        control.cell(1, 1).text = "Template owner"

        applications = document.add_table(rows=2, cols=3)
        applications.style = "Light Shading Accent 1"
        for index, value in enumerate(("Application", "Owner", "Disposition")):
            applications.cell(0, index).text = value
        for index, value in enumerate(("Template app", "Template owner", "Assess")):
            cell = applications.cell(1, index)
            cell.text = value
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "D9EAF7")
            cell._tc.get_or_add_tcPr().append(shading)
        document.save(str(self.source))
        self.bundle = self.root / "bundle"
        self.profile = export_bundle(self.source, self.bundle)
        self.contract = {
            "contract_version": "0.1",
            "document_type": "azure_migration_hld",
            "binding": template_contract_binding(self.profile),
            "fields": {
                "customer_name": {
                    "type": "placeholder",
                    "target": "body.p0",
                    "token": "CUSTOMER",
                    "required": True,
                    "data_path": "/customer/name",
                },
                "executive_summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                    "data_path": "/narrative/executive_summary",
                },
                "current_state": {
                    "type": "text",
                    "target": "body.p2",
                    "required": False,
                    "data_path": "/narrative/current_state",
                },
                "document_owner": {
                    "type": "table_cell",
                    "target": "body.t0.r1.c1",
                    "required": True,
                    "data_path": "/project/owner",
                },
                "application_inventory": {
                    "type": "table_rows",
                    "target": "body.t1",
                    "required": True,
                    "template_row": 1,
                    "columns": {"name": 0, "owner": 1, "disposition": 2},
                    "data_path": "/applications",
                },
            },
        }
        self.data = {
            "customer": {"name": "Fictional Contoso Research"},
            "project": {"owner": "Fictional Avery Morgan"},
            "narrative": {
                "executive_summary": "Fictional migration scope and target outcomes."
            },
            "applications": [
                {"name": "Atlas", "owner": "Operations", "disposition": "Rehost"},
                {"name": "Beacon", "owner": "Finance", "disposition": "Refactor"},
                {"name": "Cedar", "owner": "HR", "disposition": "Retire"},
            ],
        }

    def tearDown(self):
        self.temporary.cleanup()

    def test_compiles_typed_fields_and_dry_run_without_writing_output(self):
        patch = compile_template_patch(self.profile, self.contract, self.data)
        self.assertEqual(patch["tokens"], {"CUSTOMER": "Fictional Contoso Research"})
        self.assertEqual(
            patch["elements"]["body.p1"]["expected"], "Original executive summary"
        )
        self.assertIn("body.t0.r1.c1.p0", patch["elements"])
        self.assertEqual(patch["operations"][0]["op"], "repeat_table_rows")
        self.assertEqual(len(patch["operations"][0]["rows"]), 3)
        self.assertNotIn("body.p2", patch["elements"])

        preview = preview_template_population(self.bundle, self.contract, self.data)
        self.assertTrue(preview["valid"])
        self.assertEqual(preview["unsupported_operations"], [])
        self.assertEqual(
            {change["semantic_field"] for change in preview["changes"]},
            {
                "customer_name",
                "executive_summary",
                "document_owner",
                "application_inventory",
            },
        )
        self.assertFalse((self.root / "generated.docx").exists())

    def test_multiple_rows_preserve_template_row_formatting(self):
        source_document = Document(str(self.source))
        source_table = source_document.tables[1]
        source_row = source_table.rows[1]
        expected_tr_pr = _xml(source_row._tr.trPr)
        expected_tc_pr = [_xml(cell._tc.tcPr) for cell in source_row.cells]
        expected_p_pr = [_xml(cell.paragraphs[0]._p.pPr) for cell in source_row.cells]
        expected_r_pr = [
            _xml(cell.paragraphs[0].runs[0]._r.rPr) for cell in source_row.cells
        ]

        patch = compile_template_patch(self.profile, self.contract, self.data)
        output = self.root / "generated.docx"
        render_bundle(self.bundle, output, patch)
        generated = Document(str(output))
        table = generated.tables[1]
        self.assertEqual(getattr(table.style, "name", None), "Light Shading Accent 1")
        self.assertEqual(len(table.rows), 4)
        self.assertEqual(
            [[cell.text for cell in row.cells] for row in table.rows[1:]],
            [
                ["Atlas", "Operations", "Rehost"],
                ["Beacon", "Finance", "Refactor"],
                ["Cedar", "HR", "Retire"],
            ],
        )
        for row in table.rows[1:]:
            self.assertEqual(_xml(row._tr.trPr), expected_tr_pr)
            self.assertEqual(
                [_xml(cell._tc.tcPr) for cell in row.cells], expected_tc_pr
            )
            self.assertEqual(
                [_xml(cell.paragraphs[0]._p.pPr) for cell in row.cells], expected_p_pr
            )
            self.assertEqual(
                [_xml(cell.paragraphs[0].runs[0]._r.rPr) for cell in row.cells],
                expected_r_pr,
            )

    def test_missing_required_field_and_invalid_row_are_rejected(self):
        missing = dict(self.data)
        missing.pop("customer")
        with self.assertRaisesRegex(ValueError, "Missing required"):
            compile_template_patch(self.profile, self.contract, missing)

        invalid = dict(self.data)
        invalid["applications"] = [{"name": "Only a name"}]
        with self.assertRaisesRegex(ValueError, "missing"):
            compile_template_patch(self.profile, self.contract, invalid)

    def test_merged_and_nested_template_rows_are_rejected(self):
        merged_source = self.root / "merged.docx"
        merged = Document()
        table = merged.add_table(rows=2, cols=3)
        table.cell(1, 0).merge(table.cell(1, 1)).text = "Merged template"
        table.cell(1, 2).text = "Value"
        merged.save(str(merged_source))
        merged_bundle = self.root / "merged-bundle"
        profile = export_bundle(merged_source, merged_bundle)
        contract = {
            "contract_version": "0.1",
            "document_type": "merged_test",
            "binding": template_contract_binding(profile),
            "fields": {
                "rows": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {"a": 0, "b": 1, "c": 2},
                }
            },
        }
        with self.assertRaisesRegex(ValueError, "Merged"):
            compile_template_patch(
                profile, contract, {"rows": [{"a": "1", "b": "2", "c": "3"}]}
            )

        nested_source = self.root / "nested.docx"
        nested = Document()
        table = nested.add_table(rows=2, cols=1)
        table.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
        nested.save(str(nested_source))
        nested_bundle = self.root / "nested-bundle"
        profile = export_bundle(nested_source, nested_bundle)
        contract["binding"] = template_contract_binding(profile)
        contract["fields"]["rows"]["columns"] = {"a": 0}
        with self.assertRaisesRegex(ValueError, "Nested"):
            compile_template_patch(profile, contract, {"rows": [{"a": "1"}]})

    def test_noop_optional_population_preserves_every_package_part(self):
        contract = {
            "contract_version": "0.1",
            "document_type": "optional_only",
            "binding": template_contract_binding(self.profile),
            "fields": {
                "optional_summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": False,
                }
            },
        }
        patch = compile_template_patch(self.profile, contract, {})
        self.assertEqual(patch, {"contract": self.profile["patch_contract"]})
        output = self.root / "noop.docx"
        render_bundle(self.bundle, output, patch)
        with (
            zipfile.ZipFile(self.source) as original,
            zipfile.ZipFile(output) as generated,
        ):
            self.assertEqual(original.namelist(), generated.namelist())
            for name in original.namelist():
                self.assertEqual(original.read(name), generated.read(name), name)


if __name__ == "__main__":
    unittest.main()
