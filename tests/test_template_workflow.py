import contextlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from docxunravel import export_bundle, validate_docx_package
from docxunravel.cli import main
from docxunravel.template_contract import (
    populate_template,
    preview_template_population,
    template_contract_binding,
)


class TemplateWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.data = {
            "customer": {"name": "Northstar Research Group"},
            "project": {"scope": "A concise project summary."},
            "application_inventory": [
                {
                    "id": "APP-1",
                    "name": "Portal",
                    "owner": "Team A",
                    "business_criticality": "high",
                },
                {
                    "id": "APP-2",
                    "name": "Billing",
                    "owner": "Team B",
                    "business_criticality": "medium",
                },
                {
                    "id": "APP-3",
                    "name": "Support",
                    "owner": "Team C",
                    "business_criticality": "low",
                },
            ],
        }
        self.data_path = self.root / "data.json"
        self.data_path.write_text(json.dumps(self.data, indent=2), encoding="utf-8")
        self.source = self.root / "workflow-template.docx"
        document = Document()
        document.add_paragraph("Project for {{ CUSTOMER }}", style="Title")
        document.add_paragraph("Original project scope")
        table = document.add_table(rows=2, cols=4)
        table.style = "Table Grid"
        for index, value in enumerate(("ID", "Application", "Owner", "Criticality")):
            table.cell(0, index).text = value
        for index, value in enumerate(("APP-ID", "Template app", "Owner", "medium")):
            table.cell(1, index).text = value
        document.save(str(self.source))
        self.bundle = self.root / "bundle"
        self.profile = export_bundle(self.source, self.bundle)
        self.contract = {
            "contract_version": "0.1",
            "document_type": "project_form",
            "binding": template_contract_binding(self.profile),
            "fields": {
                "customer_brand": {
                    "type": "placeholder",
                    "target": "body.p0",
                    "token": "CUSTOMER",
                    "required": True,
                    "data_path": "/customer/name",
                },
                "executive_summary": {
                    "type": "text",
                    "target": "body.p1",
                    "required": True,
                    "data_path": "/project/scope",
                },
                "application_inventory": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {
                        "id": 0,
                        "name": 1,
                        "owner": 2,
                        "business_criticality": 3,
                    },
                    "data_path": "/application_inventory",
                },
            },
        }
        self.contract_path = self.root / "contract.json"
        self.contract_path.write_text(
            json.dumps(self.contract, indent=2), encoding="utf-8"
        )

    def tearDown(self):
        self.temporary.cleanup()

    def run_main(self, arguments):
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            result = main(arguments)
        return result, json.loads(output.getvalue())

    def test_cli_preview_patch_and_population_workflow(self):
        result, validation = self.run_main(
            ["template-validate", str(self.bundle), str(self.contract_path)]
        )
        self.assertEqual(result, 0)
        self.assertTrue(validation["valid"])
        self.assertEqual(validation["field_count"], 3)

        preview_path = self.root / "preview.json"
        result, preview = self.run_main(
            [
                "template-preview",
                str(self.bundle),
                str(self.contract_path),
                str(self.data_path),
                "--output",
                str(preview_path),
            ]
        )
        self.assertEqual(result, 0)
        self.assertTrue(preview["valid"])
        self.assertTrue(preview_path.exists())
        self.assertFalse((self.root / "generated.docx").exists())

        patch_path = self.root / "compiled-patch.json"
        result, patch_result = self.run_main(
            [
                "template-patch",
                str(self.bundle),
                str(self.contract_path),
                str(self.data_path),
                str(patch_path),
            ]
        )
        self.assertEqual(result, 0)
        self.assertTrue(patch_result["valid"])
        self.assertEqual(
            json.loads(patch_path.read_text(encoding="utf-8"))["contract"],
            self.profile["patch_contract"],
        )

        generated = self.root / "generated.docx"
        emitted_patch = self.root / "rendered-patch.json"
        fidelity = self.root / "fidelity.json"
        result, population = self.run_main(
            [
                "template-populate",
                str(self.bundle),
                str(self.contract_path),
                str(self.data_path),
                str(generated),
                "--patch-output",
                str(emitted_patch),
                "--fidelity-output",
                str(fidelity),
            ]
        )
        self.assertEqual(result, 0)
        self.assertTrue(population["preview"]["valid"])
        self.assertTrue(population["fidelity"]["passed"])
        self.assertTrue(validate_docx_package(generated)["valid"])
        self.assertTrue(emitted_patch.exists())
        self.assertTrue(fidelity.exists())
        generated_document = Document(str(generated))
        self.assertIn("Northstar Research Group", generated_document.paragraphs[0].text)
        self.assertEqual(len(generated_document.tables[0].rows), 4)

    def test_unsupported_layout_is_reported_and_output_is_not_created(self):
        source = self.root / "unsupported.docx"
        document = Document()
        table = document.add_table(rows=2, cols=4)
        table.cell(1, 0).merge(table.cell(1, 1)).text = "Merged"
        table.cell(1, 2).text = "Owner"
        table.cell(1, 3).text = "Criticality"
        document.save(str(source))
        bundle = self.root / "unsupported-bundle"
        profile = export_bundle(source, bundle)
        contract = {
            "contract_version": "0.1",
            "document_type": "unsupported_form",
            "binding": template_contract_binding(profile),
            "fields": {
                "applications": {
                    "type": "table_rows",
                    "target": "body.t0",
                    "required": True,
                    "template_row": 1,
                    "columns": {
                        "id": 0,
                        "name": 1,
                        "owner": 2,
                        "business_criticality": 3,
                    },
                    "data_path": "/application_inventory",
                }
            },
        }
        preview = preview_template_population(bundle, contract, self.data)
        self.assertFalse(preview["valid"])
        self.assertTrue(preview["unsupported_operations"])
        output = self.root / "must-not-exist.docx"
        with self.assertRaisesRegex(ValueError, "unsupported operations"):
            populate_template(bundle, output, contract, self.data)
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
