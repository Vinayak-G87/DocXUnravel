"""Markdown-to-DOCX population, rendering, and contract-aware comparison."""

from __future__ import annotations

import json
import os
import tempfile
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from docx import Document

from .core import (
    _checkbox_index,
    _profile_index,
    _profile_row_target,
    create_fidelity_report,
    profile_document,
    render_bundle,
)
from .markdown_ingestion import (
    DEFAULT_MARKDOWN_LIMITS,
    MarkdownLimits,
    markdown_to_data,
    parse_markdown_text,
    render_markdown_template,
)
from .template_contract import (
    compile_template_patch,
    preview_template_population,
    validate_template_contract,
)


def _load_object(path: str | Path) -> dict[str, Any]:
    with Path(path).open(encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise TypeError(f"JSON root must be an object: {path}")
    return value


def _pointer_segments(pointer: str) -> list[str]:
    if not pointer:
        return []
    if not pointer.startswith("/"):
        raise ValueError(f"Invalid JSON pointer: {pointer!r}")
    return [
        segment.replace("~1", "/").replace("~0", "~")
        for segment in pointer[1:].split("/")
    ]


def _assign_pointer(root: dict[str, Any], pointer: str, value: Any) -> None:
    segments = _pointer_segments(pointer)
    if not segments:
        if not isinstance(value, dict):
            raise TypeError("A root data_path requires an object value")
        root.update(value)
        return
    current = root
    for segment in segments[:-1]:
        existing = current.setdefault(segment, {})
        if not isinstance(existing, dict):
            raise ValueError(f"Conflicting Markdown data path: {pointer}")
        current = existing
    current[segments[-1]] = value


def markdown_contract_data(
    markdown: str | Path,
    contract: Mapping[str, Any],
    *,
    limits: MarkdownLimits = DEFAULT_MARKDOWN_LIMITS,
) -> dict[str, Any]:
    """Map named Markdown fields, tables, and task boxes to contract input data."""
    parsed = markdown_to_data(markdown, limits=limits)
    values: dict[str, Any] = {}
    consumed: set[tuple[str, str]] = set()

    fields = contract.get("fields")
    if not isinstance(fields, dict):
        raise TypeError("Template contract fields must be an object")
    for field_name, specification in fields.items():
        if not isinstance(specification, dict):
            raise TypeError(f"Template field {field_name!r} must be an object")
        field_type = specification.get("type")
        value: Any
        if (
            field_type in {"table_rows", "table_keyed_rows"}
            and field_name in parsed["tables"]
        ):
            value = parsed["tables"][field_name]["rows"]
            consumed.add(("tables", field_name))
        elif field_type == "checkbox_group":
            option_names = specification.get("checkboxes", {})
            value = [
                name for name in option_names if parsed["checkboxes"].get(name) is True
            ]
            for name in option_names:
                if name in parsed["checkboxes"]:
                    consumed.add(("checkboxes", name))
                if name in parsed["fields"]:
                    consumed.add(("fields", name))
        elif field_name in parsed["fields"]:
            value = parsed["fields"][field_name]
            consumed.add(("fields", field_name))
        elif field_name in parsed["sections"]:
            value = parsed["sections"][field_name]
            consumed.add(("sections", field_name))
        else:
            continue

        pointer = specification.get("data_path")
        if isinstance(pointer, str):
            _assign_pointer(values, pointer, value)
        else:
            values[field_name] = value

    unconsumed = {
        category: sorted(
            set(parsed[category])
            - {name for kind, name in consumed if kind == category}
        )
        for category in ("fields", "tables", "checkboxes", "sections")
    }
    unconsumed = {key: names for key, names in unconsumed.items() if names}
    if unconsumed:
        values["_markdown_unmapped"] = unconsumed
    return values


def compile_markdown_patch(
    bundle: str | Path,
    contract: Mapping[str, Any],
    markdown: str | Path,
    *,
    limits: MarkdownLimits = DEFAULT_MARKDOWN_LIMITS,
) -> dict[str, Any]:
    """Compile named Markdown content into a validated bundle patch."""
    bundle_path = Path(bundle)
    profile = _load_object(bundle_path / "profile.json")
    data = markdown_contract_data(markdown, contract, limits=limits)
    data.pop("_markdown_unmapped", None)
    return compile_template_patch(profile, dict(contract), data)


def preview_markdown_population(
    bundle: str | Path,
    contract: Mapping[str, Any],
    markdown: str | Path,
    *,
    limits: MarkdownLimits = DEFAULT_MARKDOWN_LIMITS,
) -> dict[str, Any]:
    """Preview Markdown population without writing a DOCX."""
    data = markdown_contract_data(markdown, contract, limits=limits)
    unmapped = data.pop("_markdown_unmapped", {})
    preview = preview_template_population(bundle, dict(contract), data)
    preview["markdown"] = {
        "source": str(markdown),
        "unmapped": unmapped,
    }
    return preview


def populate_markdown_docx(
    bundle: str | Path,
    output: str | Path,
    contract: Mapping[str, Any],
    markdown: str | Path,
    *,
    patch_output: str | Path | None = None,
    fidelity_output: str | Path | None = None,
    limits: MarkdownLimits = DEFAULT_MARKDOWN_LIMITS,
) -> dict[str, Any]:
    """Populate an existing DOCX bundle from named Markdown content."""
    bundle_path = Path(bundle)
    output_path = Path(output)
    preview = preview_markdown_population(
        bundle_path, contract, markdown, limits=limits
    )
    if not preview["valid"]:
        raise ValueError(
            "Markdown population contains unsupported operations: "
            + "; ".join(preview["unsupported_operations"])
        )
    patch = preview["patch"]
    result = render_bundle(bundle_path, output_path, patch)
    if patch_output is not None:
        Path(patch_output).write_text(
            json.dumps(patch, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    fidelity = None
    if fidelity_output is not None:
        with tempfile.TemporaryDirectory(
            prefix="docxunravel-markdown-source-"
        ) as temporary:
            source = Path(temporary) / "source.docx"
            profile = _load_object(bundle_path / "profile.json")
            render_bundle(bundle_path, source, {"contract": profile["patch_contract"]})
            fidelity = create_fidelity_report(
                source, output_path, bundle=bundle_path, patch=patch
            )
        Path(fidelity_output).write_text(
            json.dumps(fidelity, indent=2, ensure_ascii=False), encoding="utf-8"
        )
    return {
        "preview": preview,
        "render": result,
        "patch_output": str(patch_output) if patch_output is not None else None,
        "fidelity_output": str(fidelity_output)
        if fidelity_output is not None
        else None,
        "fidelity": fidelity,
    }


def _clear_document_body(document: Any) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag.rsplit("}", 1)[-1] != "sectPr":
            body.remove(child)


def render_markdown_docx(
    template: str | Path,
    data: Mapping[str, Any] | str | Path,
    output: str | Path,
    *,
    reference: str | Path | None = None,
    limits: MarkdownLimits = DEFAULT_MARKDOWN_LIMITS,
) -> dict[str, Any]:
    """Render a Markdown template and JSON data into a standalone DOCX."""
    markdown = render_markdown_template(template, data, limits=limits)
    model = parse_markdown_text(markdown, source_name=str(template), limits=limits)
    document = Document(str(reference)) if reference is not None else Document()
    if reference is not None:
        _clear_document_body(document)
    elif document.paragraphs and not document.paragraphs[0].text:
        paragraph = document.paragraphs[0]._element
        paragraph.getparent().remove(paragraph)

    for block in model["blocks"]:
        block_type = block["type"]
        if block_type == "heading":
            document.add_heading(block["text"], level=block["level"])
        elif block_type == "paragraph":
            document.add_paragraph(block["text"])
        elif block_type == "checkbox":
            glyph = "☒" if block["checked"] else "☐"
            document.add_paragraph(f"{glyph} {block['text']}")
        elif block_type == "table":
            rows = [block["columns"], *block["rows"]]
            table = document.add_table(rows=len(rows), cols=len(block["columns"]))
            try:
                table.style = "Table Grid"
            except KeyError:
                pass
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = value
        else:
            raise ValueError(f"Unsupported Markdown block type: {block_type}")

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.name}.",
        suffix=".tmp.docx",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        document.save(str(temporary_path))
        os.replace(temporary_path, output_path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise
    profile = profile_document(output_path)
    return {
        "output": str(output_path),
        "source": model["source"],
        "summary": profile["summary"],
        "reference": str(reference) if reference is not None else None,
    }


def compare_markdown_to_docx(
    markdown: str | Path,
    generated: str | Path,
    bundle: str | Path,
    contract: Mapping[str, Any],
    *,
    limits: MarkdownLimits = DEFAULT_MARKDOWN_LIMITS,
) -> dict[str, Any]:
    """Compare contract-mapped Markdown values with a generated DOCX."""
    data = markdown_contract_data(markdown, contract, limits=limits)
    unmapped = data.pop("_markdown_unmapped", {})
    original_profile = _load_object(Path(bundle) / "profile.json")
    normalized = validate_template_contract(original_profile, dict(contract))
    generated_profile = profile_document(generated)
    index = _profile_index(generated_profile)
    checkbox_index = _checkbox_index(generated_profile)
    fields: list[dict[str, Any]] = []

    from .template_contract import _MISSING, _field_value

    for name, specification in normalized["fields"].items():
        expected = _field_value(data, name, specification)
        if expected is _MISSING:
            continue
        field_type = specification["type"]
        if field_type in {"text", "table_cell"}:
            actual = index.get(specification["target"], {}).get("text")
        elif field_type == "placeholder":
            actual = (
                expected
                if isinstance(expected, str)
                and any(
                    expected in paragraph.get("text", "")
                    for paragraph in index.values()
                    if paragraph.get("type") == "paragraph"
                )
                else None
            )
        elif field_type == "checkbox_group":
            original_checkboxes = _checkbox_index(original_profile)

            def generated_checkbox(checkbox_id: str) -> dict[str, Any]:
                if checkbox_id in checkbox_index:
                    return checkbox_index[checkbox_id]
                locator = original_checkboxes.get(checkbox_id, {}).get("locator")
                return next(
                    (
                        item
                        for item in checkbox_index.values()
                        if item.get("locator") == locator
                    ),
                    {},
                )

            actual = sorted(
                option
                for option, checkbox_id in specification["checkboxes"].items()
                if generated_checkbox(checkbox_id).get("checked") is True
            )
            expected = sorted(expected) if isinstance(expected, list) else expected
        elif field_type == "table_keyed_rows":
            table = index.get(specification["target"])
            actual_rows: list[Any] = []
            if table is not None and isinstance(expected, list):
                key_specification = specification["key"]
                key_column = key_specification["column"]
                insensitive = (
                    key_specification.get("match", "exact") == "case_insensitive"
                )
                for expected_row in expected:
                    from .template_contract import _resolve_data_path

                    key_value = _resolve_data_path(
                        expected_row,
                        key_specification.get("data_path", "/key"),
                        name,
                    )
                    wanted = str(key_value).strip()
                    matched = None
                    for row_index, row in enumerate(table.get("rows", [])):
                        cells = row.get("cells", [])
                        if key_column >= len(cells):
                            continue
                        candidate = cells[key_column].get("text", "").strip()
                        if (
                            (candidate.casefold() == wanted.casefold())
                            if insensitive
                            else (candidate == wanted)
                        ):
                            matched = row_index
                            break
                    if matched is None:
                        actual_rows.append(None)
                        continue
                    actual_rows.append(
                        {
                            key: (
                                entry.get("text", "")
                                if (
                                    entry := _profile_row_target(
                                        table, matched, target["target"]
                                    )
                                )
                                else ""
                            )
                            for key, target in specification["row_targets"].items()
                        }
                    )
            from .template_contract import _resolve_data_path, _table_cell_text

            expected = [
                {
                    key: _table_cell_text(
                        _resolve_data_path(
                            row,
                            target.get("data_path", "/" + key),
                            f"{name}.{key}",
                        ),
                        name,
                        row_index,
                    )
                    for key, target in specification["row_targets"].items()
                }
                for row_index, row in enumerate(expected)
            ]
            actual = actual_rows
        else:
            table = index.get(specification["target"])
            actual_rows = []
            if table is not None and isinstance(expected, list):
                start = specification["template_row"]
                for row_offset, _ in enumerate(expected):
                    row_index = start + row_offset
                    if "columns" in specification:
                        row = table.get("rows", [])[row_index]
                        actual_rows.append(
                            {
                                key: row["cells"][column].get("text", "")
                                for key, column in specification["columns"].items()
                            }
                        )
                    else:
                        row_values: dict[str, str] = {}
                        for key, target in specification["row_targets"].items():
                            entry = _profile_row_target(
                                table, row_index, target["target"]
                            )
                            row_values[key] = entry.get("text", "") if entry else ""
                        actual_rows.append(row_values)
            if isinstance(expected, list):
                from .template_contract import _resolve_data_path, _table_cell_text

                if "columns" in specification:
                    expected = [
                        {
                            key: _table_cell_text(row[key], name, row_index)
                            for key in specification["columns"]
                        }
                        for row_index, row in enumerate(expected)
                    ]
                else:
                    expected = [
                        {
                            key: _table_cell_text(
                                _resolve_data_path(
                                    row,
                                    target.get("data_path", "/" + key),
                                    f"{name}.{key}",
                                ),
                                name,
                                row_index,
                            )
                            for key, target in specification["row_targets"].items()
                        }
                        for row_index, row in enumerate(expected)
                    ]
            actual = actual_rows
        fields.append(
            {
                "field": name,
                "type": field_type,
                "equal": actual == expected,
                "expected": expected,
                "actual": actual,
            }
        )

    return {
        "schema_version": "1.0",
        "comparison": "markdown-to-docx-contract",
        "equal": all(field["equal"] for field in fields),
        "markdown": str(markdown),
        "docx": str(generated),
        "unmapped_markdown": unmapped,
        "summary": {
            "fields": len(fields),
            "matched": sum(field["equal"] for field in fields),
            "different": sum(not field["equal"] for field in fields),
        },
        "fields": fields,
    }
