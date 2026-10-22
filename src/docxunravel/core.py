"""DocXUnravel: profile DOCX templates and render AI-generated values.

The original DOCX is always used as the rendering base, so styles, images, page
layout, numbering, and other package parts are retained. The JSON is a content
contract, not a lossy DOCX serialization format.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import posixpath
import re
import sys
import tempfile
import zipfile
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

SCHEMA_VERSION = "1.1"
TOKEN_PATTERN = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_.-]*)\s*\}\}")
FORMAT_POLICIES = {
    "preserve_first_run",
    "preserve_common_formatting",
    "replace_text_only",
    "reject_mixed_runs",
}
FORMAT_POLICY_ALIASES = {
    "first_run": "preserve_first_run",
    "reject_mixed_formatting": "reject_mixed_runs",
}
ASSET_REGISTRY_VERSION = "0.1"
RELATIONSHIP_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_RELATIONSHIP_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
WORDPROCESSING_DRAWING_NS = (
    "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
)
ASSET_RELATIONSHIP_KINDS = {
    "image": "image",
    "chart": "chart",
    "oleObject": "embedded_object",
    "package": "embedded_object",
    "hyperlink": "relationship",
    "video": "media",
    "audio": "media",
    "diagramData": "drawing",
    "diagramLayout": "drawing",
    "diagramQuickStyle": "drawing",
    "diagramColors": "drawing",
}


@dataclass(frozen=True)
class PackageLimits:
    max_entries: int = 2_048
    max_entry_size: int = 100 * 1024 * 1024
    max_total_uncompressed_size: int = 500 * 1024 * 1024
    max_compression_ratio: float = 200.0

    def validate(self) -> None:
        if self.max_entries < 1:
            raise ValueError("max_entries must be positive")
        if self.max_entry_size < 1 or self.max_total_uncompressed_size < 1:
            raise ValueError("package size limits must be positive")
        if self.max_entry_size > self.max_total_uncompressed_size:
            raise ValueError("max_entry_size cannot exceed max_total_uncompressed_size")
        if self.max_compression_ratio < 1:
            raise ValueError("max_compression_ratio must be at least 1")


DEFAULT_PACKAGE_LIMITS = PackageLimits()

UNSAFE_EDIT_TAGS = {
    qn("w:fldChar"): "field",
    qn("w:instrText"): "field_instruction",
    qn("w:hyperlink"): "hyperlink",
    qn("w:sdt"): "content_control",
    qn("w:drawing"): "drawing",
    qn("w:object"): "embedded_object",
    qn("w:pict"): "legacy_drawing",
    qn("w:tab"): "tab",
    qn("w:br"): "line_or_page_break",
    qn("w:cr"): "carriage_return",
    qn("w:noBreakHyphen"): "nonbreaking_hyphen",
    qn("w:softHyphen"): "soft_hyphen",
}


def _secure_xml_parser(*, remove_blank_text: bool = False) -> Any:
    return etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        dtd_validation=False,
        attribute_defaults=False,
        huge_tree=False,
        recover=False,
        remove_blank_text=remove_blank_text,
    )


def _parse_xml_bytes(data: bytes, part_name: str) -> Any:
    if re.search(rb"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.IGNORECASE):
        raise ValueError(f"DTD or entity declarations are not allowed in {part_name}")
    try:
        return etree.fromstring(data, _secure_xml_parser())
    except etree.XMLSyntaxError as error:
        raise ValueError(f"Invalid XML part {part_name}: {error}") from error


def inspect_docx_archive(
    path: str | Path,
    limits: PackageLimits = DEFAULT_PACKAGE_LIMITS,
    *,
    validate_xml: bool = True,
) -> dict[str, Any]:
    """Inspect a DOCX ZIP without extracting it and enforce resource bounds."""
    limits.validate()
    path = Path(path)
    try:
        with zipfile.ZipFile(path) as archive:
            all_infos = archive.infolist()
            if len(all_infos) > limits.max_entries:
                raise ValueError(
                    f"DOCX has {len(all_infos)} entries; limit is {limits.max_entries}"
                )
            all_names = [info.filename for info in all_infos]
            if len(all_names) != len(set(all_names)):
                raise ValueError("DOCX contains duplicate package entry names")
            for info in all_infos:
                member = PurePosixPath(info.filename)
                if member.is_absolute() or ".." in member.parts:
                    raise ValueError(f"Unsafe DOCX package path: {info.filename}")
                if info.flag_bits & 0x1:
                    raise ValueError(
                        f"Encrypted package entries are unsupported: {info.filename}"
                    )

            infos = [info for info in all_infos if not info.is_dir()]
            total_size = 0
            highest_ratio = 0.0
            xml_parts = 0
            for info in infos:
                if info.file_size > limits.max_entry_size:
                    raise ValueError(
                        f"Package entry exceeds size limit: {info.filename} "
                        f"({info.file_size} > {limits.max_entry_size})"
                    )
                total_size += info.file_size
                if total_size > limits.max_total_uncompressed_size:
                    raise ValueError(
                        "DOCX total uncompressed size exceeds limit "
                        f"({total_size} > {limits.max_total_uncompressed_size})"
                    )
                ratio = (
                    info.file_size / info.compress_size
                    if info.compress_size
                    else (float("inf") if info.file_size else 1.0)
                )
                highest_ratio = max(highest_ratio, ratio)
                if ratio > limits.max_compression_ratio:
                    raise ValueError(
                        f"Package entry compression ratio exceeds limit: {info.filename} "
                        f"({ratio:.2f} > {limits.max_compression_ratio})"
                    )
                if validate_xml and info.filename.endswith((".xml", ".rels")):
                    _parse_xml_bytes(archive.read(info), info.filename)
                    xml_parts += 1
    except zipfile.BadZipFile as error:
        raise ValueError(f"Invalid DOCX ZIP package: {path}") from error
    return {
        "valid": True,
        "entry_count": len(all_infos),
        "file_entry_count": len(infos),
        "total_uncompressed_size": total_size,
        "highest_compression_ratio": round(highest_ratio, 3),
        "xml_part_count": xml_parts,
        "limits": asdict(limits),
    }


def _round(value: Any) -> float | None:
    return round(value.pt, 2) if value is not None else None


def _enum(value: Any) -> str | None:
    if value is None:
        return None
    return getattr(value, "name", str(value))


def _color(font: Any) -> str | None:
    rgb = font.color.rgb
    return str(rgb) if rgb is not None else None


def _font(run: Any) -> dict[str, Any]:
    font = run.font
    style_font = run.style.font if run.style is not None else None
    return {
        "name": font.name,
        "size_pt": _round(font.size),
        "bold": run.bold,
        "italic": run.italic,
        "underline": _enum(run.underline),
        "color_rgb": _color(font),
        "highlight": _enum(font.highlight_color),
        "all_caps": font.all_caps,
        "small_caps": font.small_caps,
        "strike": font.strike,
        "subscript": font.subscript,
        "superscript": font.superscript,
        "inherited_hint": {
            "style": run.style.name if run.style is not None else None,
            "font_name": style_font.name if style_font is not None else None,
            "size_pt": _round(style_font.size) if style_font is not None else None,
        },
    }


def _paragraph_format(paragraph: Paragraph) -> dict[str, Any]:
    fmt = paragraph.paragraph_format
    return {
        "alignment": _enum(paragraph.alignment),
        "left_indent_pt": _round(fmt.left_indent),
        "right_indent_pt": _round(fmt.right_indent),
        "first_line_indent_pt": _round(fmt.first_line_indent),
        "space_before_pt": _round(fmt.space_before),
        "space_after_pt": _round(fmt.space_after),
        "line_spacing": str(fmt.line_spacing) if fmt.line_spacing is not None else None,
        "keep_together": fmt.keep_together,
        "keep_with_next": fmt.keep_with_next,
        "page_break_before": fmt.page_break_before,
    }


def _tokens(text: str) -> list[str]:
    return list(dict.fromkeys(TOKEN_PATTERN.findall(text)))


def _paragraph_profile(paragraph: Paragraph, element_id: str) -> dict[str, Any]:
    text = paragraph.text or ""
    xml_path = paragraph._p.getroottree().getpath(paragraph._p)
    part_name = str(paragraph.part.partname).lstrip("/")
    native_id = paragraph._p.get(qn("w14:paraId"))
    fingerprint = hashlib.sha256(
        (
            part_name
            + "\0"
            + text
            + "\0"
            + (paragraph.style.name if paragraph.style else "")
        ).encode()
    ).hexdigest()
    revision_tags = {
        child.tag.rsplit("}", 1)[-1]
        for child in paragraph._p.iter()
        if child.tag
        in {
            qn("w:ins"),
            qn("w:del"),
            qn("w:moveFrom"),
            qn("w:moveTo"),
            qn("w:pPrChange"),
            qn("w:rPrChange"),
        }
    }
    p_pr = paragraph._p.pPr
    numbered = bool(p_pr is not None and p_pr.numPr is not None)
    unsupported_structures = sorted(
        {
            name
            for node in paragraph._p.iter()
            if (name := UNSAFE_EDIT_TAGS.get(node.tag)) is not None
        }
    )
    return {
        "id": element_id,
        "stable_id": f"{part_name}#paraId={native_id}"
        if native_id
        else f"sha256:{fingerprint}",
        "native_id": native_id,
        "type": "paragraph",
        "text": text,
        "locator": {"part": part_name, "xpath": xml_path},
        "has_revisions": bool(revision_tags),
        "revision_types": sorted(revision_tags),
        "unsupported_edit_structures": unsupported_structures,
        "has_mixed_run_formatting": len(
            {
                json.dumps(_font(run), sort_keys=True)
                for run in paragraph.runs
                if run.text
            }
        )
        > 1,
        "style": paragraph.style.name if paragraph.style is not None else None,
        "numbered": numbered,
        "tokens": _tokens(text),
        "format": _paragraph_format(paragraph),
        "runs": [
            {
                "index": index,
                "text": run.text,
                "style": run.style.name if run.style is not None else None,
                "format": _font(run),
            }
            for index, run in enumerate(paragraph.runs)
        ],
    }


def _iter_blocks(parent: DocumentObject | _Cell) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in their actual XML order."""
    parent_element = (
        parent.element.body if isinstance(parent, DocumentObject) else parent._tc
    )
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _table_profile(table: Table, element_id: str) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for row_index, row in enumerate(table.rows):
        cells: list[dict[str, Any]] = []
        seen_cell_xml: set[int] = set()
        for cell_index, cell in enumerate(row.cells):
            cell_id = f"{element_id}.r{row_index}.c{cell_index}"
            xml_identity = id(cell._tc)
            duplicate_merged_cell = xml_identity in seen_cell_xml
            seen_cell_xml.add(xml_identity)
            cells.append(
                {
                    "id": cell_id,
                    "type": "cell",
                    "text": cell.text,
                    "locator": {
                        "part": str(cell.part.partname).lstrip("/"),
                        "xpath": cell._tc.getroottree().getpath(cell._tc),
                    },
                    "duplicate_merged_cell": duplicate_merged_cell,
                    "vertical_alignment": _enum(cell.vertical_alignment),
                    "elements": _container_profile(cell, cell_id),
                }
            )
        rows.append({"index": row_index, "cells": cells})
    return {
        "id": element_id,
        "type": "table",
        "locator": {
            "part": str(table.part.partname).lstrip("/"),
            "xpath": table._tbl.getroottree().getpath(table._tbl),
        },
        "style": table.style.name if table.style is not None else None,
        "rows": rows,
        "row_count": len(table.rows),
        "column_count": len(table.columns),
    }


def _container_profile(
    parent: DocumentObject | _Cell, prefix: str
) -> list[dict[str, Any]]:
    elements: list[dict[str, Any]] = []
    paragraph_index = 0
    table_index = 0
    for block in _iter_blocks(parent):
        if isinstance(block, Paragraph):
            elements.append(_paragraph_profile(block, f"{prefix}.p{paragraph_index}"))
            paragraph_index += 1
        else:
            elements.append(_table_profile(block, f"{prefix}.t{table_index}"))
            table_index += 1
    return elements


def _header_footer_profile(container: Any, element_id: str) -> dict[str, Any]:
    # Header/Footer has the same paragraph/table surface but no public ordered iterator.
    elements: list[dict[str, Any]] = []
    p_index = t_index = 0
    for child in container._element.iterchildren():
        if child.tag == qn("w:p"):
            elements.append(
                _paragraph_profile(
                    Paragraph(child, container), f"{element_id}.p{p_index}"
                )
            )
            p_index += 1
        elif child.tag == qn("w:tbl"):
            elements.append(
                _table_profile(Table(child, container), f"{element_id}.t{t_index}")
            )
            t_index += 1
    return {"id": element_id, "elements": elements}


def _collect_paragraphs(elements: Iterable[dict[str, Any]]) -> Iterator[dict[str, Any]]:
    for element in elements:
        if element["type"] == "paragraph":
            yield element
        else:
            for row in element["rows"]:
                for cell in row["cells"]:
                    yield from _collect_paragraphs(cell["elements"])


def _style_profile(document: DocumentObject) -> list[dict[str, Any]]:
    styles = []
    for style in document.styles:
        if style.type not in (
            WD_STYLE_TYPE.PARAGRAPH,
            WD_STYLE_TYPE.CHARACTER,
            WD_STYLE_TYPE.TABLE,
        ):
            continue
        entry: dict[str, Any] = {
            "name": style.name,
            "type": _enum(style.type),
            "builtin": style.builtin,
            "base_style": style.base_style.name
            if style.base_style is not None
            else None,
        }
        if hasattr(style, "font"):
            entry["font"] = {
                "name": style.font.name,
                "size_pt": _round(style.font.size),
                "bold": style.font.bold,
                "italic": style.font.italic,
                "color_rgb": _color(style.font),
            }
        styles.append(entry)
    return styles


def _relationship_owner(rels_name: str) -> str | None:
    if rels_name == "_rels/.rels":
        return None
    if "/_rels/" not in rels_name or not rels_name.endswith(".rels"):
        return None
    directory, filename = rels_name.split("/_rels/", 1)
    return posixpath.join(directory, filename[: -len(".rels")])


def _content_type_lookup(root: Any, package_names: set[str]) -> dict[str, str]:
    defaults: dict[str, str] = {}
    overrides: dict[str, str] = {}
    for node in root:
        if node.tag.endswith("Default"):
            defaults[node.get("Extension", "").lower()] = node.get("ContentType", "")
        elif node.tag.endswith("Override"):
            overrides[node.get("PartName", "").lstrip("/")] = node.get(
                "ContentType", ""
            )
    result: dict[str, str] = {}
    for name in package_names:
        result[name] = overrides.get(
            name, defaults.get(PurePosixPath(name).suffix.lstrip(".").lower(), "")
        )
    return result


def _integer_attribute(node: Any, name: str) -> int | None:
    if node is None or node.get(name) is None:
        return None
    try:
        return int(node.get(name))
    except (TypeError, ValueError):
        return None


def _position_metadata(container: Any, axis: str) -> dict[str, Any] | None:
    node = container.find(f".//{{{WORDPROCESSING_DRAWING_NS}}}position{axis}")
    if node is None:
        return None
    align = node.find(f"{{{WORDPROCESSING_DRAWING_NS}}}align")
    offset = node.find(f"{{{WORDPROCESSING_DRAWING_NS}}}posOffset")
    return {
        "relative_from": node.get("relativeFrom"),
        "align": align.text if align is not None else None,
        "offset_emu": int(offset.text)
        if offset is not None and offset.text and offset.text.lstrip("-").isdigit()
        else None,
    }


def _relationship_placements(root: Any, relationship_id: str) -> list[dict[str, Any]]:
    references = root.xpath(
        ".//*[@r:embed=$rid or @r:id=$rid or @r:link=$rid]",
        namespaces={"r": OFFICE_RELATIONSHIP_NS},
        rid=relationship_id,
    )
    placements: list[dict[str, Any]] = []
    for reference in references:
        ancestors = [reference, *reference.iterancestors()]
        drawing = next(
            (
                node
                for node in ancestors
                if etree.QName(node).localname in {"drawing", "object", "pict"}
            ),
            reference,
        )
        container = next(
            (
                node
                for node in ancestors
                if node.tag
                in {
                    f"{{{WORDPROCESSING_DRAWING_NS}}}inline",
                    f"{{{WORDPROCESSING_DRAWING_NS}}}anchor",
                }
            ),
            drawing,
        )
        extent = container.find(f".//{{{WORDPROCESSING_DRAWING_NS}}}extent")
        width_emu = _integer_attribute(extent, "cx")
        height_emu = _integer_attribute(extent, "cy")
        doc_properties = container.find(f".//{{{WORDPROCESSING_DRAWING_NS}}}docPr")
        placements.append(
            {
                "object_type": etree.QName(drawing).localname,
                "mode": etree.QName(container).localname,
                "xpath": root.getroottree().getpath(container),
                "dimensions": {
                    "width_emu": width_emu,
                    "height_emu": height_emu,
                    "width_pt": round(width_emu / 12700, 2)
                    if width_emu is not None
                    else None,
                    "height_pt": round(height_emu / 12700, 2)
                    if height_emu is not None
                    else None,
                },
                "horizontal": _position_metadata(container, "H"),
                "vertical": _position_metadata(container, "V"),
                "name": doc_properties.get("name")
                if doc_properties is not None
                else None,
                "description": doc_properties.get("descr")
                if doc_properties is not None
                else None,
                "title": doc_properties.get("title")
                if doc_properties is not None
                else None,
            }
        )
    return placements


def _asset_registry(source: Path) -> dict[str, Any]:
    items: list[dict[str, Any]] = []
    with zipfile.ZipFile(source) as archive:
        package_names = {
            info.filename for info in archive.infolist() if not info.is_dir()
        }
        content_types = _content_type_lookup(
            _parse_xml_bytes(
                archive.read("[Content_Types].xml"), "[Content_Types].xml"
            ),
            package_names,
        )
        owner_roots: dict[str, Any] = {}
        relationship_parts = sorted(
            name for name in package_names if name.endswith(".rels")
        )
        for rels_name in relationship_parts:
            owner = _relationship_owner(rels_name)
            if owner in package_names and owner and owner.endswith(".xml"):
                owner_roots[owner] = _parse_xml_bytes(archive.read(owner), owner)
            relationships = _parse_xml_bytes(archive.read(rels_name), rels_name)
            for relationship in relationships:
                relationship_type = relationship.get("Type", "")
                relation_name = relationship_type.rsplit("/", 1)[-1]
                target_mode = relationship.get("TargetMode", "Internal")
                target = relationship.get("Target")
                if not target:
                    continue
                package_location = None
                if target_mode != "External":
                    base = posixpath.dirname(owner) if owner else ""
                    package_location = posixpath.normpath(
                        posixpath.join(base, target.lstrip("/"))
                    )
                media_type = content_types.get(package_location or "") or None
                kind = ASSET_RELATIONSHIP_KINDS.get(relation_name)
                if (
                    kind is None
                    and media_type
                    and media_type.split("/", 1)[0]
                    in {
                        "image",
                        "audio",
                        "video",
                    }
                ):
                    kind = "media"
                if (
                    kind is None
                    and package_location
                    and package_location.startswith(
                        (
                            "word/media/",
                            "word/embeddings/",
                            "word/charts/",
                            "word/diagrams/",
                        )
                    )
                ):
                    kind = "relationship"
                if kind is None and target_mode == "External":
                    kind = "relationship"
                if kind is None:
                    continue
                content = (
                    archive.read(package_location)
                    if package_location in package_names
                    else None
                )
                relationship_id = relationship.get("Id")
                placements = (
                    _relationship_placements(owner_roots[owner], relationship_id or "")
                    if owner in owner_roots and relationship_id
                    else []
                )
                items.append(
                    {
                        "id": f"asset.{len(items)}",
                        "kind": kind,
                        "read_only": True,
                        "media_type": media_type,
                        "package_location": package_location,
                        "external_target": target
                        if target_mode == "External"
                        else None,
                        "relationship_id": relationship_id,
                        "relationship_type": relationship_type,
                        "owning_part": owner,
                        "sha256": _sha256_bytes(content)
                        if content is not None
                        else None,
                        "size": len(content) if content is not None else None,
                        "placements": placements,
                    }
                )
    counts = Counter(item["kind"] for item in items)
    return {
        "schema_version": ASSET_REGISTRY_VERSION,
        "read_only": True,
        "binary_content_included": False,
        "counts_by_kind": dict(sorted(counts.items())),
        "items": items,
    }


def profile_document(
    source: str | Path,
    limits: PackageLimits = DEFAULT_PACKAGE_LIMITS,
) -> dict[str, Any]:
    """Return a JSON-serializable, fine-grained profile of a bounded DOCX file."""
    source = Path(source)
    inspect_docx_archive(source, limits)
    with zipfile.ZipFile(source) as archive:
        source_parts = set(archive.namelist())
    assets = _asset_registry(source)
    document = Document(str(source))
    body = _container_profile(document, "body")
    headers: list[dict[str, Any]] = []
    footers: list[dict[str, Any]] = []
    seen_headers: set[int] = set()
    seen_footers: set[int] = set()

    sections = []
    for index, section in enumerate(document.sections):
        sections.append(
            {
                "index": index,
                "width_pt": _round(section.page_width),
                "height_pt": _round(section.page_height),
                "orientation": _enum(section.orientation),
                "margins_pt": {
                    "top": _round(section.top_margin),
                    "right": _round(section.right_margin),
                    "bottom": _round(section.bottom_margin),
                    "left": _round(section.left_margin),
                },
            }
        )
        for kind, container, seen, output in (
            ("header", section.header, seen_headers, headers),
            ("footer", section.footer, seen_footers, footers),
        ):
            part_name = str(container.part.partname).lstrip("/")
            # Accessing an undefined header/footer makes python-docx create an
            # in-memory part. Do not profile a part absent from the source ZIP.
            identity = id(container._element)
            if part_name in source_parts and identity not in seen:
                seen.add(identity)
                output.append(_header_footer_profile(container, f"{kind}{len(output)}"))

    all_elements = list(body)
    for container in headers + footers:
        all_elements.extend(container["elements"])
    paragraphs = list(_collect_paragraphs(all_elements))
    tokens = sorted(
        {token for paragraph in paragraphs for token in paragraph["tokens"]}
    )
    style_counts = Counter(p["style"] or "<none>" for p in paragraphs)
    font_counts = Counter(
        run["format"]["name"]
        or run["format"]["inherited_hint"]["font_name"]
        or "<inherited>"
        for paragraph in paragraphs
        for run in paragraph["runs"]
        if run["text"]
    )
    properties = document.core_properties

    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": {
            "filename": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        },
        "metadata": {
            "title": properties.title,
            "subject": properties.subject,
            "author": properties.author,
            "keywords": properties.keywords,
            "category": properties.category,
            "comments": properties.comments,
        },
        "summary": {
            "paragraph_count": len(paragraphs),
            "table_count": sum(1 for element in body if element["type"] == "table"),
            "inline_shape_count": len(document.inline_shapes),
            "asset_count": len(assets["items"]),
            "tokens": tokens,
            "style_usage": dict(style_counts.most_common()),
            "font_usage_by_run": dict(font_counts.most_common()),
        },
        "sections": sections,
        "styles": _style_profile(document),
        "body": body,
        "headers": headers,
        "footers": footers,
        "assets": assets,
        "replacement_contract": {
            "tokens": {token: "" for token in tokens},
            "elements": {
                paragraph["id"]: paragraph["text"]
                for paragraph in paragraphs
                if paragraph["text"].strip()
            },
        },
        "capabilities": {
            "preserved_by_using_original_as_base": [
                "page layout",
                "styles",
                "images",
                "relationships",
                "numbering",
                "themes",
            ],
            "profiled": [
                "ordered paragraphs and tables",
                "nested tables",
                "headers and footers",
                "runs and common run formatting",
                "paragraph formatting",
                "styles",
                "tokens",
                "read-only assets and relationship-backed content",
            ],
            "limited": [
                "text boxes and drawing-layer text",
                "tracked changes",
                "comments",
                "content controls",
                "fields",
                "footnotes/endnotes",
                "exact computed style inheritance",
            ],
        },
    }


def _walk_paragraphs(
    parent: DocumentObject | _Cell, prefix: str
) -> Iterator[tuple[str, Paragraph]]:
    p_index = t_index = 0
    for block in _iter_blocks(parent):
        if isinstance(block, Paragraph):
            yield f"{prefix}.p{p_index}", block
            p_index += 1
        else:
            table_id = f"{prefix}.t{t_index}"
            for row_index, row in enumerate(block.rows):
                for cell_index, cell in enumerate(row.cells):
                    yield from _walk_paragraphs(
                        cell, f"{table_id}.r{row_index}.c{cell_index}"
                    )
            t_index += 1


def _walk_header_footer(container: Any, prefix: str) -> Iterator[tuple[str, Paragraph]]:
    p_index = t_index = 0
    for child in container._element.iterchildren():
        if child.tag == qn("w:p"):
            yield f"{prefix}.p{p_index}", Paragraph(child, container)
            p_index += 1
        elif child.tag == qn("w:tbl"):
            table = Table(child, container)
            table_id = f"{prefix}.t{t_index}"
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    yield from _walk_paragraphs(
                        cell, f"{table_id}.r{row_index}.c{cell_index}"
                    )
            t_index += 1


def _text_nodes(paragraph: Paragraph) -> list[Any]:
    return list(paragraph._p.iter(qn("w:t")))


def _replace_text_nodes(paragraph: Paragraph, old: str, new: str) -> int:
    """Replace text even when Word split it across formatting runs."""
    count = 0
    while True:
        nodes = _text_nodes(paragraph)
        combined = "".join(node.text or "" for node in nodes)
        start = combined.find(old)
        if start < 0:
            return count
        end = start + len(old)
        offsets: list[tuple[int, int]] = []
        for node_index, node in enumerate(nodes):
            offsets.extend(
                (node_index, char_index) for char_index in range(len(node.text or ""))
            )
        start_node, start_char = offsets[start]
        end_node, end_char = offsets[end - 1]
        if start_node == end_node:
            text = nodes[start_node].text or ""
            nodes[start_node].text = text[:start_char] + new + text[end_char + 1 :]
        else:
            first_text = nodes[start_node].text or ""
            last_text = nodes[end_node].text or ""
            nodes[start_node].text = first_text[:start_char] + new
            for index in range(start_node + 1, end_node):
                nodes[index].text = ""
            nodes[end_node].text = last_text[end_char + 1 :]
        count += 1


def _set_paragraph_text(paragraph: Paragraph, value: str) -> None:
    nodes = _text_nodes(paragraph)
    if nodes:
        nodes[0].text = value
        for node in nodes[1:]:
            node.text = ""
    elif value:
        paragraph.add_run(value)


def render_document(
    source: str | Path,
    output: str | Path,
    data: dict[str, Any],
    *,
    verify_sha256: str | None = None,
) -> dict[str, Any]:
    """Render token and element values into a copy of the original DOCX."""
    source = Path(source)
    if verify_sha256:
        actual = hashlib.sha256(source.read_bytes()).hexdigest()
        if actual != verify_sha256:
            raise ValueError("Source DOCX does not match the profile SHA-256")

    document = Document(str(source))
    paragraphs: dict[str, Paragraph] = dict(_walk_paragraphs(document, "body"))
    seen_headers: set[int] = set()
    seen_footers: set[int] = set()
    header_index = footer_index = 0
    for section in document.sections:
        for kind, container, seen in (
            ("header", section.header, seen_headers),
            ("footer", section.footer, seen_footers),
        ):
            identity = id(container._element)
            if identity in seen:
                continue
            seen.add(identity)
            index = header_index if kind == "header" else footer_index
            paragraphs.update(_walk_header_footer(container, f"{kind}{index}"))
            if kind == "header":
                header_index += 1
            else:
                footer_index += 1

    element_values = data.get("elements", {})
    token_values = data.get("tokens", {})
    unknown_elements = sorted(set(element_values) - set(paragraphs))
    if unknown_elements:
        raise KeyError(f"Unknown element IDs: {', '.join(unknown_elements)}")

    changed_elements = 0
    for element_id, value in element_values.items():
        _set_paragraph_text(paragraphs[element_id], "" if value is None else str(value))
        changed_elements += 1

    token_replacements = 0
    for paragraph in paragraphs.values():
        for token, value in token_values.items():
            token_replacements += _replace_text_nodes(
                paragraph,
                "{{" + str(token) + "}}",
                "" if value is None else str(value),
            )
            # Also accept whitespace inside existing template braces.
            pattern = re.compile(r"\{\{\s*" + re.escape(str(token)) + r"\s*\}\}")
            while True:
                text = "".join(node.text or "" for node in _text_nodes(paragraph))
                match = pattern.search(text)
                if not match:
                    break
                token_replacements += _replace_text_nodes(
                    paragraph, match.group(0), "" if value is None else str(value)
                )

    document.save(str(output))
    return {
        "output": str(output),
        "changed_elements": changed_elements,
        "token_replacements": token_replacements,
    }


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _semantic_profile_hash(profile: dict[str, Any]) -> str:
    semantic_profile = {
        key: profile[key]
        for key in (
            "schema_version",
            "source",
            "body",
            "headers",
            "footers",
            "package",
        )
    }
    return _sha256_bytes(
        json.dumps(semantic_profile, sort_keys=True, separators=(",", ":")).encode()
    )


def export_bundle(
    source: str | Path,
    bundle: str | Path,
    limits: PackageLimits = DEFAULT_PACKAGE_LIMITS,
) -> dict[str, Any]:
    """Extract a DOCX losslessly within limits and write its semantic profile."""
    source = Path(source)
    bundle = Path(bundle)
    if bundle.exists() and any(bundle.iterdir()):
        raise FileExistsError(f"Bundle directory is not empty: {bundle}")
    package_dir = bundle / "package"
    package_dir.mkdir(parents=True, exist_ok=True)
    inspection = inspect_docx_archive(source, limits)
    profile = profile_document(source, limits)
    entries: list[dict[str, Any]] = []
    extracted_total = 0

    with zipfile.ZipFile(source) as archive:
        for info in archive.infolist():
            member = PurePosixPath(info.filename)
            if member.is_absolute() or ".." in member.parts:
                raise ValueError(f"Unsafe DOCX package path: {info.filename}")
            if info.is_dir():
                continue
            destination = package_dir.joinpath(*member.parts)
            destination.parent.mkdir(parents=True, exist_ok=True)
            digest = hashlib.sha256()
            extracted_size = 0
            with (
                archive.open(info) as source_part,
                destination.open("wb") as output_part,
            ):
                while chunk := source_part.read(1024 * 1024):
                    extracted_size += len(chunk)
                    extracted_total += len(chunk)
                    if extracted_size > limits.max_entry_size:
                        raise ValueError(
                            f"Extracted entry exceeds size limit: {info.filename}"
                        )
                    if extracted_total > limits.max_total_uncompressed_size:
                        raise ValueError("Extracted package exceeds total size limit")
                    digest.update(chunk)
                    output_part.write(chunk)
            if extracted_size != info.file_size:
                raise ValueError(f"Extracted size mismatch for {info.filename}")
            entries.append(
                {
                    "path": info.filename,
                    "sha256": digest.hexdigest(),
                    "size": extracted_size,
                    "compress_type": info.compress_type,
                    "date_time": list(info.date_time),
                    "external_attr": info.external_attr,
                }
            )

    profile["package"] = {
        "format": "Open Packaging Conventions",
        "directory": "package",
        "entry_count": len(entries),
        "entries": entries,
        "inspection": inspection,
    }
    profile_sha256 = _semantic_profile_hash(profile)
    profile["patch_contract"] = {
        "schema_version": SCHEMA_VERSION,
        "source_sha256": profile["source"]["sha256"],
        "profile_sha256": profile_sha256,
    }
    (bundle / "profile.json").write_text(
        json.dumps(profile, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return profile


def _profile_index(profile: dict[str, Any]) -> dict[str, dict[str, Any]]:
    index: dict[str, dict[str, Any]] = {}

    def visit(elements: Iterable[dict[str, Any]]) -> None:
        for element in elements:
            index[element["id"]] = element
            if element["type"] == "table":
                for row in element["rows"]:
                    for cell in row["cells"]:
                        index[cell["id"]] = cell
                        visit(cell["elements"])

    visit(profile.get("body", []))
    for container in profile.get("headers", []) + profile.get("footers", []):
        visit(container["elements"])

    # Native Word paragraph IDs survive ordinary insertions and re-profiling.
    # Only expose an alias when it is unique within this profile.
    stable_entries: dict[str, list[dict[str, Any]]] = {}
    for entry in list(index.values()):
        stable_id = entry.get("stable_id")
        if stable_id:
            stable_entries.setdefault(stable_id, []).append(entry)
    for stable_id, entries in stable_entries.items():
        if len(entries) == 1:
            index[stable_id] = entries[0]
    return index


def _xml_text_nodes(element: Any) -> list[Any]:
    return list(element.iter(qn("w:t")))


def _element_unsupported_structures(element: Any) -> list[str]:
    return sorted(
        {
            name
            for node in element.iter()
            if (name := UNSAFE_EDIT_TAGS.get(node.tag)) is not None
        }
    )


def _element_has_revisions(element: Any) -> bool:
    revision_tags = {
        qn("w:ins"),
        qn("w:del"),
        qn("w:moveFrom"),
        qn("w:moveTo"),
        qn("w:pPrChange"),
        qn("w:rPrChange"),
        qn("w:tblPrChange"),
    }
    if any(node.tag in revision_tags for node in element.iter()):
        return True
    ancestor = element.getparent()
    while ancestor is not None:
        if ancestor.tag in revision_tags:
            return True
        ancestor = ancestor.getparent()
    return False


def _replace_xml_text(element: Any, old: str, new: str) -> int:
    count = 0
    while old:
        nodes = _xml_text_nodes(element)
        combined = "".join(node.text or "" for node in nodes)
        start = combined.find(old)
        if start < 0:
            break
        end = start + len(old)
        positions = [
            (node_index, character_index)
            for node_index, node in enumerate(nodes)
            for character_index in range(len(node.text or ""))
        ]
        first_node, first_character = positions[start]
        last_node, last_character = positions[end - 1]
        if first_node == last_node:
            value = nodes[first_node].text or ""
            nodes[first_node].text = (
                value[:first_character] + new + value[last_character + 1 :]
            )
        else:
            first_value = nodes[first_node].text or ""
            last_value = nodes[last_node].text or ""
            nodes[first_node].text = first_value[:first_character] + new
            for node_index in range(first_node + 1, last_node):
                nodes[node_index].text = ""
            nodes[last_node].text = last_value[last_character + 1 :]
        count += 1
    return count


def _set_node_text(node: Any, value: str) -> None:
    node.text = value
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    if value.startswith(" ") or value.endswith(" "):
        node.set(xml_space, "preserve")
    else:
        node.attrib.pop(xml_space, None)


def _set_xml_text(element: Any, value: str) -> None:
    nodes = _xml_text_nodes(element)
    if nodes:
        _set_node_text(nodes[0], value)
        for node in nodes[1:]:
            _set_node_text(node, "")
        return
    run = etree.SubElement(element, qn("w:r"))
    _set_node_text(etree.SubElement(run, qn("w:t")), value)


def _replace_text_only(element: Any, value: str) -> None:
    """Distribute replacement across existing text nodes without changing run XML."""
    nodes = _xml_text_nodes(element)
    if not nodes:
        _set_xml_text(element, value)
        return
    offset = 0
    for index, node in enumerate(nodes):
        if index == len(nodes) - 1:
            chunk = value[offset:]
        else:
            original_length = len(node.text or "")
            chunk = value[offset : offset + original_length]
            offset += len(chunk)
        _set_node_text(node, chunk)


def _common_run_properties(element: Any) -> list[Any]:
    runs: list[Any] = []
    for node in _xml_text_nodes(element):
        if not (node.text or ""):
            continue
        run = node.getparent()
        if run is not None and run.tag == qn("w:r") and run not in runs:
            runs.append(run)
    if not runs:
        return []
    property_maps = []
    for run in runs:
        properties = run.find(qn("w:rPr"))
        property_maps.append(
            {(child.tag, etree.tostring(child)): child for child in properties}
            if properties is not None
            else {}
        )
    common_keys = set(property_maps[0])
    for property_map in property_maps[1:]:
        common_keys &= set(property_map)
    return [
        property_maps[0][key] for key in sorted(common_keys, key=lambda item: item[0])
    ]


def _replace_preserving_common_formatting(element: Any, value: str) -> None:
    common_properties = _common_run_properties(element)
    nodes = _xml_text_nodes(element)
    if not nodes:
        _set_xml_text(element, value)
        return
    target_run = nodes[0].getparent()
    if target_run is None or target_run.tag != qn("w:r"):
        raise ValueError("Cannot preserve common formatting outside a normal Word run")
    existing_properties = target_run.find(qn("w:rPr"))
    if existing_properties is not None:
        target_run.remove(existing_properties)
    if common_properties:
        properties = etree.Element(qn("w:rPr"))
        for child in common_properties:
            properties.append(_parse_xml_bytes(etree.tostring(child), "run properties"))
        target_run.insert(0, properties)
    _set_xml_text(element, value)


def _apply_paragraph_replacement(element: Any, value: str, policy: str) -> None:
    if policy == "preserve_first_run" or policy == "reject_mixed_runs":
        _set_xml_text(element, value)
    elif policy == "preserve_common_formatting":
        _replace_preserving_common_formatting(element, value)
    elif policy == "replace_text_only":
        _replace_text_only(element, value)
    else:
        raise ValueError(f"Unsupported format policy: {policy!r}")


def _apply_run_replacement(paragraph: Any, specification: dict[str, Any]) -> None:
    run_index = specification["index"]
    runs = paragraph.findall(qn("w:r"))
    if run_index >= len(runs):
        raise ValueError(f"Run index {run_index} no longer exists")
    run = runs[run_index]
    current = "".join(node.text or "" for node in _xml_text_nodes(run))
    expected = specification.get("expected")
    if expected is not None and current != expected:
        raise ValueError(f"Expected run text conflict at index {run_index}")
    _set_xml_text(run, specification["value"])


def _new_paragraph(specification: dict[str, Any]) -> Any:
    paragraph = etree.Element(qn("w:p"))
    style = specification.get("style")
    if style:
        properties = etree.SubElement(paragraph, qn("w:pPr"))
        style_element = etree.SubElement(properties, qn("w:pStyle"))
        style_element.set(qn("w:val"), str(style))
    run = etree.SubElement(paragraph, qn("w:r"))
    etree.SubElement(run, qn("w:t")).text = str(specification.get("text", ""))
    return paragraph


def _append_cell_content(cell: Any, content: Any) -> None:
    if isinstance(content, dict) and "elements" in content:
        for element in content["elements"]:
            cell.append(_new_element(element))
    else:
        text = content.get("text", "") if isinstance(content, dict) else content
        cell.append(_new_paragraph({"text": "" if text is None else str(text)}))
    if len(cell) == 0 or cell[-1].tag != qn("w:p"):
        cell.append(_new_paragraph({"text": ""}))


def _new_table(specification: dict[str, Any]) -> Any:
    table = etree.Element(qn("w:tbl"))
    properties = etree.SubElement(table, qn("w:tblPr"))
    style = specification.get("style")
    if style:
        style_element = etree.SubElement(properties, qn("w:tblStyle"))
        style_element.set(qn("w:val"), str(style))
    rows = specification.get("rows", [])
    column_count = max((len(row) for row in rows), default=0)
    grid = etree.SubElement(table, qn("w:tblGrid"))
    for _ in range(column_count):
        etree.SubElement(grid, qn("w:gridCol"))
    for row_specification in rows:
        row = etree.SubElement(table, qn("w:tr"))
        for cell_specification in row_specification:
            cell = etree.SubElement(row, qn("w:tc"))
            etree.SubElement(cell, qn("w:tcPr"))
            _append_cell_content(cell, cell_specification)
    return table


def _new_element(specification: dict[str, Any]) -> Any:
    element_type = specification.get("type")
    if element_type == "paragraph":
        return _new_paragraph(specification)
    if element_type == "table":
        return _new_table(specification)
    raise ValueError(f"Unsupported inserted element type: {element_type!r}")


def _replacement_value(value: Any) -> tuple[str, str | None, str]:
    if value is None or isinstance(value, str):
        return ("" if value is None else value), None, "preserve_first_run"
    if not isinstance(value, dict):
        raise TypeError("Element replacements must be strings, null, or objects")
    unknown = set(value) - {"value", "expected", "format_policy"}
    if unknown:
        raise ValueError(f"Unknown replacement fields: {', '.join(sorted(unknown))}")
    replacement = value.get("value")
    expected = value.get("expected")
    policy = value.get("format_policy", "preserve_first_run")
    policy = FORMAT_POLICY_ALIASES.get(policy, policy)
    if replacement is not None and not isinstance(replacement, str):
        raise TypeError("Replacement value must be a string or null")
    if expected is not None and not isinstance(expected, str):
        raise TypeError("Expected value must be a string or null")
    if policy not in FORMAT_POLICIES:
        raise ValueError(f"Unsupported format policy: {policy!r}")
    return ("" if replacement is None else replacement), expected, policy


def _validate_element_specification(specification: Any, depth: int = 0) -> None:
    if not isinstance(specification, dict):
        raise TypeError("Inserted elements must be JSON objects")
    if depth > 12:
        raise ValueError("Inserted element nesting exceeds 12 levels")
    element_type = specification.get("type")
    if element_type == "paragraph":
        if not isinstance(specification.get("text", ""), str):
            raise TypeError("Paragraph text must be a string")
        if specification.get("style") is not None and not isinstance(
            specification["style"], str
        ):
            raise TypeError("Paragraph style must be a string")
        return
    if element_type != "table":
        raise ValueError(f"Unsupported inserted element type: {element_type!r}")
    rows = specification.get("rows")
    if not isinstance(rows, list) or not rows:
        raise ValueError("Inserted tables require at least one row")
    width = None
    for row in rows:
        if not isinstance(row, list) or not row:
            raise ValueError("Every table row must be a non-empty array")
        if width is None:
            width = len(row)
        elif len(row) != width:
            raise ValueError("All generated table rows must have the same width")
        for cell in row:
            if isinstance(cell, dict) and "elements" in cell:
                if not isinstance(cell["elements"], list):
                    raise TypeError("Cell elements must be an array")
                for child in cell["elements"]:
                    _validate_element_specification(child, depth + 1)
            elif not isinstance(cell, (str, int, float, bool, type(None), dict)):
                raise TypeError("Unsupported table cell value")


def _validate_repeat_table_rows(
    operation: dict[str, Any], target: dict[str, Any]
) -> None:
    if target.get("type") != "table":
        raise ValueError("repeat_table_rows requires a table target")
    template_row = operation.get("template_row")
    rows = operation.get("rows")
    target_rows = target.get("rows", [])
    if (
        not isinstance(template_row, int)
        or isinstance(template_row, bool)
        or template_row < 0
        or template_row >= len(target_rows)
    ):
        raise ValueError("repeat_table_rows template_row does not exist")
    if not isinstance(rows, list):
        raise TypeError("repeat_table_rows rows must be an array")
    if len(rows) > 10_000:
        raise ValueError("repeat_table_rows exceeds the 10,000 row limit")
    width = target.get("column_count", 0)
    for row in rows:
        if not isinstance(row, list) or len(row) != width:
            raise ValueError(
                f"Every repeated row must contain exactly {width} string values"
            )
        if any(not isinstance(value, str) for value in row):
            raise TypeError("Repeated table cell values must be strings")
    template = target_rows[template_row]
    cells = template.get("cells", [])
    if any(cell.get("duplicate_merged_cell") for cell in cells):
        raise ValueError("Merged template rows are unsupported for repeat_table_rows")
    for cell in cells:
        elements = cell.get("elements", [])
        if any(element.get("type") == "table" for element in elements):
            raise ValueError("Nested tables are unsupported in repeated template rows")
        paragraphs = [
            element for element in elements if element.get("type") == "paragraph"
        ]
        if len(paragraphs) != 1:
            raise ValueError(
                "Each repeated template-row cell must contain exactly one paragraph"
            )
        paragraph = paragraphs[0]
        if paragraph.get("has_revisions") or paragraph.get(
            "unsupported_edit_structures"
        ):
            raise ValueError(
                "Repeated template rows cannot contain unsafe editable structures"
            )


def validate_patch(
    profile: dict[str, Any], data: dict[str, Any]
) -> list[dict[str, str]]:
    """Validate an AI patch before any package part is modified."""
    if not isinstance(data, dict):
        raise TypeError("Patch root must be an object")
    unknown = set(data) - {
        "contract",
        "tokens",
        "elements",
        "runs",
        "operations",
        "options",
    }
    if unknown:
        raise ValueError(f"Unknown patch fields: {', '.join(sorted(unknown))}")
    stored_contract = profile.get("patch_contract")
    if not isinstance(stored_contract, dict):
        raise ValueError("Bundle profile has no valid patch contract")
    if stored_contract.get("source_sha256") != profile.get("source", {}).get("sha256"):
        raise ValueError("Bundle source hash and patch contract disagree")
    if stored_contract.get("profile_sha256") != _semantic_profile_hash(profile):
        raise ValueError("Semantic profile is stale or has been modified")
    contract = data.get("contract")
    if contract != stored_contract:
        raise ValueError("Patch contract does not match this bundle profile")
    tokens = data.get("tokens", {})
    elements = data.get("elements", {})
    runs = data.get("runs", {})
    operations = data.get("operations", [])
    if (
        not isinstance(tokens, dict)
        or not isinstance(elements, dict)
        or not isinstance(runs, dict)
        or not isinstance(operations, list)
    ):
        raise TypeError(
            "tokens, elements, and runs must be objects; operations must be an array"
        )
    for token, value in tokens.items():
        if not isinstance(token, str) or not TOKEN_PATTERN.fullmatch(
            "{{" + token + "}}"
        ):
            raise ValueError(f"Invalid token name: {token!r}")
        if value is not None and not isinstance(value, str):
            raise TypeError(f"Token {token!r} must have a string or null value")
    index = _profile_index(profile)
    warnings: list[dict[str, str]] = []
    for element_id, value in elements.items():
        entry = index.get(element_id)
        if entry is None:
            raise KeyError(f"Unknown element ID: {element_id}")
        if entry.get("type") != "paragraph":
            raise ValueError(f"Text replacement requires a paragraph ID: {element_id}")
        replacement, expected, policy = _replacement_value(value)
        if expected is not None and expected != entry.get("text", ""):
            raise ValueError(f"Expected text conflict for {element_id}")
        if entry.get("has_revisions"):
            raise ValueError(f"Tracked revisions prevent safe editing of {element_id}")
        if entry.get("unsupported_edit_structures"):
            structures = ", ".join(entry["unsupported_edit_structures"])
            raise ValueError(
                f"Unsupported OOXML structures prevent full replacement of {element_id}: {structures}"
            )
        if entry.get("has_mixed_run_formatting") and replacement != entry.get("text"):
            if policy == "reject_mixed_runs":
                raise ValueError(f"Mixed run formatting rejected for {element_id}")
            warning_messages = {
                "preserve_first_run": "Replacement uses the first text run's formatting.",
                "preserve_common_formatting": "Only formatting common to all text runs is retained on replacement text.",
                "replace_text_only": "Existing run boundaries are retained; changed text length may shift semantic emphasis.",
            }
            warnings.append(
                {
                    "code": "MIXED_RUN_FORMATTING",
                    "element": element_id,
                    "message": warning_messages[policy],
                }
            )

    for paragraph_id, replacements in runs.items():
        entry = index.get(paragraph_id)
        if entry is None or entry.get("type") != "paragraph":
            raise KeyError(f"Unknown run-replacement paragraph ID: {paragraph_id}")
        if paragraph_id in elements:
            raise ValueError(
                f"Paragraph {paragraph_id} cannot have full and run-level replacements"
            )
        if entry.get("has_revisions") or entry.get("unsupported_edit_structures"):
            raise ValueError(
                f"Paragraph {paragraph_id} cannot be safely edited at run level"
            )
        if not isinstance(replacements, list) or not replacements:
            raise ValueError("Run replacements must be non-empty arrays")
        seen_run_indexes: set[int] = set()
        profile_runs = entry.get("runs", [])
        for replacement in replacements:
            if not isinstance(replacement, dict):
                raise TypeError("Run replacement must be an object")
            if set(replacement) - {"index", "expected", "value"}:
                raise ValueError("Run replacement contains unknown fields")
            run_index = replacement.get("index")
            if (
                not isinstance(run_index, int)
                or isinstance(run_index, bool)
                or run_index < 0
            ):
                raise TypeError("Run index must be a non-negative integer")
            if run_index in seen_run_indexes:
                raise ValueError(
                    f"Duplicate run replacement index {run_index} in {paragraph_id}"
                )
            seen_run_indexes.add(run_index)
            if run_index >= len(profile_runs):
                raise ValueError(
                    f"Run index {run_index} does not exist in {paragraph_id}"
                )
            expected = replacement.get("expected")
            value = replacement.get("value")
            if expected is not None and expected != profile_runs[run_index]["text"]:
                raise ValueError(
                    f"Expected run text conflict for {paragraph_id} run {run_index}"
                )
            if not isinstance(value, str):
                raise TypeError("Run replacement value must be a string")

    operation_targets: dict[str, str] = {}
    deleted_entries: list[dict[str, Any]] = []
    repeated_entries: list[dict[str, Any]] = []
    for operation in operations:
        if not isinstance(operation, dict):
            raise TypeError("Each operation must be an object")
        name = operation.get("op")
        allowed_operations = {
            "insert_before",
            "insert_after",
            "append_to_cell",
            "delete",
            "repeat_table_rows",
        }
        if name not in allowed_operations:
            raise ValueError(f"Unsupported operation: {name!r}")
        allowed_fields = (
            {"op", "target", "template_row", "rows"}
            if name == "repeat_table_rows"
            else {"op", "target", "element"}
        )
        if set(operation) - allowed_fields:
            raise ValueError("Operation contains unknown fields")
        target_id = operation.get("target")
        if not isinstance(target_id, str):
            raise TypeError("Operation target must be a string")
        target = index.get(target_id)
        if target is None:
            raise KeyError(f"Unknown operation target: {target_id}")
        canonical_target = target["id"]
        if canonical_target in operation_targets:
            raise ValueError(
                f"Conflicting structural operations target {canonical_target}: "
                f"{operation_targets[canonical_target]} and {name}"
            )
        operation_targets[canonical_target] = str(name)
        if target.get("has_revisions"):
            raise ValueError(
                f"Tracked revisions prevent structural editing of {operation.get('target')}"
            )
        if target.get("unsupported_edit_structures"):
            raise ValueError(
                f"Unsupported OOXML structure at operation target {target_id}"
            )
        if name == "append_to_cell" and target.get("type") != "cell":
            raise ValueError("append_to_cell requires a table cell target")
        if name in {"insert_before", "insert_after", "delete"} and target.get(
            "type"
        ) not in {"paragraph", "table"}:
            raise ValueError(f"{name} requires a paragraph or table target")
        if name == "delete":
            deleted_entries.append(target)
            if canonical_target in {
                index[key]["id"] for key in elements
            } or canonical_target in {index[key]["id"] for key in runs}:
                raise ValueError(
                    f"Cannot edit and delete {canonical_target} in one patch"
                )
        elif name == "repeat_table_rows":
            _validate_repeat_table_rows(operation, target)
            repeated_entries.append(target)
        else:
            if "element" not in operation:
                raise ValueError(f"{name} requires an element")
            _validate_element_specification(operation["element"])

    edited_entries = [index[key] for key in elements] + [index[key] for key in runs]
    for repeated in repeated_entries:
        for edited in edited_entries:
            if edited["id"].startswith(repeated["id"] + "."):
                raise ValueError(
                    f"Cannot edit {edited['id']} while repeating rows in {repeated['id']}"
                )
        for operation_target, operation_name in operation_targets.items():
            if operation_target != repeated["id"] and operation_target.startswith(
                repeated["id"] + "."
            ):
                raise ValueError(
                    f"Cannot apply {operation_name} inside repeated table {repeated['id']}"
                )

    affected_entries = edited_entries + [
        index[operation["target"]]
        for operation in operations
        if operation["op"] != "delete"
    ]
    for deleted in deleted_entries:
        deleted_locator = deleted.get("locator", {})
        deleted_path = deleted_locator.get("xpath", "")
        for affected in affected_entries:
            affected_locator = affected.get("locator", {})
            if deleted_locator.get("part") == affected_locator.get(
                "part"
            ) and affected_locator.get("xpath", "").startswith(deleted_path + "/"):
                raise ValueError(
                    f"Cannot modify {affected['id']} inside deleted element {deleted['id']}"
                )
    return warnings


def _verify_bundle_package(profile: dict[str, Any], package_dir: Path) -> None:
    manifest = profile.get("package", {}).get("entries", [])
    if not manifest:
        raise ValueError("Profile does not contain a package manifest")
    expected_paths = {entry["path"] for entry in manifest}
    actual_paths = {
        path.relative_to(package_dir).as_posix()
        for path in package_dir.rglob("*")
        if path.is_file()
    }
    if actual_paths != expected_paths:
        raise ValueError("Extracted package parts do not match the bundle manifest")
    for entry in manifest:
        content = (package_dir / entry["path"]).read_bytes()
        if _sha256_bytes(content) != entry["sha256"]:
            raise ValueError(f"Bundle package part has changed: {entry['path']}")


def _repeat_table_template_row(
    table: Any, operation: dict[str, Any]
) -> tuple[Any, list[Any]]:
    rows = [child for child in table if child.tag == qn("w:tr")]
    template_index = operation["template_row"]
    if template_index >= len(rows):
        raise ValueError("OOXML template row no longer exists")
    template = rows[template_index]
    if any(node.tag in {qn("w:gridSpan"), qn("w:vMerge")} for node in template.iter()):
        raise ValueError(
            "Merged-cell template rows are unsupported for repeat_table_rows"
        )
    cells = [child for child in template if child.tag == qn("w:tc")]
    expected_width = len(operation["rows"][0]) if operation["rows"] else len(cells)
    if len(cells) != expected_width:
        raise ValueError("Template-row OOXML cell count does not match patch rows")
    for cell in cells:
        if any(node.tag == qn("w:tbl") for node in cell.iterdescendants()):
            raise ValueError("Nested tables are unsupported in repeated template rows")
        paragraphs = [child for child in cell if child.tag == qn("w:p")]
        if len(paragraphs) != 1:
            raise ValueError(
                "Each repeated template-row cell must contain exactly one paragraph"
            )
        if _element_has_revisions(paragraphs[0]) or _element_unsupported_structures(
            paragraphs[0]
        ):
            raise ValueError(
                "Repeated template rows cannot contain unsafe editable structures"
            )
    return template, cells


def _apply_repeat_table_rows(table: Any, operation: dict[str, Any]) -> None:
    template, _ = _repeat_table_template_row(table, operation)
    for values in operation["rows"]:
        clone = copy.deepcopy(template)
        cells = [child for child in clone if child.tag == qn("w:tc")]
        for cell, value in zip(cells, values):
            paragraph = next(child for child in cell if child.tag == qn("w:p"))
            _apply_paragraph_replacement(paragraph, value, "preserve_first_run")
        template.addprevious(clone)
    template.getparent().remove(template)


def _validate_repeat_table_in_bundle(
    package_dir: Path, entry: dict[str, Any], operation: dict[str, Any]
) -> None:
    locator = entry.get("locator", {})
    parser = _secure_xml_parser(remove_blank_text=False)
    tree = etree.parse(str(package_dir / locator["part"]), parser)
    namespaces = {key: value for key, value in tree.getroot().nsmap.items() if key}
    matches = tree.xpath(locator["xpath"], namespaces=namespaces)
    if len(matches) != 1:
        raise ValueError(
            f"Locator for {entry['id']} matched {len(matches)} OOXML nodes"
        )
    _repeat_table_template_row(matches[0], operation)


def preview_bundle(bundle: str | Path, data: dict[str, Any]) -> dict[str, Any]:
    bundle = Path(bundle)
    package_dir = bundle / "package"
    profile = _load_json(bundle / "profile.json")
    _verify_bundle_package(profile, package_dir)
    warnings = validate_patch(profile, data)
    index = _profile_index(profile)
    changes = []
    for element_id, value in data.get("elements", {}).items():
        replacement, _, policy = _replacement_value(value)
        changes.append(
            {
                "kind": "replace_text",
                "element": element_id,
                "before": index[element_id]["text"],
                "after": replacement,
                "format_policy": policy,
            }
        )
    for paragraph_id, replacements in data.get("runs", {}).items():
        for replacement in replacements:
            changes.append(
                {
                    "kind": "replace_run_text",
                    "element": paragraph_id,
                    "run_index": replacement["index"],
                    "before": index[paragraph_id]["runs"][replacement["index"]]["text"],
                    "after": replacement["value"],
                }
            )
    for operation in data.get("operations", []):
        if operation["op"] == "repeat_table_rows":
            target = index[operation["target"]]
            _validate_repeat_table_in_bundle(package_dir, target, operation)
            template = target["rows"][operation["template_row"]]
            changes.append(
                {
                    "kind": operation["op"],
                    "target": operation["target"],
                    "template_row": operation["template_row"],
                    "before": [cell.get("text", "") for cell in template["cells"]],
                    "after": operation["rows"],
                }
            )
        else:
            changes.append(
                {
                    "kind": operation["op"],
                    "target": operation["target"],
                    "element_type": operation.get("element", {}).get("type"),
                }
            )
    return {
        "valid": True,
        "source_sha256": profile["source"]["sha256"],
        "token_keys": sorted(data.get("tokens", {})),
        "changes": changes,
        "warnings": warnings,
    }


def validate_docx_package(
    path: str | Path,
    limits: PackageLimits = DEFAULT_PACKAGE_LIMITS,
) -> dict[str, Any]:
    """Validate bounded OPC structure, XML, content types, and relationships."""
    inspection = inspect_docx_archive(path, limits)
    required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        names = [info.filename for info in infos if not info.is_dir()]
        if len(names) != len(set(names)):
            raise ValueError("DOCX contains duplicate package entry names")
        name_set = set(names)
        missing = required - name_set
        if missing:
            raise ValueError(
                f"DOCX is missing required parts: {', '.join(sorted(missing))}"
            )
        parsed: dict[str, Any] = {}
        for name in names:
            if name.endswith((".xml", ".rels")):
                parsed[name] = _parse_xml_bytes(archive.read(name), name)

        content_types = parsed["[Content_Types].xml"]
        defaults = {
            node.get("Extension", "").lower()
            for node in content_types
            if node.tag.endswith("Default")
        }
        overrides = {
            node.get("PartName", "").lstrip("/")
            for node in content_types
            if node.tag.endswith("Override")
        }
        for name in names:
            if name in {"[Content_Types].xml"} or name.endswith(".rels"):
                continue
            extension = PurePosixPath(name).suffix.lstrip(".").lower()
            if name not in overrides and extension not in defaults:
                raise ValueError(f"No content type declared for package part: {name}")

        relationship_count = 0
        for rels_name, root in parsed.items():
            if not rels_name.endswith(".rels"):
                continue
            ids: set[str] = set()
            base = (
                "" if rels_name == "_rels/.rels" else rels_name.split("/_rels/", 1)[0]
            )
            for relationship in root:
                relationship_id = relationship.get("Id")
                if not relationship_id or relationship_id in ids:
                    raise ValueError(
                        f"Missing or duplicate relationship ID in {rels_name}"
                    )
                ids.add(relationship_id)
                relationship_count += 1
                if relationship.get("TargetMode") == "External":
                    continue
                target = relationship.get("Target")
                if not target:
                    raise ValueError(f"Relationship {relationship_id} has no target")
                resolved = posixpath.normpath(posixpath.join(base, target.lstrip("/")))
                if resolved.startswith("../") or resolved not in name_set:
                    raise ValueError(
                        f"Broken relationship {relationship_id} in {rels_name}: {target}"
                    )
    # python-docx provides an additional practical openability check.
    Document(str(path))
    return {
        "valid": True,
        "entry_count": len(names),
        "xml_part_count": len(parsed),
        "relationship_count": relationship_count,
        "total_uncompressed_size": inspection["total_uncompressed_size"],
        "highest_compression_ratio": inspection["highest_compression_ratio"],
        "limits": inspection["limits"],
    }


def _part_hashes(path: str | Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {
            info.filename: _sha256_bytes(archive.read(info))
            for info in archive.infolist()
            if not info.is_dir()
        }


def _expected_modified_parts(
    profile: dict[str, Any], patch: dict[str, Any] | None
) -> set[str]:
    if not patch:
        return set()
    index = _profile_index(profile)
    parts: set[str] = set()
    for target_id in list(patch.get("elements", {})) + list(patch.get("runs", {})):
        entry = index.get(target_id)
        if entry and entry.get("locator"):
            parts.add(entry["locator"]["part"])
    for operation in patch.get("operations", []):
        entry = index.get(operation.get("target"))
        if entry and entry.get("locator"):
            parts.add(entry["locator"]["part"])
    token_keys = set(patch.get("tokens", {}))
    if token_keys:
        for entry in index.values():
            if (
                entry.get("type") == "paragraph"
                and token_keys.intersection(entry.get("tokens", []))
                and entry.get("locator")
            ):
                parts.add(entry["locator"]["part"])
    return parts


def _profile_paragraph_records(profile: dict[str, Any]) -> list[dict[str, Any]]:
    elements = list(profile.get("body", []))
    for container in profile.get("headers", []) + profile.get("footers", []):
        elements.extend(container["elements"])
    return list(_collect_paragraphs(elements))


def _optional_part_hash(hashes: dict[str, str], name: str) -> str | None:
    return hashes.get(name)


def _inserted_texts(specification: dict[str, Any]) -> list[str]:
    if specification.get("type") == "paragraph":
        return [str(specification.get("text", ""))]
    texts: list[str] = []
    for row in specification.get("rows", []):
        for cell in row:
            if isinstance(cell, dict) and "elements" in cell:
                for element in cell["elements"]:
                    texts.extend(_inserted_texts(element))
            else:
                texts.append(
                    str(
                        cell.get("text", "")
                        if isinstance(cell, dict)
                        else ("" if cell is None else cell)
                    )
                )
    return texts


def _entry_paragraph_texts(entry: dict[str, Any]) -> list[str]:
    if entry.get("type") == "paragraph":
        return [entry.get("text", "")]
    if entry.get("type") == "table":
        return [
            text
            for row in entry.get("rows", [])
            for cell in row.get("cells", [])
            for element in cell.get("elements", [])
            for text in _entry_paragraph_texts(element)
        ]
    if entry.get("type") == "cell":
        return [
            text
            for element in entry.get("elements", [])
            for text in _entry_paragraph_texts(element)
        ]
    return []


def _expected_semantic_texts(
    profile: dict[str, Any], patch: dict[str, Any] | None
) -> list[str]:
    records = _profile_paragraph_records(profile)
    if not patch:
        return [record["text"] for record in records]
    index = _profile_index(profile)
    text_by_id = {record["id"]: record["text"] for record in records}
    for element_id, value in patch.get("elements", {}).items():
        replacement, _, _ = _replacement_value(value)
        text_by_id[index[element_id]["id"]] = replacement
    for paragraph_id, replacements in patch.get("runs", {}).items():
        entry = index[paragraph_id]
        run_texts = [run["text"] for run in entry.get("runs", [])]
        for replacement in replacements:
            run_texts[replacement["index"]] = replacement["value"]
        text_by_id[entry["id"]] = "".join(run_texts)
    for record in records:
        canonical_id = record["id"]
        text = text_by_id[canonical_id]
        for token, value in patch.get("tokens", {}).items():
            text = re.sub(
                r"\{\{\s*" + re.escape(token) + r"\s*\}\}",
                "" if value is None else value,
                text,
            )
        text_by_id[canonical_id] = text

    additions: list[str] = []
    for operation in patch.get("operations", []):
        target = index[operation["target"]]
        if operation["op"] == "delete":
            removed = Counter(_entry_paragraph_texts(target))
            for record_id, text in list(text_by_id.items()):
                if removed[text]:
                    del text_by_id[record_id]
                    removed[text] -= 1
        elif operation["op"] == "repeat_table_rows":
            template = target["rows"][operation["template_row"]]
            removed = Counter(
                text
                for cell in template["cells"]
                for text in _entry_paragraph_texts(cell)
            )
            for record_id, text in list(text_by_id.items()):
                if removed[text]:
                    del text_by_id[record_id]
                    removed[text] -= 1
            additions.extend(value for row in operation["rows"] for value in row)
        else:
            additions.extend(_inserted_texts(operation["element"]))
    return list(text_by_id.values()) + additions


def create_fidelity_report(
    original: str | Path,
    generated: str | Path,
    *,
    bundle: str | Path | None = None,
    patch: dict[str, Any] | None = None,
    visual_report: dict[str, Any] | str | Path | None = None,
    limits: PackageLimits = DEFAULT_PACKAGE_LIMITS,
) -> dict[str, Any]:
    """Create a machine-readable report without conflating integrity and visuals."""
    original = Path(original)
    generated = Path(generated)
    original_integrity = validate_docx_package(original, limits)
    generated_integrity = validate_docx_package(generated, limits)
    original_hashes = _part_hashes(original)
    generated_hashes = _part_hashes(generated)
    original_profile = profile_document(original, limits)
    generated_profile = profile_document(generated, limits)

    bundle_profile = None
    if bundle is not None:
        bundle_path = Path(bundle)
        bundle_profile = _load_json(bundle_path / "profile.json")
        _verify_bundle_package(bundle_profile, bundle_path / "package")
        if bundle_profile["source"]["sha256"] != _sha256_bytes(original.read_bytes()):
            raise ValueError("Fidelity bundle does not belong to the original DOCX")
        if patch is not None:
            validate_patch(bundle_profile, patch)
    expected_modified_parts = _expected_modified_parts(
        bundle_profile or original_profile, patch
    )
    original_parts = set(original_hashes)
    generated_parts = set(generated_hashes)
    changed_parts = sorted(
        name
        for name in original_parts & generated_parts
        if original_hashes[name] != generated_hashes[name]
    )
    missing_parts = sorted(original_parts - generated_parts)
    added_parts = sorted(generated_parts - original_parts)
    unexpected_changed_parts = sorted(set(changed_parts) - expected_modified_parts)
    untouched_parts = sorted(original_parts - expected_modified_parts)
    untouched_preserved = sorted(
        name
        for name in untouched_parts
        if generated_hashes.get(name) == original_hashes[name]
    )

    original_records = _profile_paragraph_records(original_profile)
    generated_records = _profile_paragraph_records(generated_profile)
    original_text = [record["text"] for record in original_records]
    generated_text = [record["text"] for record in generated_records]
    expected_text = _expected_semantic_texts(bundle_profile or original_profile, patch)
    semantic_exact = original_text == generated_text
    semantic_expected = Counter(expected_text) == Counter(generated_text)
    unexpected_removed_text = list(
        (Counter(expected_text) - Counter(generated_text)).elements()
    )
    unexpected_added_text = list(
        (Counter(generated_text) - Counter(expected_text)).elements()
    )
    original_text_digest = _sha256_bytes("\n\0".join(original_text).encode())
    generated_text_digest = _sha256_bytes("\n\0".join(generated_text).encode())
    expected_text_digest = _sha256_bytes("\n\0".join(expected_text).encode())

    style_hashes: dict[str, dict[str, Any]] = {
        "styles": {
            "original": _optional_part_hash(original_hashes, "word/styles.xml"),
            "generated": _optional_part_hash(generated_hashes, "word/styles.xml"),
        },
        "numbering": {
            "original": _optional_part_hash(original_hashes, "word/numbering.xml"),
            "generated": _optional_part_hash(generated_hashes, "word/numbering.xml"),
        },
    }
    for value in style_hashes.values():
        value["preserved"] = value["original"] == value["generated"]

    unsupported = sorted(
        {
            structure
            for record in original_records
            for structure in record.get("unsupported_edit_structures", [])
        }
    )
    if isinstance(visual_report, (str, Path)):
        visual = _load_json(visual_report)
    else:
        visual = visual_report
    visual_summary = {
        "performed": visual is not None,
        "passed": visual.get("passed") if visual else None,
        "mode": visual.get("mode") if visual else None,
        "page_count_match": visual.get("page_count_match") if visual else None,
        "unexpected_changed_pages": visual.get("unexpected_changed_pages", [])
        if visual
        else [],
        "report": visual,
    }
    no_op = not patch or not any(
        patch.get(key) for key in ("tokens", "elements", "runs", "operations")
    )
    semantic_passed = semantic_exact if no_op else semantic_expected
    package_passed = (
        not missing_parts
        and not added_parts
        and not unexpected_changed_parts
        and len(untouched_preserved) == len(untouched_parts)
    )
    styles_passed = all(value["preserved"] for value in style_hashes.values())
    visual_passed = visual_summary["passed"] is not False
    return {
        "schema_version": "1.0",
        "mode": "no-op" if no_op else "patch",
        "passed": package_passed
        and semantic_passed
        and styles_passed
        and visual_passed,
        "package_integrity": {
            "passed": original_integrity["valid"] and generated_integrity["valid"],
            "original": original_integrity,
            "generated": generated_integrity,
        },
        "untouched_part_hash_preservation": {
            "passed": package_passed,
            "expected_modified_parts": sorted(expected_modified_parts),
            "changed_parts": changed_parts,
            "unexpected_changed_parts": unexpected_changed_parts,
            "missing_parts": missing_parts,
            "added_parts": added_parts,
            "untouched_count": len(untouched_parts),
            "preserved_count": len(untouched_preserved),
        },
        "semantic_text_preservation": {
            "passed": semantic_passed,
            "exact_match": semantic_exact,
            "matches_expected_patch": semantic_expected,
            "original_paragraph_count": len(original_records),
            "generated_paragraph_count": len(generated_records),
            "original_text_sha256": original_text_digest,
            "generated_text_sha256": generated_text_digest,
            "expected_text_sha256": expected_text_digest,
            "unexpected_removed_text": unexpected_removed_text,
            "unexpected_added_text": unexpected_added_text,
            "intentional_changes_allowed": not no_op,
        },
        "style_and_numbering_preservation": {
            "passed": styles_passed,
            **style_hashes,
        },
        "visual_differences": visual_summary,
        "unsupported_ooxml_structures": unsupported,
    }


def write_fidelity_report(
    output: str | Path,
    *args: Any,
    **kwargs: Any,
) -> dict[str, Any]:
    report = create_fidelity_report(*args, **kwargs)
    Path(output).write_text(json.dumps(report, indent=2), encoding="utf-8")
    return report


def render_bundle(
    bundle: str | Path, output: str | Path, data: dict[str, Any]
) -> dict[str, Any]:
    """Patch extracted OOXML parts and rebuild a functionally equivalent DOCX."""
    bundle = Path(bundle)
    package_dir = bundle / "package"
    profile = _load_json(bundle / "profile.json")
    _verify_bundle_package(profile, package_dir)
    warnings = validate_patch(profile, data)
    index = _profile_index(profile)
    modified: dict[str, tuple[Any, Any]] = {}

    def locate(element_id: str) -> Any:
        if element_id not in index:
            raise KeyError(f"Unknown element ID: {element_id}")
        locator = index[element_id].get("locator")
        if not locator:
            raise ValueError(f"Element has no OOXML locator: {element_id}")
        part = locator["part"]
        if part not in modified:
            parser = _secure_xml_parser(remove_blank_text=False)
            tree = etree.parse(str(package_dir / part), parser)
            namespaces = {
                key: value for key, value in tree.getroot().nsmap.items() if key
            }
            modified[part] = (tree, namespaces)
        tree, namespaces = modified[part]
        matches = tree.xpath(locator["xpath"], namespaces=namespaces)
        if len(matches) != 1:
            raise ValueError(
                f"Locator for {element_id} matched {len(matches)} nodes in {part}"
            )
        return matches[0]

    changed_elements = 0
    for element_id, value in data.get("elements", {}).items():
        replacement, expected, policy = _replacement_value(value)
        target = locate(element_id)
        if _element_has_revisions(target):
            raise ValueError(f"Tracked revisions prevent safe editing of {element_id}")
        unsupported = _element_unsupported_structures(target)
        if unsupported:
            raise ValueError(
                f"Unsupported OOXML structures prevent replacement of {element_id}: "
                + ", ".join(unsupported)
            )
        current_text = "".join(node.text or "" for node in _xml_text_nodes(target))
        if expected is not None and current_text != expected:
            raise ValueError(f"OOXML text conflict for {element_id}")
        _apply_paragraph_replacement(target, replacement, policy)
        changed_elements += 1

    changed_runs = 0
    for paragraph_id, replacements in data.get("runs", {}).items():
        target = locate(paragraph_id)
        if _element_has_revisions(target) or _element_unsupported_structures(target):
            raise ValueError(
                f"Paragraph {paragraph_id} cannot be safely edited at run level"
            )
        for replacement in replacements:
            _apply_run_replacement(target, replacement)
            changed_runs += 1

    token_replacements = 0
    token_values = data.get("tokens", {})
    paragraph_entries = {
        entry["locator"]["part"] + "|" + entry["locator"]["xpath"]: entry
        for entry in index.values()
        if entry.get("type") == "paragraph" and entry.get("locator")
    }
    for entry in paragraph_entries.values() if token_values else ():
        paragraph = locate(entry["id"])
        for token, value in token_values.items():
            pattern = re.compile(r"\{\{\s*" + re.escape(str(token)) + r"\s*\}\}")
            while True:
                text = "".join(node.text or "" for node in _xml_text_nodes(paragraph))
                match = pattern.search(text)
                if not match:
                    break
                if _element_has_revisions(paragraph):
                    raise ValueError(
                        f"Tracked revisions prevent token editing of {entry['id']}"
                    )
                token_replacements += _replace_xml_text(
                    paragraph, match.group(0), "" if value is None else str(value)
                )

    # Resolve all original structural targets before insertion changes positional XPath indexes.
    resolved_operations = [
        (operation, locate(operation.get("target", "")))
        for operation in data.get("operations", [])
    ]
    structural_changes = 0
    for operation, target in resolved_operations:
        operation_name = operation.get("op")
        if operation_name == "repeat_table_rows":
            _apply_repeat_table_rows(target, operation)
            structural_changes += 1
            continue
        if _element_has_revisions(target):
            raise ValueError(
                f"Tracked revisions prevent structural editing of {operation.get('target')}"
            )
        unsupported = _element_unsupported_structures(target)
        if unsupported:
            raise ValueError(
                f"Unsupported OOXML structures prevent structural editing of "
                f"{operation.get('target')}: {', '.join(unsupported)}"
            )
        if operation_name == "delete":
            target.getparent().remove(target)
        elif operation_name in ("insert_before", "insert_after"):
            new_element = _new_element(operation["element"])
            if operation_name == "insert_before":
                target.addprevious(new_element)
            else:
                target.addnext(new_element)
        elif operation_name == "append_to_cell":
            if target.tag != qn("w:tc"):
                raise ValueError("append_to_cell requires a table cell target")
            # Keep Word's required final paragraph after any nested table.
            final_paragraph = (
                target[-1] if len(target) and target[-1].tag == qn("w:p") else None
            )
            new_element = _new_element(operation["element"])
            if final_paragraph is not None and new_element.tag == qn("w:tbl"):
                final_paragraph.addprevious(new_element)
            else:
                target.append(new_element)
        else:
            raise ValueError(f"Unsupported operation: {operation_name!r}")
        structural_changes += 1

    manifest = profile.get("package", {}).get("entries", [])
    if not manifest:
        raise ValueError("Profile does not contain a package manifest")
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.name}.", suffix=".tmp", dir=output.parent
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with zipfile.ZipFile(temporary_path, "w") as archive:
            for entry in manifest:
                path = entry["path"]
                if path in modified:
                    tree, _ = modified[path]
                    content = etree.tostring(
                        tree, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                else:
                    content = (package_dir / path).read_bytes()
                    if _sha256_bytes(content) != entry["sha256"]:
                        raise ValueError(f"Untouched package part has changed: {path}")
                info = zipfile.ZipInfo(path, tuple(entry["date_time"]))
                info.compress_type = entry["compress_type"]
                info.external_attr = entry["external_attr"]
                archive.writestr(info, content)
        integrity = validate_docx_package(temporary_path)
        os.replace(temporary_path, output)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise

    return {
        "output": str(output),
        "changed_elements": changed_elements,
        "changed_runs": changed_runs,
        "token_replacements": token_replacements,
        "structural_changes": structural_changes,
        "modified_parts": sorted(modified),
        "integrity": integrity,
        "warnings": warnings,
    }


def _load_json(path: str | Path) -> dict[str, Any]:
    with Path(path).open(encoding="utf-8") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise TypeError("JSON root must be an object")
    return value


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="DocXUnravel: losslessly profile and render agent-friendly DOCX templates"
    )
    commands = parser.add_subparsers(dest="command", required=True)

    profile_parser = commands.add_parser(
        "profile", help="Create a detailed JSON document profile"
    )
    profile_parser.add_argument("source")
    profile_parser.add_argument("output")

    extract_parser = commands.add_parser(
        "extract", help="Create a lossless OOXML package plus semantic JSON bundle"
    )
    extract_parser.add_argument("source")
    extract_parser.add_argument("bundle")
    extract_parser.add_argument(
        "--max-entries", type=int, default=DEFAULT_PACKAGE_LIMITS.max_entries
    )
    extract_parser.add_argument(
        "--max-entry-size", type=int, default=DEFAULT_PACKAGE_LIMITS.max_entry_size
    )
    extract_parser.add_argument(
        "--max-total-size",
        type=int,
        default=DEFAULT_PACKAGE_LIMITS.max_total_uncompressed_size,
    )
    extract_parser.add_argument(
        "--max-compression-ratio",
        type=float,
        default=DEFAULT_PACKAGE_LIMITS.max_compression_ratio,
    )

    preview_parser = commands.add_parser(
        "preview", help="Validate a patch and show its human-readable change plan"
    )
    preview_parser.add_argument("bundle")
    preview_parser.add_argument("data")

    bundle_parser = commands.add_parser(
        "render-bundle", help="Patch OOXML in an extracted bundle and rebuild DOCX"
    )
    bundle_parser.add_argument("bundle")
    bundle_parser.add_argument("data")
    bundle_parser.add_argument("output")

    fidelity_parser = commands.add_parser(
        "fidelity-report",
        help="Report package, semantic, style, numbering, and optional visual fidelity",
    )
    fidelity_parser.add_argument("original")
    fidelity_parser.add_argument("generated")
    fidelity_parser.add_argument("output")
    fidelity_parser.add_argument("--bundle")
    fidelity_parser.add_argument("--patch")
    fidelity_parser.add_argument("--visual-report")

    visual_parser = commands.add_parser(
        "visual", help="Optionally render two DOCX files and compare their pages"
    )
    visual_parser.add_argument("original")
    visual_parser.add_argument("generated")
    visual_parser.add_argument("output_dir")
    visual_parser.add_argument("--renderer")
    visual_parser.add_argument("--difference-threshold", type=float, default=0.001)
    visual_parser.add_argument(
        "--expected-difference-threshold", type=float, default=0.15
    )
    visual_parser.add_argument("--pixel-threshold", type=int, default=12)
    visual_parser.add_argument("--dimension-tolerance-pt", type=float, default=0.5)
    visual_parser.add_argument("--expected-pages", default="")
    visual_parser.add_argument("--no-op", action="store_true")
    visual_parser.add_argument("--dpi", type=int, default=144)
    visual_parser.add_argument("--timeout", type=int, default=120)

    render_parser = commands.add_parser(
        "render", help="Render JSON values into the original DOCX"
    )
    render_parser.add_argument("source")
    render_parser.add_argument("data", help="JSON containing tokens and/or elements")
    render_parser.add_argument("output")
    render_parser.add_argument(
        "--profile", help="Profile JSON whose source SHA-256 must match"
    )

    template_validate_parser = commands.add_parser(
        "template-validate",
        help="Validate a semantic template contract against an extracted bundle",
    )
    template_validate_parser.add_argument("bundle")
    template_validate_parser.add_argument("contract")

    template_preview_parser = commands.add_parser(
        "template-preview",
        help="Preview canonical migration discovery population without modifying a DOCX",
    )
    template_preview_parser.add_argument("bundle")
    template_preview_parser.add_argument("contract")
    template_preview_parser.add_argument("discovery")
    template_preview_parser.add_argument("--output")

    template_patch_parser = commands.add_parser(
        "template-patch",
        help="Compile canonical migration discovery data into a validated patch",
    )
    template_patch_parser.add_argument("bundle")
    template_patch_parser.add_argument("contract")
    template_patch_parser.add_argument("discovery")
    template_patch_parser.add_argument("output")

    template_populate_parser = commands.add_parser(
        "template-populate",
        help="Populate a bundle from canonical migration discovery data",
    )
    template_populate_parser.add_argument("bundle")
    template_populate_parser.add_argument("contract")
    template_populate_parser.add_argument("discovery")
    template_populate_parser.add_argument("output")
    template_populate_parser.add_argument("--patch-output")
    template_populate_parser.add_argument("--fidelity-output")

    args = parser.parse_args(argv)
    if args.command == "profile":
        profile = profile_document(args.source)
        Path(args.output).write_text(
            json.dumps(profile, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        print(
            json.dumps(
                {"profile": args.output, "summary": profile["summary"]}, indent=2
            )
        )
        return 0

    if args.command == "extract":
        limits = PackageLimits(
            max_entries=args.max_entries,
            max_entry_size=args.max_entry_size,
            max_total_uncompressed_size=args.max_total_size,
            max_compression_ratio=args.max_compression_ratio,
        )
        profile = export_bundle(args.source, args.bundle, limits)
        print(
            json.dumps(
                {
                    "bundle": args.bundle,
                    "entries": profile["package"]["entry_count"],
                    "summary": profile["summary"],
                },
                indent=2,
            )
        )
        return 0

    if args.command.startswith("template-"):
        from .template_contract import (
            compile_migration_template_patch,
            populate_migration_template,
            preview_migration_template_population,
            validate_template_contract,
        )

        bundle = Path(args.bundle)
        contract = _load_json(args.contract)
        profile = _load_json(bundle / "profile.json")
        if args.command == "template-validate":
            normalized = validate_template_contract(profile, contract)
            result = {
                "valid": True,
                "document_type": normalized["document_type"],
                "contract_version": normalized["contract_version"],
                "field_count": len(normalized["fields"]),
                "source_sha256": normalized["binding"]["source_sha256"],
            }
            print(json.dumps(result, indent=2))
            return 0
        discovery = _load_json(args.discovery)
        if args.command == "template-preview":
            result = preview_migration_template_population(bundle, contract, discovery)
            if args.output:
                Path(args.output).write_text(
                    json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8"
                )
            print(json.dumps(result, indent=2, ensure_ascii=False))
            return 0 if result["valid"] else 2
        if args.command == "template-patch":
            patch = compile_migration_template_patch(profile, contract, discovery)
            Path(args.output).write_text(
                json.dumps(patch, indent=2, ensure_ascii=False), encoding="utf-8"
            )
            print(
                json.dumps(
                    {"valid": True, "patch": args.output},
                    indent=2,
                )
            )
            return 0
        result = populate_migration_template(
            bundle,
            args.output,
            contract,
            discovery,
            patch_output=args.patch_output,
            fidelity_output=args.fidelity_output,
        )
        print(json.dumps(result, indent=2, ensure_ascii=False))
        return 0

    if args.command == "fidelity-report":
        patch = _load_json(args.patch) if args.patch else None
        report = write_fidelity_report(
            args.output,
            args.original,
            args.generated,
            bundle=args.bundle,
            patch=patch,
            visual_report=args.visual_report,
        )
        print(json.dumps(report, indent=2))
        return 0 if report["passed"] else 2

    if args.command == "visual":
        from .visual_regression import _page_set, verify_visuals

        try:
            report = verify_visuals(
                args.original,
                args.generated,
                args.output_dir,
                renderer=args.renderer,
                timeout=args.timeout,
                difference_threshold=args.difference_threshold,
                expected_difference_threshold=args.expected_difference_threshold,
                pixel_threshold=args.pixel_threshold,
                dimension_tolerance_pt=args.dimension_tolerance_pt,
                expected_changed_pages=_page_set(args.expected_pages),
                no_op=args.no_op,
                dpi=args.dpi,
            )
        except RuntimeError as error:
            print(
                json.dumps({"available": False, "error": str(error)}, indent=2),
                file=sys.stderr,
            )
            return 3
        print(json.dumps(report, indent=2))
        return 0 if report["passed"] else 2

    data = _load_json(args.data)
    if args.command == "preview":
        print(json.dumps(preview_bundle(args.bundle, data), indent=2))
        return 0
    if args.command == "render-bundle":
        print(json.dumps(render_bundle(args.bundle, args.output, data), indent=2))
        return 0

    expected_hash = None
    if args.profile:
        expected_hash = _load_json(args.profile).get("source", {}).get("sha256")
    print(
        json.dumps(
            render_document(
                args.source, args.output, data, verify_sha256=expected_hash
            ),
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
